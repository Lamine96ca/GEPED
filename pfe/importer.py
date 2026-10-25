"""
Import en masse des inscriptions (Excel .xlsx ou CSV) par la scolarité.

Le fichier est préparé par la scolarité : UNE LIGNE = UNE INSCRIPTION. Le système importe
les lignes telles quelles, sans rien déduire. Il vérifie seulement :
  - les colonnes et valeurs obligatoires : matricule, prenom, nom, programme, promotion,
    session, licence (une inscription est toujours dans une licence) ;
  - que le programme, la promotion, la session et la licence existent en base ;
  - qu'il n'y a pas de doublon : une même inscription (matricule, session, licence) n'est
    enregistrée qu'une fois. Déjà en base ou répétée dans le fichier, la ligne est IGNORÉE
    (ce n'est pas une erreur) et l'import continue avec les autres lignes ;
  - les lignes vides sont ignorées.

Un étudiant peut donc apparaître sur plusieurs lignes (une par année). Un matricule inconnu
crée l'étudiant (avec les informations de sa première ligne) ; un étudiant déjà en base
n'est pas modifié (seul son statut préinscrit est retiré si l'import n'est pas une
préinscription). Exception : en PRÉINSCRIPTION, un étudiant retrouvé par son INE voit sa
fiche mise à jour avec sa première ligne (valeurs renseignées seulement), sans que son
MATRICULE ni son statut ne changent. Tout passe par bulk_create / bulk_update.

Matricules : les INSCRITS ont déjà le leur, importé tel quel (colonne obligatoire). Pour une
PRÉINSCRIPTION, le matricule n'est pas demandé : il est généré à partir de la session, du
code du programme (programme du fichier rapproché de la base), d'un compteur et de la
première lettre du prénom (voir GenerateurMatricules). Un préinscrit sans matricule est
reconnu par son INE, sinon son e-mail, sinon nom + prénom + date de naissance, pour ne pas
le créer deux fois.
"""
import csv
import datetime
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation

from django.db import transaction

from .models import Etudiant, Inscription, Licence, Promotion, Service, Session

OBLIGATOIRES = ['matricule', 'prenom', 'nom', 'programme', 'promotion', 'session', 'licence']

COLONNES_IMPORT = OBLIGATOIRES + [
    'cycle', 'sexe', 'date_de_naissance', 'lieu_de_naissance', 'nationalite', 'telephone', 'email',
    'pere', 'mere', 'tuteur', 'contact_de_tuteur', 'adresse', 'ine', 'pvbac', 'annee_obtention_bac',
    'annee_orientation', 'option_de_bac', 'moyenne_au_bac', 'ecole_origine', 'centre_de_bac', 'ville',
]

# En-têtes alternatifs rencontrés dans les fichiers existants
ALIAS = {
    'natinalite': 'nationalite',
    'prenoms': 'prenom',
    'noms': 'nom',
    'filiere': 'programme',
    'annee': 'session',
    'annee_academique': 'session',
    'niveau': 'licence',
    'mail': 'email',
    'e_mail': 'email',
    'date_naissance': 'date_de_naissance',
    'lieu_naissance': 'lieu_de_naissance',
}

CHAMPS_ETUDIANT = [
    'matricule', 'prenom', 'nom', 'sexe', 'date_de_naissance', 'lieu_de_naissance',
    'nationalite', 'telephone', 'email', 'pere', 'mere', 'tuteur', 'contact_de_tuteur',
    'adresse', 'ine', 'pvbac', 'annee_obtention_bac', 'annee_orientation', 'option_de_bac',
    'moyenne_au_bac', 'ecole_origine', 'centre_de_bac', 'ville',
]

LONGUEURS = {f.name: f.max_length for f in Etudiant._meta.fields if getattr(f, 'max_length', None)}

DOMAINE_EMAIL_GENERE = 'etudiant.gped'
MAX_ERREURS_AFFICHEES = 300


def normaliser_entete(texte):
    texte = unicodedata.normalize('NFKD', str(texte or '')).encode('ascii', 'ignore').decode()
    texte = texte.strip().lower().replace(' ', '_').replace('-', '_').replace("'", '_')
    return ALIAS.get(texte, texte)


# ---------------------------------------------------------------- lecture

def lire_lignes(fichier):
    """Renvoie (entetes, itérateur de tuples). Accepte .xlsx et .csv."""
    nom = fichier.name.lower()
    if nom.endswith('.xlsx'):
        lignes = _feuille_xlsx(fichier)
    elif nom.endswith('.csv'):
        texte = fichier.read().decode('utf-8-sig', errors='replace')
        try:
            dialecte = csv.Sniffer().sniff(texte[:4096], delimiters=';,\t')
        except csv.Error:
            dialecte = csv.excel
            dialecte.delimiter = ';'
        lignes = csv.reader(io.StringIO(texte), dialecte)
    else:
        raise ValueError("Format non pris en charge : utilisez un fichier .xlsx ou .csv.")

    entetes = next(lignes, None)
    if not entetes:
        raise ValueError("Le fichier est vide.")
    return [normaliser_entete(e) for e in entetes], lignes


_XML_INTERDITS = re.compile(rb'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def reparer_xlsx(contenu):
    """
    Répare en mémoire un .xlsx produit par un autre logiciel qu'Excel : caractères invisibles interdits
    en XML (ex. \\x02 collé dans un intitulé) remplacés par une espace, et attribut mal orthographié
    « biltinId » du fichier des styles. Renvoie (octets réparés, nombre de corrections).
    """
    import zipfile
    entree, sortie, corrections = zipfile.ZipFile(io.BytesIO(contenu)), io.BytesIO(), 0
    with zipfile.ZipFile(sortie, 'w', zipfile.ZIP_DEFLATED) as z:
        for info in entree.infolist():
            brut = entree.read(info.filename)
            if info.filename.endswith(('.xml', '.rels')):
                brut, n = _XML_INTERDITS.subn(b' ', brut)
                corrections += n
                if info.filename == 'xl/styles.xml':
                    brut, n = re.subn(rb'\bbiltinId=', b'builtinId=', brut)
                    corrections += n
            z.writestr(info, brut)
    return sortie.getvalue(), corrections


def _feuille_xlsx(fichier):
    """Lignes de la feuille active ; un fichier légèrement endommagé est réparé automatiquement."""
    import openpyxl
    fichier.seek(0)
    contenu = fichier.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
        fichier.gped_reparations = 0
    except Exception:
        try:
            repare, corrections = reparer_xlsx(contenu)
            wb = openpyxl.load_workbook(io.BytesIO(repare), read_only=True, data_only=True)
        except Exception:
            raise ValueError("Ce fichier Excel est endommagé et n'a pas pu être lu. Ouvrez-le dans Excel, faites "
                             "« Enregistrer sous » au format .xlsx, puis réessayez.")
        fichier.gped_reparations = corrections         # signalé dans le résultat de l'import
    return wb.active.iter_rows(values_only=True)


# ---------------------------------------------------------------- conversion

def texte(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def en_date(v):
    if v in (None, ''):
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    s = texte(v)
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d.%m.%Y', '%Y/%m/%d'):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"date invalide « {s} » (format attendu JJ/MM/AAAA)")


def en_sexe(v):
    s = texte(v).upper()
    return 'F' if s.startswith('F') else 'M'


def en_moyenne(v):
    s = texte(v).replace(',', '.')
    if not s:
        return None
    try:
        d = Decimal(s).quantize(Decimal('0.01'))
    except InvalidOperation:
        raise ValueError(f"moyenne au bac invalide « {s} »")
    if not (0 <= d < 100):
        raise ValueError(f"moyenne au bac hors limites « {s} »")
    return d


def cle(texte_brut):
    """Clé de rapprochement : sans accents, majuscules, espaces simples."""
    s = unicodedata.normalize('NFKD', texte(texte_brut)).encode('ascii', 'ignore').decode()
    return ' '.join(s.upper().split())


def cle_compacte(v):
    """Clé sans espaces : « p 18 » = « P18 », « Licence 4 » = « LICENCE4 »."""
    return cle(v).replace(' ', '')


def cle_session(v):
    """« 2025 - 2026 », « 2025-2026 », « 2025/2026 » → « 2025-2026 »."""
    annees = re.findall(r'\d{4}', texte(v))
    return '-'.join(annees) if annees else cle_compacte(v)


def cle_licence(v):
    """« Licence 4 », « LICENCE 4 », « L4 » → « LICENCE4 »."""
    k = cle_compacte(v)
    m = re.fullmatch(r'L(\d)', k)
    return f'LICENCE{m.group(1)}' if m else k


# Cycle : code ou libellé, sans tenir compte des accents ni de la casse
CYCLES = {
    'DUT': 'DUT',
    'TC': 'Tronc commun', 'TRONC COMMUN': 'Tronc commun',
    'LIC': 'Licence', 'LICENCE': 'Licence', 'L': 'Licence',
    'ING': 'Ingénieur', 'INGENIEUR': 'Ingénieur',
    'MAS': 'Master', 'MASTER': 'Master', 'M': 'Master',
    'DOC': 'Doctorat', 'DOCTORAT': 'Doctorat',
}


def en_cycle(v):
    s = cle(v)
    if not s:
        return None
    if s not in CYCLES:
        raise ValueError(f"cycle inconnu « {texte(v)} » (DUT, Tronc commun, Licence, Ingénieur, Master, Doctorat)")
    return CYCLES[s]


# ---------------------------------------------------------------- import

class RapportImport:
    def __init__(self):
        self.lignes = 0
        self.etudiants_crees = 0
        self.etudiants_existants = 0
        self.etudiants_mis_a_jour = 0   # préinscription : fiches retrouvées par l'INE, mises à jour (matricule conservé)
        self.inscriptions_creees = 0
        self.erreurs = []           # [(numéro de ligne, message)] : lignes rejetées
        self.doublons = []          # [(numéro de ligne, message)] : inscriptions déjà présentes, ignorées
        self.matricules_generes = []  # [(numéro de ligne, étudiant, programme, matricule)] : préinscrits
        self.simulation = False
        self.preinscription = False
        self.colonnes_ignorees = []
        self.reparations = 0        # caractères invisibles corrigés dans un fichier endommagé

    @property
    def nb_erreurs(self):
        return len(self.erreurs)

    @property
    def nb_doublons(self):
        return len(self.doublons)

    def erreur(self, ligne, message):
        self.erreurs.append((ligne, message))


def _index_programmes():
    """Programmes indexés par nom et par code (insensible à la casse et aux accents)."""
    index = {}
    for s in Service.objects.filter(programme=True):
        index[cle(s.nom)] = s
        if s.code:
            index.setdefault(cle(s.code), s)
    return index


def _lire_etudiant(ligne):
    """Informations de l'étudiant présentes sur la ligne, converties et contrôlées."""
    data = {}
    for champ in CHAMPS_ETUDIANT:
        if champ not in ligne:
            continue
        v = ligne[champ]
        if champ == 'date_de_naissance':
            data[champ] = en_date(v)
        elif champ == 'sexe':
            data[champ] = en_sexe(v)
        elif champ == 'moyenne_au_bac':
            data[champ] = en_moyenne(v)
        else:
            s = texte(v)
            if champ in ('annee_obtention_bac', 'annee_orientation'):
                s = s[:4]
            if champ in LONGUEURS and len(s) > LONGUEURS[champ]:
                raise ValueError(f"« {champ} » dépasse {LONGUEURS[champ]} caractères")
            data[champ] = s or None
    return data


def code_lettre(prenom):
    """Code de la 1re lettre du prénom, sans accent : ord(lettre) - 30 (A = 35 … Z = 60)."""
    s = unicodedata.normalize('NFKD', texte(prenom)).encode('ascii', 'ignore').decode().strip()
    lettre = s[0].upper() if s else 'X'
    return str(ord(lettre) - 30)


class GenerateurMatricules:
    """
    Matricule d'un PRÉINSCRIT = SESSION + PROGRAMME + COMPTEUR + LETTRE
      SESSION   : 2 derniers caractères de l'année de la session (« 2025 - 2026 » → 26)
      PROGRAMME : code du programme en base (complété à 2 chiffres : 4 → 04)
      COMPTEUR  : nombre d'inscriptions du programme + 1, sur 4 chiffres, puis +1 à chaque
                  matricule généré dans l'import
      LETTRE    : code de la 1re lettre du prénom (ord - 30)
    Ex. Mamadou, GENIE INFORMATIQUE (04), 2025 - 2026, 2399 inscriptions → 26 04 2400 47.
    Un matricule déjà pris (en base ou dans le fichier) fait passer au compteur suivant.
    """
    def __init__(self, deja_pris=()):
        self.compteurs = {}                # id programme -> dernier compteur utilisé
        self.pris = set(deja_pris)
        self._prefixes_lus = set()

    def generer(self, session, programme, prenom):
        session_code = str(session.annee)[-2:] if session and session.annee else '00'
        service_code = str(programme.code or '00').zfill(2)
        prefixe = session_code + service_code
        if prefixe not in self._prefixes_lus:
            self.pris.update(Etudiant.objects.filter(matricule__startswith=prefixe)
                             .values_list('matricule', flat=True))
            self._prefixes_lus.add(prefixe)
        if programme.pk not in self.compteurs:
            self.compteurs[programme.pk] = Inscription.objects.filter(filiere_id=programme.pk).count()
        lettre = code_lettre(prenom)
        while True:
            self.compteurs[programme.pk] += 1
            matricule = f"{prefixe}{str(self.compteurs[programme.pk]).zfill(4)}{lettre}"
            if matricule not in self.pris:
                self.pris.add(matricule)
                return matricule


def _cle_etudiant(data):
    """
    Identifie l'étudiant d'une ligne : par son matricule ; pour un préinscrit sans matricule,
    par son INE, sinon son e-mail, sinon nom + prénom + date de naissance (None si rien ne
    l'identifie).
    """
    if data.get('matricule'):
        return ('matricule', data['matricule'])
    if data.get('ine'):
        return ('ine', data['ine'])
    if data.get('email'):
        return ('email', data['email'])
    if data.get('date_de_naissance'):
        return ('personne', cle(data['nom']), cle(data['prenom']), data['date_de_naissance'])
    return None


# Préinscription d'un étudiant déjà en base (même INE) : sa fiche est mise à jour, son matricule jamais
CHAMPS_NON_MIS_A_JOUR = {'matricule'}


def _mettre_a_jour_existants(valides, existants):
    """
    Met à jour la fiche des étudiants déjà en base avec leur PREMIÈRE ligne du fichier :
    seules les valeurs renseignées sont reprises (une case vide n'efface rien), le matricule
    n'est jamais modifié et l'e-mail seulement s'il n'est pas déjà celui d'un autre étudiant.
    Le statut (inscrit / préinscrit) ne change pas. Renvoie le nombre de fiches modifiées.
    """
    premieres = {}                      # pk -> (étudiant, données de sa première ligne)
    for _, data, _, k in valides:
        if k in existants:
            premieres.setdefault(existants[k].pk, (existants[k], data))
    emails = list({d['email'] for _, d in premieres.values() if d.get('email')})
    proprietaires = {}                  # e-mail -> pk de l'étudiant qui l'utilise
    for i in range(0, len(emails), 500):
        for pk, email in Etudiant.objects.filter(email__in=emails[i:i + 500]).values_list('pk', 'email'):
            proprietaires[email.lower()] = pk
    modifies, champs_modifies = [], set()
    for e, data in premieres.values():
        champs = []
        for champ, valeur in data.items():
            if champ in CHAMPS_NON_MIS_A_JOUR or valeur in (None, ''):
                continue
            if champ == 'email' and proprietaires.setdefault(valeur, e.pk) != e.pk:
                continue                # e-mail déjà pris par un autre étudiant : on garde l'ancien
            if getattr(e, champ) != valeur:
                setattr(e, champ, valeur)
                champs.append(champ)
        if champs:
            modifies.append(e)
            champs_modifies.update(champs)
    if modifies:
        Etudiant.objects.bulk_update(modifies, sorted(champs_modifies), batch_size=500)
    return len(modifies)


def importer_etudiants(fichier, simulation=False, cycle=None, preinscription=False):
    """
    Une ligne du fichier = une inscription (matricule, session, licence, programme, promotion).

    `cycle` : cycle appliqué aux inscriptions (une colonne « cycle » du fichier l'emporte).
    `preinscription` : True = préinscrits : le matricule n'est pas demandé, il est GÉNÉRÉ
                       (voir GenerateurMatricules) ; False = inscrits : ils ont déjà leur
                       matricule, importé tel quel.
    """
    rapport = RapportImport()
    rapport.simulation = simulation
    rapport.preinscription = preinscription
    entetes, lignes = lire_lignes(fichier)
    rapport.reparations = getattr(fichier, 'gped_reparations', 0)

    obligatoires = [c for c in OBLIGATOIRES if not (preinscription and c == 'matricule')]
    rapport.colonnes_ignorees = [e for e in entetes if e and e not in COLONNES_IMPORT]
    manquantes = [c for c in obligatoires if c not in entetes]
    if manquantes:
        raise ValueError("Colonnes obligatoires manquantes : " + ", ".join(f"« {c} »" for c in manquantes) + ".")

    programmes = _index_programmes()
    promotions = {cle_compacte(p.nom): p for p in Promotion.objects.all()}
    sessions = {cle_session(s.annee): s for s in Session.objects.all()}
    licences = {cle_licence(l.nom): l for l in Licence.objects.all()}

    # 1) Lecture et contrôle de chaque ligne ---------------------------------
    valides = []                       # (n° ligne, données étudiant, données inscription, clé étudiant)
    vues = {}                          # (clé étudiant, session, licence) -> n° de ligne
    for numero, brut in enumerate(lignes, start=2):
        if brut is None or all(texte(c) == '' for c in brut):
            continue                   # ligne vide : ignorée
        rapport.lignes += 1
        ligne = {entetes[i]: brut[i] for i in range(min(len(entetes), len(brut))) if entetes[i]}
        try:
            vides = [c for c in obligatoires if not texte(ligne.get(c))]
            if vides:
                raise ValueError("valeur manquante : " + ", ".join(vides))
            data = _lire_etudiant(ligne)
            if data.get('email'):
                data['email'] = data['email'].lower()
            if data.get('ine'):
                data['ine'] = data['ine'].upper()

            # Le programme du fichier est rapproché de celui de la base (nom ou code)
            programme = programmes.get(cle(ligne['programme']))
            if programme is None:
                raise ValueError(f"programme inconnu « {texte(ligne['programme'])} »")
            promotion = promotions.get(cle_compacte(ligne['promotion']))
            if promotion is None:
                raise ValueError(f"promotion inconnue « {texte(ligne['promotion'])} »")
            session = sessions.get(cle_session(ligne['session']))
            if session is None:
                raise ValueError(f"session inconnue « {texte(ligne['session'])} »")
            licence = licences.get(cle_licence(ligne['licence']))
            if licence is None:
                raise ValueError(f"licence inconnue « {texte(ligne['licence'])} »")
            insc = {'session': session, 'licence': licence, 'filiere': programme, 'promotion': promotion,
                    'cycle': en_cycle(ligne.get('cycle')) or cycle}
        except ValueError as e:
            rapport.erreur(numero, str(e))
            continue

        cle_etu = _cle_etudiant(data) or ('ligne', numero)
        doublon = (cle_etu, session.pk, licence.pk)
        if doublon in vues:
            # Doublon : la ligne est ignorée (pas une erreur) et l'import continue
            qui = data.get('matricule') or f"{data['prenom']} {data['nom']}"
            rapport.doublons.append((numero, f"{qui} · {licence} · {session} : déjà présente à la ligne "
                                             f"{vues[doublon]} du fichier"))
            continue
        vues[doublon] = numero
        valides.append((numero, data, insc, cle_etu))

    if not valides:
        return rapport

    # 2) Écriture en base ---------------------------------------------------
    with transaction.atomic():
        # Étudiants déjà en base, retrouvés par la même clé que dans le fichier. Préinscription :
        # l'INE passe en premier (un INE connu = cet étudiant, même si le fichier donne un matricule)
        candidats = {}                 # clé de la ligne -> identifiants à essayer, dans l'ordre
        for _, data, _, k in valides:
            if k not in candidats:
                ine = ('ine', data['ine']) if preinscription and data.get('ine') else None
                candidats[k] = [ine, k] if ine and ine != k else [k]
        cles = {c for liste_cles in candidats.values() for c in liste_cles}
        matricules = [k[1] for k in cles if k[0] == 'matricule']
        ines = [k[1] for k in cles if k[0] == 'ine']
        emails = [k[1] for k in cles if k[0] == 'email']
        dates = list({k[3] for k in cles if k[0] == 'personne'})
        trouves = {}                   # identifiant -> Etudiant
        for i in range(0, len(matricules), 500):
            for e in Etudiant.objects.filter(matricule__in=matricules[i:i + 500]):
                trouves[('matricule', e.matricule)] = e
        for i in range(0, len(ines), 500):
            for e in Etudiant.objects.filter(ine__in=ines[i:i + 500]):
                trouves[('ine', e.ine)] = e
        for i in range(0, len(emails), 500):
            for e in Etudiant.objects.filter(email__in=emails[i:i + 500]):
                trouves[('email', e.email.lower())] = e
        for i in range(0, len(dates), 500):
            for e in Etudiant.objects.filter(date_de_naissance__in=dates[i:i + 500]):
                k = ('personne', cle(e.nom), cle(e.prenom), e.date_de_naissance)
                if k in cles:
                    trouves[k] = e
        existants = {}                 # clé étudiant -> Etudiant
        par_ine = set()                # clés des lignes dont l'étudiant a été retrouvé par son INE
        for k, liste_cles in candidats.items():
            c = next((c for c in liste_cles if c in trouves), None)
            if c is not None:
                existants[k] = trouves[c]
                if c[0] == 'ine':
                    par_ine.add(k)
        rapport.etudiants_existants = len({e.pk for e in existants.values()})

        # Nouveaux étudiants : informations de leur première ligne ; matricule généré pour un
        # préinscrit qui n'en a pas ; l'e-mail doit être libre
        generateur = GenerateurMatricules(deja_pris=matricules)
        nouveaux, generes = {}, {}     # clé -> données ; clé -> (n° de ligne, programme) si matricule généré
        for numero, data, insc, k in valides:
            if k in existants or k in nouveaux:
                continue
            if not data.get('matricule'):
                data['matricule'] = generateur.generer(insc['session'], insc['filiere'], data['prenom'])
                generes[k] = (numero, insc['filiere'])
            data['email'] = (data.get('email') or f"{data['matricule']}@{DOMAINE_EMAIL_GENERE}").lower()
            data['preinscrit'] = preinscription
            if preinscription:              # préinscrits : boursiers et en formation initiale par défaut
                data.setdefault('boursier', 'boursier')
                data.setdefault('type_formation', 'initiale')
            nouveaux[k] = data
        emails_nouveaux = [d['email'] for d in nouveaux.values()]
        pris = set()
        for i in range(0, len(emails_nouveaux), 500):
            pris.update(e.lower() for e in Etudiant.objects.filter(email__in=emails_nouveaux[i:i + 500])
                        .values_list('email', flat=True))
        rejetes, vus_email = set(), {}
        for k, d in nouveaux.items():
            if d['email'] in pris or d['email'] in vus_email:
                rejetes.add(k)
            vus_email.setdefault(d['email'], k)

        crees = {k: d for k, d in nouveaux.items() if k not in rejetes}
        Etudiant.objects.bulk_create([Etudiant(**d) for d in crees.values()], batch_size=500)
        rapport.etudiants_crees = len(crees)
        rapport.matricules_generes = sorted(
            (generes[k][0], f"{d['prenom']} {d['nom']}", generes[k][1], d['matricule'])
            for k, d in crees.items() if k in generes)

        # bulk_create ne renvoie pas les id sous MySQL : on les relit par matricule
        ids = {k: e.pk for k, e in existants.items()}
        par_matricule = {d['matricule']: k for k, d in crees.items()}
        a_relire = list(par_matricule)
        for i in range(0, len(a_relire), 500):
            for m, pk in Etudiant.objects.filter(matricule__in=a_relire[i:i + 500]).values_list('matricule', 'id'):
                ids[par_matricule[m]] = pk
        etiquettes = {k: e.matricule or f"{e.prenom} {e.nom}" for k, e in existants.items()}
        etiquettes.update({k: d['matricule'] for k, d in crees.items()})

        # Inscriptions déjà en base pour ces étudiants : (étudiant, session, licence)
        etu_ids = list(set(ids.values()))
        deja = set()
        for i in range(0, len(etu_ids), 500):
            deja.update(Inscription.objects.filter(etudiant_id__in=etu_ids[i:i + 500])
                        .values_list('etudiant_id', 'session_id', 'licence_id'))

        a_creer = []
        for numero, data, insc, k in valides:
            if k in rejetes:
                rapport.erreur(numero, f"e-mail {nouveaux[k]['email']} déjà utilisé par un autre étudiant")
                continue
            etu_id = ids[k]
            if (etu_id, insc['session'].pk, insc['licence'].pk) in deja:
                rapport.doublons.append((numero, f"{etiquettes[k]} · {insc['licence']} · {insc['session']} : "
                                                 f"déjà inscrit(e) dans gePed"))
                continue
            a_creer.append(Inscription(etudiant_id=etu_id, **insc))
            deja.add((etu_id, insc['session'].pk, insc['licence'].pk))   # même étudiant sur deux lignes
        Inscription.objects.bulk_create(a_creer, batch_size=500)
        rapport.inscriptions_creees = len(a_creer)
        rapport.erreurs.sort()
        rapport.doublons.sort()

        # Préinscription : un étudiant retrouvé par son INE voit sa fiche mise à jour, matricule conservé
        if preinscription and par_ine:
            rapport.etudiants_mis_a_jour = _mettre_a_jour_existants(
                valides, {k: e for k, e in existants.items() if k in par_ine})

        # Import « inscription » : seule modification d'un étudiant existant, il n'est plus préinscrit
        if not preinscription and existants:
            pks = [e.pk for e in existants.values() if e.preinscrit]
            Etudiant.objects.filter(pk__in=pks).update(preinscrit=False)

        if simulation:
            transaction.set_rollback(True)

    return rapport


def modele_import_xlsx():
    """Fichier Excel vide avec les en-têtes attendus (et une ligne d'exemple)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Inscriptions'
    ws.append(COLONNES_IMPORT)
    ws.append(['2601000001', 'Mamadou', 'DIALLO', 'GENIE INFORMATIQUE', 'P20', '2025 - 2026', 'LICENCE 1',
               'Licence', 'M', '31/10/2005', 'Conakry', 'Guinéenne', '620000000', '', 'Nom du père',
               'Nom de la mère', '', '', 'Conakry', '', '', '2025', '2025', 'Sciences mathématiques', 12.5,
               '', '', 'Conakry'])
    for cell in ws[1]:
        obligatoire = cell.value in OBLIGATOIRES
        cell.font = Font(bold=True, color='1F2937' if obligatoire else 'FFFFFF')
        cell.fill = PatternFill('solid', fgColor='FBBD08' if obligatoire else '41627C')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, len(str(col[0].value)) + 4)
    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()
