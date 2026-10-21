"""
Import des soutenances (Excel .xlsx ou CSV), notamment des anciens étudiants : scolarité et
chefs de programme. Une ligne = la soutenance d'un étudiant. Colonnes repérées par leur nom :
  - obligatoires : matricule, programme, session, promotion ;
  - facultatives : cycle (vide : celui du programme), note, note_session, projet, theme,
    encadreur, chef_departement, maitre_stage, entreprise, president_jury, examinateur,
    rapporteur, date, heure, salle, statut.

Contrôles : étudiant connu (pas préinscrit), programme, session et promotion existants, notes
sur 10 (« ABS » = 0). Une ligne répétée (même étudiant, même cycle) est ignorée (doublon).
Les lignes qui portent le même projet (même programme, session et promotion) forment un seul
projet de soutenance ; sans nom de projet mais avec un thème, un encadreur… un projet est créé
pour l'étudiant seul.

Une soutenance déjà enregistrée est mise à jour (une case vide garde l'ancienne valeur).
Les notes importées sont VERROUILLÉES d'office : un nouvel import ne peut plus les changer
(seules les informations du projet de soutenance changent) ; la scolarité les modifie depuis
la fiche, avec un motif. Chef de programme : une ligne d'un autre programme → tout est refusé.
"""
import datetime
import io
import re
from collections import Counter

from django.db import transaction
from django.utils import timezone

from .import_matieres import ImportRefuse
from .importer import _index_programmes, cle, cle_compacte, cle_session, en_cycle, en_date, lire_lignes, texte
from .models import Etudiant, ProjetSoutenance, Promotion, Session, Soutenance
from .notes import en_note, recalculer
from .soutenances import CHAMPS_NOTE, CYCLES_SOUTENANCE, enregistrer_notes, notes_de

OBLIGATOIRES = ['matricule', 'programme', 'session', 'promotion']
CHAMPS_PROJET = ['theme', 'encadreur', 'chef_departement', 'maitre_stage', 'entreprise', 'president_jury',
                 'examinateur', 'rapporteur', 'salle']
COLONNES = OBLIGATOIRES + ['cycle', 'note', 'note_session', 'projet'] + CHAMPS_PROJET + ['date', 'heure', 'statut']
ALIAS = {
    'note_soutenance': 'note', 'note_de_soutenance': 'note', 'moyenne': 'note', 'moyenne_soutenance': 'note',
    'note_de_session': 'note_session', 'rattrapage': 'note_session', 'session_rattrapage': 'note_session',
    'groupe': 'projet', 'projet_de_soutenance': 'projet', 'intitule': 'projet', 'numero_projet': 'projet',
    'theme_de_soutenance': 'theme', 'theme_du_memoire': 'theme', 'sujet': 'theme', 'memoire': 'theme',
    'encadrant': 'encadreur', 'directeur_de_memoire': 'encadreur',
    'chef_de_departement': 'chef_departement', 'chef_dept': 'chef_departement',
    'maitre_de_stage': 'maitre_stage', 'structure': 'entreprise', 'structure_d_accueil': 'entreprise',
    'president': 'president_jury', 'president_du_jury': 'president_jury',
    'date_soutenance': 'date', 'date_de_soutenance': 'date', 'heure_de_soutenance': 'heure',
}
STATUTS = {cle(lib): code for code, lib in ProjetSoutenance.STATUTS} | {cle(code): code for code, _ in
                                                                         ProjetSoutenance.STATUTS}
LONGUEURS = {f.name: f.max_length for f in ProjetSoutenance._meta.fields if getattr(f, 'max_length', None)}


def en_heure(v):
    if v in (None, ''):
        return None
    if isinstance(v, datetime.datetime):
        return v.time().replace(second=0, microsecond=0)
    if isinstance(v, datetime.time):
        return v.replace(second=0, microsecond=0)
    if isinstance(v, float) and 0 <= v < 1:                       # heure Excel (fraction de jour)
        minutes = round(v * 24 * 60)
        return datetime.time(minutes // 60, minutes % 60)
    m = re.fullmatch(r'(\d{1,2})\s*[hH:.]\s*(\d{2})?', texte(v))
    if not m or int(m.group(1)) > 23 or int(m.group(2) or 0) > 59:
        raise ValueError(f"heure invalide « {texte(v)} » (ex. 09:30 ou 14h)")
    return datetime.time(int(m.group(1)), int(m.group(2) or 0))


class RapportSoutenances:
    def __init__(self):
        self.lignes = 0
        self.creees = 0
        self.mises_a_jour = 0
        self.inchangees = 0
        self.projets_crees = 0
        self.projets_mis_a_jour = 0
        self.erreurs = []
        self.doublons = []
        self.avertissements = []
        self.simulation = False
        self.colonnes_ignorees = []
        self.reparations = 0

    @property
    def nb_erreurs(self):
        return len(self.erreurs)

    @property
    def nb_doublons(self):
        return len(self.doublons)

    def erreur(self, ligne, message):
        self.erreurs.append((ligne, message))


def _lire(ligne, programmes, sessions, promotions):
    """Contrôle d'une ligne → dictionnaire prêt à enregistrer (lève ValueError)."""
    vides = [c for c in OBLIGATOIRES if not texte(ligne.get(c))]
    if vides:
        raise ValueError("valeur manquante : " + ", ".join(vides))
    programme = programmes.get(cle(ligne['programme']))
    if programme is None:
        raise ValueError(f"programme inconnu « {texte(ligne['programme'])} »")
    session = sessions.get(cle_session(ligne['session']))
    if session is None:
        raise ValueError(f"session inconnue « {texte(ligne['session'])} »")
    promotion = promotions.get(cle_compacte(ligne['promotion']))
    if promotion is None:
        raise ValueError(f"promotion inconnue « {texte(ligne['promotion'])} »")
    cycle = en_cycle(ligne.get('cycle')) or programme.cycle
    if not cycle:
        raise ValueError(f"cycle manquant : le programme {programme} n'en a pas, ajoutez la colonne « cycle »")
    notes = {}
    for c in CHAMPS_NOTE:
        try:
            notes[c] = en_note(ligne.get(c))
        except ValueError as e:
            raise ValueError(f"{c} : {e}")
    infos = {}
    for c in CHAMPS_PROJET:
        v = ' '.join(texte(ligne.get(c)).split())
        if len(v) > LONGUEURS.get(c, 10 ** 6):
            raise ValueError(f"{c} trop long ({len(v)} caractères, {LONGUEURS[c]} au plus)")
        if v:
            infos[c] = v
    if texte(ligne.get('date')):
        infos['date'] = en_date(ligne.get('date'))
    if texte(ligne.get('heure')):
        infos['heure'] = en_heure(ligne.get('heure'))
    if texte(ligne.get('statut')):
        statut = STATUTS.get(cle(ligne['statut']))
        if statut is None:
            raise ValueError(f"statut inconnu « {texte(ligne['statut'])} » (programmée, soutenue, reportée, ajournée)")
        infos['statut'] = statut
    intitule = ' '.join(texte(ligne.get('projet')).split())
    if len(intitule) > LONGUEURS['intitule']:
        raise ValueError(f"nom de projet trop long ({len(intitule)} caractères)")
    return {'matricule': texte(ligne['matricule']), 'programme': programme, 'session': session,
            'promotion': promotion, 'cycle': cycle, 'notes': notes, 'infos': infos, 'intitule': intitule}


def importer_soutenances(fichier, profil, utilisateur=None, simulation=False):
    rapport = RapportSoutenances()
    rapport.simulation = simulation
    entetes, lignes = lire_lignes(fichier)
    rapport.reparations = getattr(fichier, 'gped_reparations', 0)
    entetes = [ALIAS.get(e, e) for e in entetes]
    rapport.colonnes_ignorees = [e for e in entetes if e and e not in COLONNES]
    manquantes = [c for c in OBLIGATOIRES if c not in entetes]
    if manquantes:
        raise ValueError("Colonnes obligatoires manquantes : " + ", ".join(f"« {c} »" for c in manquantes) + ".")

    programmes = _index_programmes()
    sessions = {cle_session(s.annee): s for s in Session.objects.all()}
    promotions = {cle_compacte(p.nom): p for p in Promotion.objects.all()}
    autorises = None if profil.est_scolarite else set(profil.programme_ids)

    # 1) Lecture et contrôle des lignes ---------------------------------------------
    lues = []
    for numero, brut in enumerate(lignes, start=2):
        if brut is None or all(texte(c) == '' for c in brut):
            continue
        rapport.lignes += 1
        ligne = {entetes[i]: brut[i] for i in range(min(len(entetes), len(brut))) if entetes[i]}
        try:
            lues.append((numero, _lire(ligne, programmes, sessions, promotions)))
        except ValueError as e:
            rapport.erreur(numero, str(e))

    hors_perimetre = [(n, d['programme'].nom) for n, d in lues
                      if autorises is not None and d['programme'].pk not in autorises]
    if hors_perimetre:
        details = ', '.join(f"ligne {n} ({p})" for n, p in hors_perimetre[:10])
        if len(hors_perimetre) > 10:
            details += f"… et {len(hors_perimetre) - 10} autre(s)"
        raise ImportRefuse(f"Import refusé : le fichier contient des soutenances d'un programme qui n'est pas le "
                           f"vôtre : {details}. Aucune soutenance n'a été enregistrée.")

    # 2) Étudiants et doublons -----------------------------------------------------
    matricules = list({d['matricule'] for _, d in lues})
    etudiants = {}
    for i in range(0, len(matricules), 500):
        etudiants.update({e.matricule: e for e in Etudiant.objects.filter(matricule__in=matricules[i:i + 500])})
    retenues, vues = [], {}
    for numero, d in lues:
        e = etudiants.get(d['matricule'])
        if e is None:
            rapport.erreur(numero, f"matricule inconnu « {d['matricule']} »")
            continue
        if e.preinscrit:
            rapport.erreur(numero, f"{d['matricule']} est préinscrit : finalisez d'abord son inscription")
            continue
        identite = (e.pk, d['cycle'])
        if identite in vues:
            rapport.doublons.append((numero, f"{d['matricule']} · {d['cycle']} : déjà présent à la ligne "
                                             f"{vues[identite]} du fichier"))
            continue
        vues[identite] = numero
        d['etudiant'] = e
        retenues.append((numero, d))
    hors_cycle = Counter(d['cycle'] for _, d in retenues if d['cycle'] not in CYCLES_SOUTENANCE)
    for cycle, n in sorted(hors_cycle.items()):
        rapport.avertissements.append(f"{n} soutenance{'s' if n > 1 else ''} en {cycle} : la soutenance concerne "
                                      f"d'habitude les cycles Ingénieur et Master (lignes acceptées).")

    # 3) Écriture -----------------------------------------------------------------
    if retenues:
        with transaction.atomic():
            _enregistrer(retenues, entetes, rapport, utilisateur, getattr(fichier, 'name', ''))
            if simulation:
                transaction.set_rollback(True)
    rapport.erreurs.sort()
    rapport.doublons.sort()
    return rapport


def _enregistrer(retenues, entetes, rapport, utilisateur, nom_fichier):
    ids = [d['etudiant'].pk for _, d in retenues]
    existantes = {}
    for i in range(0, len(ids), 500):
        for s in Soutenance.objects.filter(etudiant_id__in=ids[i:i + 500]):
            existantes[(s.etudiant_id, s.cycle)] = s
    groupes = {(d['programme'].pk, d['session'].pk, d['promotion'].pk) for _, d in retenues}
    projets = {}
    for p in ProjetSoutenance.objects.filter(filiere_id__in={g[0] for g in groupes},
                                             session_id__in={g[1] for g in groupes}):
        projets[(cle(p.intitule), p.filiere_id, p.session_id, p.promotion_id)] = p
    projets_crees, projets_modifies = set(), set()

    def projet_de(d):
        e, infos = d['etudiant'], d['infos']
        intitule = d['intitule'] or (f"{e.nom} {e.prenom} ({e.matricule})"[:120] if infos else '')
        if not intitule:
            return None
        k = (cle(intitule), d['programme'].pk, d['session'].pk, d['promotion'].pk)
        p = projets.get(k)
        nouveau = p is None
        if nouveau:
            p = ProjetSoutenance(intitule=intitule, filiere=d['programme'], session=d['session'],
                                 promotion=d['promotion'], cycle=d['cycle'])
            if 'statut' not in infos and any(v is not None for v in d['notes'].values()):
                p.statut = ProjetSoutenance.SOUTENUE
            projets[k] = p
        change = False
        for c, v in infos.items():                           # une case vide garde l'ancienne valeur
            if getattr(p, c) != v:
                setattr(p, c, v)
                change = True
        if nouveau or change:
            p.save()
            (projets_crees if nouveau or p.pk in projets_crees else projets_modifies).add(p.pk)
        return p

    nouvelles, a_mettre_a_jour, changements = [], [], []
    for numero, d in retenues:
        e = d['etudiant']
        projet = projet_de(d)
        s = existantes.get((e.pk, d['cycle']))
        if s is None:
            s = Soutenance(etudiant=e, cycle=d['cycle'], filiere=d['programme'], session=d['session'],
                           promotion=d['promotion'], projet=projet, **d['notes'])
            s.calculer()
            nouvelles.append(s)
            continue
        modifiee = False
        for champ, valeur in (('filiere', d['programme']), ('session', d['session']), ('promotion', d['promotion'])):
            if getattr(s, f'{champ}_id') != valeur.pk:
                setattr(s, champ, valeur)
                modifiee = True
        if projet is not None and s.projet_id != projet.pk:
            s.projet = projet
            modifiee = True
        # Colonnes de notes absentes du fichier ou cases vides : on garde ce qui était enregistré
        valeurs = {c: d['notes'][c] if c in entetes and d['notes'][c] is not None else getattr(s, c)
                   for c in CHAMPS_NOTE}
        if valeurs != notes_de(s):
            if s.verrouillee:
                rapport.avertissements.append(f"Ligne {numero} : la note de {e.matricule} est verrouillée, elle "
                                              f"n'a pas été modifiée (seules les informations du projet changent).")
            else:
                changements.append((s, valeurs))
                modifiee = True
        if modifiee:
            a_mettre_a_jour.append(s)
        else:
            rapport.inchangees += 1
    rapport.projets_crees, rapport.projets_mis_a_jour = len(projets_crees), len(projets_modifies)

    Soutenance.objects.bulk_create(nouvelles, batch_size=500)
    if a_mettre_a_jour:
        Soutenance.objects.bulk_update(a_mettre_a_jour, ['filiere', 'session', 'promotion', 'projet'], batch_size=500)
    enregistrer_notes(changements, utilisateur, motif=f"Import du fichier {nom_fichier}".strip())
    rapport.creees, rapport.mises_a_jour = len(nouvelles), len(a_mettre_a_jour)
    # Notes importées : verrouillées d'office (seulement les soutenances du fichier)
    importees = {(d['etudiant'].pk, d['cycle']) for _, d in retenues}
    a_verrouiller = [pk for pk, etu, cycle in Soutenance.objects.filter(
        etudiant_id__in=ids, note_finale__isnull=False, verrouillee=False).values_list('pk', 'etudiant_id', 'cycle')
        if (etu, cycle) in importees]
    for i in range(0, len(a_verrouiller), 500):
        Soutenance.objects.filter(pk__in=a_verrouiller[i:i + 500]).update(
            verrouillee=True, verrouillee_le=timezone.now(), verrouillee_par=utilisateur)
    recalculer(ids)


def modele_soutenances_xlsx():
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Soutenances'
    ws.append(COLONNES)
    exemples = [
        ['2104000123', 'GENIE INFORMATIQUE', '2021 - 2022', 'P18', '', 8.5, '', 'Projet 1',
         'Gestion des inscriptions en ligne', 'Dr Mamadou BARRY', 'M. Ibrahima SOW', '', '', 'Pr Alpha DIALLO',
         'Dr Fatoumata CAMARA', 'M. Ousmane BAH', 'Amphi A', '15/07/2022', '09:00', 'soutenue'],
        ['2104000217', 'GENIE INFORMATIQUE', '2021 - 2022', 'P18', '', 8.5, '', 'Projet 1'],
        ['2004000328', 'ENERGETIQUE', '2020 - 2021', 'P17', '', 'ABS', 6],
    ]
    for ligne in exemples:
        ws.append(ligne)
    for cell in ws[1]:
        obligatoire = cell.value in OBLIGATOIRES
        cell.font = Font(bold=True, color='1F2937' if obligatoire else 'FFFFFF')
        cell.fill = PatternFill('solid', fgColor='FBBD08' if obligatoire else '41627C')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or '')) for c in col) + 2))
    ws.freeze_panes = 'A2'
    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()
