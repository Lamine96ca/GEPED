import csv
import datetime
import io
import re
from pathlib import Path

from django.contrib import messages
from django.contrib.auth import update_session_auth_hash
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.paginator import Paginator
from django.db import IntegrityError, transaction
from django.db.models import Count, Q, Sum
from django.http import Http404, HttpResponse, JsonResponse, StreamingHttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils.text import slugify
from django.views.decorators.http import require_POST

from . import liste
from django.conf import settings

from .attestations import CONTENUS as CONTENUS_ATTESTATION, infos_publiques, pdf_attestation, titre_document
from .forms import (AttestationFinEtudesForm, AttestationInscriptionForm, AttestationNiveauForm,
                    AttestationOrientationForm, AttestationSoutenanceForm, CoordonneesForm, EtudiantForm, ImportForm,
                    ImportMatieresForm,
                    InscriptionForm, MatiereForm, MotDePasseForm)
from .import_matieres import OBLIGATOIRES as OBLIGATOIRES_MATIERES, importer_matieres, modele_matieres_xlsx
from .importer import OBLIGATOIRES, importer_etudiants, modele_import_xlsx
from .models import (Attestation, Etudiant, Inscription, Licence, Promotion, Promotion_Licence_Session, Semestre,
                     Service, Session)
from .photos import compresser, enregistrer_photo, retirer_photo, url_photo
from .roles import (acces_requis, gestion_matieres_requise, get_profil, inscriptions_visibles, matieres_visibles,
                    programmes_du_profil, scolarite_requise)

TAILLES_PAGE = (25, 50, 100, 250, 500)


# ------------------------------------------------------------------ accueil

def _pct(n, total):
    return round(100 * n / total) if total else 0


def _repartition_par_sexe(qs, champ, defaut, trier_par_effectif=False, limite=None):
    """
    Effectifs groupés par `champ`, avec la répartition garçons / filles.
    Une seule requête GROUP BY (champ, sexe). `largeur` = barre relative au groupe le plus grand.
    """
    groupes = {}
    for cle, sexe, n in qs.values_list(champ, 'etudiant__sexe').annotate(n=Count('id')).order_by():
        g = groupes.setdefault(cle, {'nom': cle or defaut, 'M': 0, 'F': 0})
        g['F' if sexe == 'F' else 'M'] += n
    lignes = list(groupes.values())
    for g in lignes:
        g['n'] = g['M'] + g['F']
        g['pct_f'] = _pct(g['F'], g['n'])
        g['pct_m'] = 100 - g['pct_f'] if g['n'] else 0
    lignes.sort(key=(lambda g: -g['n']) if trier_par_effectif else (lambda g: str(g['nom'])))
    if limite:
        lignes = lignes[:limite]
    plus_grand = max((g['n'] for g in lignes), default=0)
    for g in lignes:
        g['largeur'] = max(_pct(g['n'], plus_grand), 2)
    return lignes


@login_required
@acces_requis
def tableau_de_bord(request):
    profil = get_profil(request)
    session = liste.session_par_defaut()
    qs = inscriptions_visibles(profil)
    if session:
        qs = qs.filter(session=session)

    par_sexe = dict(qs.values_list('etudiant__sexe').annotate(n=Count('id')).order_by())
    garcons, filles = par_sexe.get('M', 0), par_sexe.get('F', 0)
    total = garcons + filles
    redoublants = qs.filter(redouble=True).count()
    preinscrits = qs.filter(etudiant__preinscrit=True).count()

    par_programme = _repartition_par_sexe(qs, 'filiere__nom', 'Sans programme', trier_par_effectif=True, limite=10)
    par_niveau = _repartition_par_sexe(qs, 'licence__nom', 'Sans niveau')
    par_cycle = [
        {'code': code or '?', 'nom': liste.CYCLE_LIBELLES.get(code, 'Non renseigné'), 'n': n, 'pct': _pct(n, total)}
        for code, n in qs.values_list('cycle').annotate(n=Count('id')).order_by('-n')
    ]

    return render(request, 'pfe/tableau_de_bord.html', {
        'session': session,
        'total': total,
        'garcons': garcons,
        'filles': filles,
        'pct_garcons': _pct(garcons, total),
        'pct_filles': _pct(filles, total),
        'redoublants': redoublants,
        'pct_redoublants': _pct(redoublants, total),
        'preinscrits': preinscrits,
        'pct_preinscrits': _pct(preinscrits, total),
        'nb_programmes': qs.exclude(filiere=None).values('filiere').distinct().count(),
        'par_programme': par_programme,
        'par_niveau': par_niveau,
        'par_cycle': par_cycle,
        'actif': 'tableau',
    })


# ------------------------------------------------------------------ liste

def _options_filtres(profil):
    programmes = Service.objects.filter(programme=True).order_by('nom')
    if profil.role == 'programme':
        programmes = programmes.filter(id__in=profil.programme_ids)
    return {
        'sessions': Session.objects.order_by('-annee'),
        'programmes': programmes,
        'licences': Licence.objects.order_by('nom'),
        'semestres': Semestre.objects.order_by('nom'),
        'promotions': Promotion.objects.order_by('nom'),
        'cycles': Inscription.CYCLE_CHOICES,
        'cycles_libelles': liste.CYCLE_LIBELLES,
    }


def _page_liste(request, mode):
    profil = get_profil(request)
    return render(request, 'pfe/etudiants_liste.html', {
        **_options_filtres(profil),
        'tailles_page': TAILLES_PAGE,
        'mode': mode,
        'peut_inscrire': mode == 'preinscriptions' and profil.est_scolarite,
        'actif': mode,
    })


@login_required
@acces_requis
def etudiants_liste(request):
    return _page_liste(request, 'etudiants')


@login_required
@acces_requis
def preinscriptions_liste(request):
    """Même liste que « Étudiants », limitée aux préinscrits ; la scolarité peut les inscrire."""
    return _page_liste(request, 'preinscriptions')


@login_required
@scolarite_requise
@require_POST
def valider_preinscriptions(request):
    """
    Inscrit les préinscrits sélectionnés (statut préinscrit retiré).
    Reçoit une ou plusieurs inscriptions (`ids`). Réponse JSON pour la liste, redirection sinon.
    """
    ids = [i for i in request.POST.getlist('ids') if i.isdigit()][:1000]
    # Liste Python (pas de sous-requête) : MySQL refuse un UPDATE de pfe_etudiant qui relit pfe_etudiant
    etudiants = list(Inscription.objects.filter(pk__in=ids, etudiant__preinscrit=True)
                     .values_list('etudiant_id', flat=True))
    # Demande d'inscription déposée en ligne : ses coordonnées et sa photo sont appliquées en même temps
    from .espace_etudiant import appliquer
    from .models import DemandeEtudiant
    for demande in DemandeEtudiant.objects.filter(etudiant_id__in=etudiants, type=DemandeEtudiant.INSCRIPTION,
                                                  statut=DemandeEtudiant.EN_ATTENTE).select_related('etudiant'):
        appliquer(demande, request.user)
    Etudiant.objects.filter(pk__in=etudiants).update(preinscrit=False)
    n = len(set(etudiants))

    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return JsonResponse({'inscrits': n})
    if n:
        messages.success(request, "L'étudiant est maintenant inscrit." if n == 1 else f"{n} étudiants sont maintenant inscrits.")
    else:
        messages.info(request, "Aucun préinscrit à inscrire.")
    suivant = request.POST.get('suivant', '')
    if suivant.startswith('/') and not suivant.startswith('//'):
        return redirect(suivant)
    return redirect('preinscriptions_liste')


@login_required
@acces_requis
def etudiants_donnees(request):
    """JSON paginé pour le tableau. Lignes renvoyées sous forme de listes (compact)."""
    profil = get_profil(request)
    qs, _ = liste.filtrer(profil, request.GET)
    qs = liste.trier(qs, request.GET)

    try:
        taille = int(request.GET.get('taille', 50))
    except ValueError:
        taille = 50
    taille = taille if taille in TAILLES_PAGE else 50
    try:
        page = max(1, int(request.GET.get('page', 1)))
    except ValueError:
        page = 1

    total = qs.count()
    nb_pages = max(1, -(-total // taille))
    page = min(page, nb_pages)
    debut = (page - 1) * taille
    champs = [c for _, c in liste.COLONNES_TABLE]
    lignes = list(qs.values_list(*champs)[debut:debut + taille])

    return JsonResponse({
        'colonnes': [k for k, _ in liste.COLONNES_TABLE],
        'lignes': lignes,
        'total': total,
        'page': page,
        'pages': nb_pages,
        'taille': taille,
    })


def _valeur_export(v):
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'Oui' if v else 'Non'
    if isinstance(v, datetime.date):
        return v.strftime('%d/%m/%Y')
    return v


class _Echo:
    def write(self, valeur):
        return valeur


@login_required
@acces_requis
def etudiants_export(request):
    profil = get_profil(request)
    qs, session_id = liste.filtrer(profil, request.GET)
    qs = liste.trier(qs, request.GET)
    entetes = [e for e, _ in liste.COLONNES_EXPORT]
    champs = [c for _, c in liste.COLONNES_EXPORT]
    i_cycle = champs.index('cycle')

    def convertir(ligne):
        ligne = list(ligne)
        ligne[i_cycle] = liste.CYCLE_LIBELLES.get(ligne[i_cycle], ligne[i_cycle])   # code -> libellé
        return [_valeur_export(v) for v in ligne]

    lignes = (convertir(l) for l in qs.values_list(*champs).iterator(chunk_size=2000))

    session = Session.objects.filter(pk=session_id).first() if session_id else None
    suffixe = (session.annee.replace('/', '-').replace(' ', '') if session else 'toutes-sessions')
    nom_fichier = f"etudiants_{suffixe}_{datetime.date.today():%Y%m%d}"

    if request.GET.get('format') == 'csv':
        ecrivain = csv.writer(_Echo(), delimiter=';')

        def flux():
            yield '﻿'                       # BOM : accents corrects dans Excel
            yield ecrivain.writerow(entetes)
            for ligne in lignes:
                yield ecrivain.writerow(ligne)

        reponse = StreamingHttpResponse(flux(), content_type='text/csv; charset=utf-8')
        reponse['Content-Disposition'] = f'attachment; filename="{nom_fichier}.csv"'
        return reponse

    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet('Etudiants')
    ws.freeze_panes = 'A2'
    gras, fond = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='41627C')
    ligne_entete = []
    for e in entetes:
        cell = WriteOnlyCell(ws, value=e)
        cell.font, cell.fill = gras, fond
        ligne_entete.append(cell)
    ws.append(ligne_entete)
    for ligne in lignes:
        ws.append(ligne)
    sortie = io.BytesIO()
    wb.save(sortie)
    reponse = HttpResponse(sortie.getvalue(),
                           content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    reponse['Content-Disposition'] = f'attachment; filename="{nom_fichier}.xlsx"'
    return reponse


# ------------------------------------------------------------------ fiche

def _inscription_visible(request, pk):
    qs = inscriptions_visibles(get_profil(request)).select_related(
        'etudiant', 'session', 'licence', 'semestre', 'filiere', 'promotion')
    return get_object_or_404(qs, pk=pk)


@login_required
@acces_requis
def etudiant_fiche(request, pk):
    inscription = _inscription_visible(request, pk)
    historique = inscriptions_visibles(get_profil(request)).filter(
        etudiant=inscription.etudiant).select_related('session', 'licence', 'semestre', 'filiere', 'promotion'
                                                       ).order_by('-session__annee')
    from .models import PieceDossier
    profil = get_profil(request)
    return render(request, 'pfe/etudiant_fiche.html', {
        'inscription': inscription,
        'etudiant': inscription.etudiant,
        'historique': historique,
        'actif': 'etudiants',
        'photo_url': url_photo(inscription.etudiant),
        'peut_changer_photo': _peut_changer_photo(request, inscription.etudiant),
        # Dossier (pièces) : scolarité seulement (documents personnels)
        'pieces': (inscription.etudiant.pieces.select_related('session', 'utilisateur')
                   if profil.est_scolarite else None),
        'types_pieces': PieceDossier.TYPES,
    })


# ------------------------------------------------------------------ dossier (pièces)

def _retour_fiche(request, etudiant):
    suivant = request.POST.get('suivant', '')
    if suivant.startswith('/') and not suivant.startswith('//'):
        return redirect(suivant)
    insc = etudiant.etudiant_inscription.order_by('-session__annee').first()
    return redirect('etudiant_fiche', pk=insc.pk) if insc else redirect('etudiants_liste')


@login_required
@scolarite_requise
@require_POST
def piece_ajouter(request, pk):
    """Ajoute une pièce au dossier (scan ou fichier), compressée en PDF ; remplace celle du même type."""
    from . import pieces
    from .models import PieceDossier
    etudiant = get_object_or_404(Etudiant, pk=pk)
    type_piece = request.POST.get('type')
    if type_piece not in dict(PieceDossier.TYPES):
        messages.error(request, "Choisissez le type de pièce.")
        return _retour_fiche(request, etudiant)
    try:
        piece = pieces.enregistrer(etudiant, type_piece, request.FILES.getlist('pages'), utilisateur=request.user,
                                   libelle=(request.POST.get('libelle') or '').strip())
    except ValidationError as e:
        messages.error(request, ' '.join(e.messages))
    else:
        messages.success(request, f"{piece.get_type_display()} ajouté au dossier ({piece.pages} page"
                                  f"{'s' if piece.pages > 1 else ''}, {max(1, piece.taille // 1024)} Ko).")
    return _retour_fiche(request, etudiant)


@login_required
@scolarite_requise
def piece_voir(request, pk):
    from django.http import FileResponse
    from . import pieces
    from .models import PieceDossier
    piece = get_object_or_404(PieceDossier.objects.select_related('etudiant'), pk=pk)
    nom = f"{piece.type}_{piece.etudiant.matricule or piece.etudiant.pk}.pdf"
    reponse = FileResponse(pieces.ouvrir(piece), content_type='application/pdf', filename=nom)
    reponse['Content-Disposition'] = f'inline; filename="{nom}"'
    reponse['X-Content-Type-Options'] = 'nosniff'
    reponse['Cache-Control'] = 'private, no-store'
    return reponse


@login_required
@scolarite_requise
@require_POST
def piece_supprimer(request, pk):
    from . import pieces
    from .models import PieceDossier
    piece = get_object_or_404(PieceDossier.objects.select_related('etudiant'), pk=pk)
    etudiant, libelle = piece.etudiant, piece.get_type_display()
    pieces.supprimer(piece)
    messages.success(request, f"{libelle} retiré du dossier.")
    return _retour_fiche(request, etudiant)


# ------------------------------------------------------------------ photos

@login_required
def media_protege(request, chemin):
    """
    Fichiers déposés (photos des étudiants et du personnel) : jamais publics. Visibles par le personnel
    ayant un rôle gePed ; un compte sans rôle (étudiant) ne voit que SA propre photo.
    """
    import mimetypes
    from django.http import FileResponse
    from django.utils._os import safe_join
    try:
        fichier = Path(safe_join(settings.MEDIA_ROOT, chemin))
    except Exception:                                   # « ../ » ou chemin absolu : refusé
        raise Http404()
    if not fichier.is_file():
        raise Http404()
    if not get_profil(request).a_acces:
        fiche = _fiche_du_compte(request.user)
        if fiche is None or not fiche.photo or fiche.photo.name != chemin:
            raise PermissionDenied("Ce fichier n'est pas accessible avec votre compte.")
    reponse = FileResponse(open(fichier, 'rb'), content_type=mimetypes.guess_type(fichier.name)[0]
                           or 'application/octet-stream')
    reponse['X-Content-Type-Options'] = 'nosniff'
    reponse['Cache-Control'] = 'private, max-age=86400'
    return reponse


def _peut_changer_photo(request, etudiant):
    """La scolarité, ou l'étudiant lui-même (compte lié à sa fiche)."""
    return get_profil(request).est_scolarite or (
        etudiant.utilisateur_id is not None and etudiant.utilisateur_id == request.user.id)


def _changer_photo(request, obj, base_nom):
    """POST « photo » (fichier) ou « retirer » : enregistre et répond en JSON {url} ou {erreur}."""
    if request.POST.get('retirer'):
        retirer_photo(obj)
        return JsonResponse({'url': None})
    fichier = request.FILES.get('photo')
    if not fichier:
        return JsonResponse({'erreur': "Aucune photo reçue."}, status=400)
    try:
        enregistrer_photo(obj, compresser(fichier), base_nom)
    except ValidationError as e:
        return JsonResponse({'erreur': ' '.join(e.messages)}, status=400)
    return JsonResponse({'url': obj.photo.url, 'poids': obj.photo.size})


@login_required
@require_POST
def etudiant_photo(request, pk):
    etudiant = get_object_or_404(Etudiant, pk=pk)
    if not _peut_changer_photo(request, etudiant):
        raise PermissionDenied("Seule la scolarité (ou l'étudiant lui-même) peut changer cette photo.")
    return _changer_photo(request, etudiant, etudiant.matricule or f"etudiant-{etudiant.pk}")


def _fiche_du_compte(user):
    """Fiche dont la photo est celle de l'utilisateur connecté : personnel, sinon étudiant."""
    return getattr(user, 'profil_personnel', None) or getattr(user, 'profil_etudiant', None)


@login_required
@require_POST
def mon_compte_photo(request):
    fiche = _fiche_du_compte(request.user)
    if fiche is None:
        return JsonResponse({'erreur': "Aucune fiche n'est liée à votre compte : photo impossible."}, status=400)
    return _changer_photo(request, fiche, request.user.username)


# ------------------------------------------------------------------ scolarité

def _enregistrer(request, form_etu, form_insc, etudiant_existant=None):
    if not (form_etu.is_valid() and form_insc.is_valid()):
        return None
    if not form_insc.valider_unicite(form_etu.instance):
        return None
    try:
        with transaction.atomic():
            etudiant = form_etu.save()
            inscription = form_insc.save(commit=False)
            inscription.etudiant = etudiant
            inscription.save()
    except IntegrityError:
        form_insc.add_error(None, "Enregistrement impossible : cette inscription existe déjà.")
        return None
    return inscription


@login_required
@scolarite_requise
def inscription_creer(request):
    # ?etudiant=<id> : réinscription d'un étudiant existant (nouvelle session / licence)
    etudiant = None
    if request.GET.get('etudiant'):
        etudiant = get_object_or_404(Etudiant, pk=request.GET['etudiant'])
    if request.method == 'POST':
        form_etu = EtudiantForm(request.POST, request.FILES, instance=etudiant, prefix='etu')
        form_insc = InscriptionForm(request.POST, prefix='insc')
        inscription = _enregistrer(request, form_etu, form_insc)
        if inscription:
            messages.success(request, f"{inscription.etudiant.prenom} {inscription.etudiant.nom} a été inscrit(e).")
            if 'continuer' in request.POST:
                return redirect('inscription_creer')
            # La fiche propose aussitôt d'imprimer la fiche d'inscription
            return redirect(f"{reverse('etudiant_fiche', args=[inscription.pk])}?imprimer=1")
    else:
        form_etu = EtudiantForm(instance=etudiant, prefix='etu')
        form_insc = InscriptionForm(prefix='insc', initial=_prevision_reinscription(etudiant))
    contexte = {
        'form_etu': form_etu, 'form_insc': form_insc, 'etudiant': etudiant,
        'titre': 'Réinscrire un étudiant' if etudiant else 'Inscrire un étudiant',
        'actif': 'inscrire', 'promotions_prevues': _promotions_prevues(),
    }
    if etudiant:
        # Réinscription : toutes ses notes, année par année, pour décider ; inscription déjà faite ?
        from .views_demandes import SEUIL_ROUGE, notes_empilees
        session = liste.session_par_defaut()
        contexte.update(annees=notes_empilees(etudiant), seuil_rouge=SEUIL_ROUGE,
                        derniere=_derniere_inscription(etudiant),
                        deja_inscrit=etudiant.etudiant_inscription.filter(session=session).select_related('licence').first()
                        if session else None)
    return render(request, 'pfe/inscription_form.html', contexte)


def _derniere_inscription(etudiant):
    return (etudiant.etudiant_inscription.select_related('session', 'licence', 'filiere', 'promotion')
            .order_by('-session__annee', '-licence__nom').first())


def _prevision_reinscription(etudiant):
    """
    Réinscription : inscription de la session en cours préremplie d'après la dernière inscription
    (programme, promotion, cycle) et le niveau prévu pour la promotion dans la session
    (Promotion_Licence_Session). Tout reste modifiable.
    """
    if etudiant is None:
        return {}
    from .espace_etudiant import niveaux_prevus
    session = liste.session_par_defaut()
    d = _derniere_inscription(etudiant)
    if d is None or session is None or d.session_id == session.pk:
        return {}
    niveaux = niveaux_prevus(d, session)
    licence = niveaux[0] if niveaux else None
    return {k: v for k, v in {
        'session': session.pk, 'cycle': d.cycle, 'filiere': d.filiere_id, 'promotion': d.promotion_id,
        'licence': licence.pk if licence else None, 'redouble': bool(licence and licence.pk == d.licence_id),
    }.items() if v not in (None, '')}


def _promotions_prevues():
    """{session: {licence: promotion}} d'après Promotion_Licence_Session (promotion placée selon le niveau)."""
    prevues = {}
    for session_id, licence_id, promotion_id in Promotion_Licence_Session.objects.values_list(
            'session_id', 'licence_id', 'promotion_id'):
        prevues.setdefault(str(session_id), {}).setdefault(str(licence_id), promotion_id)
    return prevues


@login_required
@scolarite_requise
def inscription_fiche_pdf(request, pk):
    """Fiche d'inscription (ou de réinscription) à imprimer, avec la photo de l'étudiant."""
    from .fiche_inscription import reponse_fiche
    return reponse_fiche(get_object_or_404(Inscription.objects.select_related(
        'etudiant', 'session', 'licence', 'semestre', 'filiere', 'promotion'), pk=pk))


@login_required
@scolarite_requise
def inscription_modifier(request, pk):
    inscription = get_object_or_404(Inscription.objects.select_related('etudiant'), pk=pk)
    if request.method == 'POST':
        form_etu = EtudiantForm(request.POST, request.FILES, instance=inscription.etudiant, prefix='etu')
        form_insc = InscriptionForm(request.POST, instance=inscription, prefix='insc')
        if _enregistrer(request, form_etu, form_insc):
            messages.success(request, "Modifications enregistrées.")
            return redirect('etudiant_fiche', pk=inscription.pk)
    else:
        form_etu = EtudiantForm(instance=inscription.etudiant, prefix='etu')
        form_insc = InscriptionForm(instance=inscription, prefix='insc')
    return render(request, 'pfe/inscription_form.html', {
        'form_etu': form_etu, 'form_insc': form_insc, 'inscription': inscription,
        'titre': 'Modifier l’inscription', 'actif': 'etudiants',
    })


@login_required
@scolarite_requise
@require_POST
def inscription_supprimer(request, pk):
    inscription = get_object_or_404(Inscription.objects.select_related('etudiant', 'session'), pk=pk)
    etudiant = inscription.etudiant
    inscription.delete()
    messages.success(request, f"Inscription de {etudiant.prenom} {etudiant.nom} ({inscription.session}) supprimée.")
    return redirect('etudiants_liste')


@login_required
@scolarite_requise
def importer(request):
    rapport = None
    if request.method == 'POST':
        form = ImportForm(request.POST, request.FILES)
        if form.is_valid():
            d = form.cleaned_data
            try:
                rapport = importer_etudiants(d['fichier'], simulation=d['simulation'], cycle=d['cycle'] or None,
                                             preinscription=d['preinscription'])
            except ValueError as e:
                form.add_error('fichier', str(e))
            except IntegrityError as e:
                form.add_error(None, f"Import annulé, conflit en base : {e}")
            else:
                if not rapport.simulation and rapport.inscriptions_creees:
                    messages.success(request, "Import terminé et enregistré.")
    else:
        form = ImportForm()
    return render(request, 'pfe/importer.html', {
        'form': form, 'rapport': rapport, 'actif': 'importer', 'obligatoires': OBLIGATOIRES,
    })


@login_required
@scolarite_requise
def modele_import(request):
    reponse = HttpResponse(modele_import_xlsx(),
                           content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    reponse['Content-Disposition'] = 'attachment; filename="modele_import_etudiants.xlsx"'
    return reponse


# ------------------------------------------------------------------ statistiques

@login_required
@acces_requis
def statistiques(request):
    """Statistiques des inscriptions (H/F par niveau, programme… voir pfe/statistiques.py)."""
    from . import statistiques as st
    from .models import Resultat, Soutenance
    profil = get_profil(request)
    g = request.GET
    sessions = Session.objects.order_by('-annee')
    if 'session' in g:
        session_id = _entier(g.get('session'))
    else:
        courante = sessions.filter(en_cours=True).first() or sessions.first()
        session_id = courante.pk if courante else None
    programmes = list(programmes_du_profil(profil))
    programme_id = _entier(g.get('programme'))
    if programme_id not in {p.pk for p in programmes}:
        programme_id = None
    cycles = dict(Inscription.CYCLE_CHOICES)
    cycle = g.get('cycle') if g.get('cycle') in cycles else ''

    base = inscriptions_visibles(profil).filter(etudiant__preinscrit=False)
    resultats = Resultat.objects.filter(niveau=Resultat.ANNEE, filiere__isnull=False)
    sout = Soutenance.objects.all()
    if profil.role == 'programme':
        resultats = resultats.filter(filiere_id__in=profil.programme_ids)
        sout = sout.filter(filiere_id__in=profil.programme_ids)
    if cycle:
        base, resultats, sout = base.filter(cycle=cycle), resultats.filter(cycle=cycle), sout.filter(cycle=cycle)
    if programme_id:
        base = base.filter(filiere_id=programme_id)
        resultats, sout = resultats.filter(filiere_id=programme_id), sout.filter(filiere_id=programme_id)
    if session_id:
        resultats, sout = resultats.filter(session_id=session_id), sout.filter(session_id=session_id)
    s = st.calculer(base, session_id, resultats, sout)
    session_choisie = next((x for x in sessions if x.pk == session_id), None)
    programme_choisi = next((p for p in programmes if p.pk == programme_id), None)
    filtres_texte = ' · '.join(x for x in (f"Session {session_choisie}" if session_choisie else 'Toutes les sessions',
                                           cycles.get(cycle, ''), str(programme_choisi or '')) if x)
    if g.get('export') == 'xlsx':
        reponse = HttpResponse(st.classeur(s, filtres_texte),
                               content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        reponse['Content-Disposition'] = 'attachment; filename="statistiques_inscriptions.xlsx"'
        return reponse
    params = g.copy()
    params.pop('export', None)
    return render(request, 'pfe/statistiques.html', {
        **s, 'sessions': sessions, 'liste_programmes': programmes, 'cycles': Inscription.CYCLE_CHOICES,
        'filtres': {'session': session_id, 'programme': programme_id, 'cycle': cycle},
        'filtres_texte': filtres_texte, 'params': params.urlencode(), 'actif': 'statistiques',
        'blocs_reussite': [('Réussite par programme', s['reussite']['programmes'], 'col-xl-7'),
                           ('Réussite par niveau', s['reussite']['niveaux'], 'col-xl-5')],
    })


# ------------------------------------------------------------------ matières

def _entier(valeur):
    return int(valeur) if str(valeur or '').isdigit() else None


# Colonnes triables de la liste des matières : clé → (titre, champs de tri, classe CSS de la colonne)
TRIS_MATIERES = {
    'nom': ('Matière', ['nom'], ''),
    'code': ('Code', ['code'], ''),
    'credit': ('Crédits', ['credit'], 'text-end'),
    'cycle': ('Cycle', ['cycle'], ''),
    'semestre': ('Semestre', ['semestre_id'], ''),
    'licence': ('Licence', ['licence_id'], ''),
    'programme': ('Programme', ['filiere__nom'], ''),
    'session': ('Session', ['session__annee'], ''),
    'promotion': ('Promo', ['promotion__nom'], 'd-none d-xl-table-cell'),
}


@login_required
@acces_requis
def matieres_liste(request):
    """Scolarité et autres services : toutes les matières ; chef de programme : les siennes."""
    profil = get_profil(request)
    programmes = programmes_du_profil(profil)
    qs = matieres_visibles(profil).select_related('filiere', 'licence', 'semestre', 'session', 'promotion')

    sessions = Session.objects.order_by('-annee')
    g = request.GET
    if 'session' in g:
        session_id = _entier(g.get('session'))                  # vide = toutes les sessions
    else:
        courante = sessions.filter(en_cours=True).first() or sessions.first()
        session_id = courante.pk if courante else None
    filtres = {'session': session_id, 'programme': _entier(g.get('programme')),
               'licence': _entier(g.get('licence')), 'semestre': _entier(g.get('semestre')),
               'cycle': g.get('cycle') if g.get('cycle') in dict(Inscription.CYCLE_CHOICES) else None}
    if filtres['cycle']:
        qs = qs.filter(cycle=filtres['cycle'])
    if filtres['session']:
        qs = qs.filter(session_id=filtres['session'])
    if filtres['programme']:
        qs = qs.filter(filiere_id=filtres['programme'])
    if filtres['licence']:
        qs = qs.filter(licence_id=filtres['licence'])
    if filtres['semestre']:
        qs = qs.filter(semestre_id=filtres['semestre'])
    q = (g.get('q') or '').strip()
    if q:
        qs = qs.filter(Q(nom__icontains=q) | Q(code__icontains=q))

    # Crédits par programme et par semestre (sur la sélection)
    totaux = list(qs.values('filiere__nom', 'licence__nom', 'semestre__nom', 'semestre_id')
                  .annotate(nb=Count('id'), credits=Sum('credit'))
                  .order_by('filiere__nom', 'semestre_id'))
    total = qs.aggregate(nb=Count('id'), credits=Sum('credit'))

    # Tri (clic sur un titre de colonne) ; par défaut : programme, licence, semestre, nom
    tri = g.get('tri') if g.get('tri') in TRIS_MATIERES else None
    ordre = 'desc' if g.get('ordre') == 'desc' else 'asc'
    if tri:
        champs = [('-' if ordre == 'desc' else '') + c for c in TRIS_MATIERES[tri][1]]
        qs = qs.order_by(*champs, 'nom', 'pk')
    else:
        qs = qs.order_by('filiere__nom', 'licence_id', 'semestre_id', 'nom', 'pk')
    page = Paginator(qs, 50).get_page(g.get('page'))
    params = g.copy()
    params.pop('page', None)
    sans_tri = params.copy()
    for cle_ in ('tri', 'ordre'):
        sans_tri.pop(cle_, None)
    colonnes = []
    for cle_, (libelle, _, classe) in TRIS_MATIERES.items():
        actif = cle_ == tri
        lien = sans_tri.copy()
        lien['tri'] = cle_
        lien['ordre'] = 'desc' if actif and ordre == 'asc' else 'asc'
        colonnes.append({'cle': cle_, 'libelle': libelle, 'classe': classe, 'actif': actif, 'ordre': ordre,
                         'url': lien.urlencode()})
    return render(request, 'pfe/matieres_liste.html', {
        'avec_descriptions': any(m.description for m in page),
        'page': page, 'totaux': totaux, 'total': total, 'filtres': filtres, 'q': q, 'tri': tri, 'ordre': ordre,
        'colonnes': colonnes,
        'params': params.urlencode(), 'sessions': sessions, 'programmes': programmes,
        'licences': Licence.objects.order_by('id'),
        'semestres': Semestre.objects.select_related('licence').order_by('licence_id', 'id'),
        'cycles': Inscription.CYCLE_CHOICES,
        'peut_gerer': profil.gere_matieres, 'actif': 'matieres',
    })


def _formulaire_matiere(request, matiere=None):
    profil = get_profil(request)
    programmes = list(programmes_du_profil(profil))
    if request.method == 'POST':
        form = MatiereForm(request.POST, instance=matiere, programmes=programmes_du_profil(profil))
        if form.is_valid():
            m = form.save()
            messages.success(request, f"Matière « {m.nom} » {'modifiée' if matiere else 'ajoutée'}.")
            if 'continuer' in request.POST:
                return redirect('matiere_creer')
            return redirect(f"{reverse('matieres_liste')}?session={m.session_id}&programme={m.filiere_id}")
    else:
        form = MatiereForm(instance=matiere, programmes=programmes_du_profil(profil))
    semestres_par_licence = {s.pk: s.licence_id for s in Semestre.objects.all()}
    return render(request, 'pfe/matiere_form.html', {
        'form': form, 'matiere': matiere, 'actif': 'matieres', 'un_seul_programme': len(programmes) == 1,
        'semestres_par_licence': semestres_par_licence,
        'cycles_des_programmes': {p.pk: p.get_cycle_display() if p.cycle else '' for p in programmes},
        'titre': f"Modifier « {matiere.nom} »" if matiere else 'Ajouter une matière',
    })


@login_required
@gestion_matieres_requise
def matiere_creer(request):
    return _formulaire_matiere(request)


@login_required
@gestion_matieres_requise
def matiere_modifier(request, pk):
    # Hors de son périmètre (autre programme) : 404, la matière « n'existe pas » pour lui
    return _formulaire_matiere(request, get_object_or_404(matieres_visibles(get_profil(request)), pk=pk))


def _utilisations(matiere):
    """Nombre d'éléments (notes, etc.) qui dépendent de la matière : la supprimer les effacerait."""
    total = 0
    for rel in matiere._meta.related_objects:
        if rel.one_to_many or rel.one_to_one:
            total += rel.related_model.objects.filter(**{rel.field.name: matiere}).count()
    return total


@login_required
@gestion_matieres_requise
@require_POST
def matiere_supprimer(request, pk):
    matiere = get_object_or_404(matieres_visibles(get_profil(request)), pk=pk)
    retour = request.POST.get('suivant') or reverse('matieres_liste')
    if not retour.startswith('/') or retour.startswith('//'):
        retour = reverse('matieres_liste')
    utilisations = _utilisations(matiere)
    if utilisations:
        messages.error(request, f"Suppression impossible : « {matiere.nom} » a déjà {utilisations} note(s) "
                                f"ou élément(s) enregistré(s).")
        return redirect(retour)
    matiere.delete()
    messages.success(request, f"Matière « {matiere.nom} » supprimée.")
    return redirect(retour)


@login_required
@gestion_matieres_requise
def matieres_importer(request):
    profil = get_profil(request)
    rapport = None
    if request.method == 'POST':
        form = ImportMatieresForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                rapport = importer_matieres(form.cleaned_data['fichier'], profil,
                                            simulation=form.cleaned_data['simulation'])
            except ValueError as e:                      # colonnes manquantes, programme hors périmètre…
                form.add_error('fichier', str(e))
            except IntegrityError as e:
                form.add_error(None, f"Import annulé, conflit en base : {e}")
            else:
                if not rapport.simulation and (rapport.creees or rapport.mises_a_jour):
                    messages.success(request, "Import des matières terminé et enregistré.")
    else:
        form = ImportMatieresForm()
    return render(request, 'pfe/matieres_importer.html', {
        'form': form, 'rapport': rapport, 'actif': 'matieres', 'obligatoires': OBLIGATOIRES_MATIERES,
        'mes_programmes': None if profil.est_scolarite else programmes_du_profil(profil),
    })


@login_required
@gestion_matieres_requise
def matieres_modele(request):
    reponse = HttpResponse(modele_matieres_xlsx(),
                           content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    reponse['Content-Disposition'] = 'attachment; filename="modele_import_matieres.xlsx"'
    return reponse


# ------------------------------------------------------------------ mon compte

@login_required
def mon_compte(request):
    """Chaque utilisateur modifie uniquement son e-mail, son téléphone et son mot de passe."""
    form_coord = CoordonneesForm(user=request.user)
    form_mdp = MotDePasseForm(request.user)

    if request.method == 'POST' and 'coordonnees' in request.POST:
        form_coord = CoordonneesForm(request.POST, user=request.user)
        if form_coord.is_valid():
            form_coord.save()
            messages.success(request, "Vos coordonnées ont été mises à jour.")
            return redirect('mon_compte')
    elif request.method == 'POST' and 'mot_de_passe' in request.POST:
        form_mdp = MotDePasseForm(request.user, request.POST)
        if form_mdp.is_valid():
            user = form_mdp.save()
            update_session_auth_hash(request, user)      # reste connecté après le changement
            messages.success(request, "Votre mot de passe a été modifié.")
            return redirect('mon_compte')

    return render(request, 'pfe/mon_compte.html', {
        'form_coord': form_coord,
        'form_mdp': form_mdp,
        'personnel': form_coord.personnel,
        'actif': 'compte',
        'fiche_photo': _fiche_du_compte(request.user) is not None,
    })


# ------------------------------------------------------------------ attestations

TYPES_ATTESTATION = [
    # (clé d'URL, code du modèle, libellé, disponible)
    ('orientation', Attestation.ORIENTATION, 'Orientation', True),
    ('inscription', Attestation.INSCRIPTION, 'Inscription', True),
    ('niveau', Attestation.NIVEAU, 'Niveau', True),
    ('fin-etudes', Attestation.FIN_ETUDES, "Fin d'études", True),
    ('soutenance', Attestation.SOUTENANCE, 'Soutenance', True),
]
FORMULAIRES_ATTESTATION = {
    Attestation.ORIENTATION: AttestationOrientationForm,
    Attestation.INSCRIPTION: AttestationInscriptionForm,
    Attestation.NIVEAU: AttestationNiveauForm,
    Attestation.FIN_ETUDES: AttestationFinEtudesForm,
    Attestation.SOUTENANCE: AttestationSoutenanceForm,
}
COMPLEMENT_ATTESTATION = {'orientation': "d'orientation", 'inscription': "d'inscription",
                          'niveau': 'de niveau', 'fin-etudes': "de fin d'études", 'soutenance': 'de soutenance'}


@login_required
@scolarite_requise
def attestations(request, type_url='inscription'):
    """
    Attestations d'un type : formulaire (ajout ou modification, bouton « Enregistrer »),
    aide à la recherche d'un étudiant, et liste des attestations enregistrées.
    """
    types = {t[0]: t for t in TYPES_ATTESTATION}
    if type_url not in types:
        raise Http404
    _, code, libelle, disponible = types[type_url]

    # Formulaire : ajout, ou modification avec ?modifier=<id>
    instance = None
    if disponible and request.GET.get('modifier'):
        instance = get_object_or_404(Attestation, pk=request.GET['modifier'], type=code)
    form = None
    if disponible:
        if request.method == 'POST':
            form = FORMULAIRES_ATTESTATION[code](request.POST, instance=instance)
            if form.is_valid():
                attestation = form.save(session_actuelle=liste.session_par_defaut())
                if getattr(form, 'manquants', None):
                    messages.warning(request, "Informations absentes de la fiche de l'étudiant, imprimées « — » : "
                                              f"{', '.join(form.manquants)}. Complétez la fiche puis réimprimez.")
                messages.success(request, f"Attestation n° {attestation.num_ref} "
                                          f"{'modifiée' if instance else 'enregistrée'} pour "
                                          f"{attestation.etudiant.prenom} {attestation.etudiant.nom}.")
                return redirect(f"{reverse('attestations_type', args=[type_url])}?imprimer={attestation.pk}")
        else:
            form = FORMULAIRES_ATTESTATION[code](instance=instance, initial={
                'matricule': request.GET.get('matricule', ''), 'licence': request.GET.get('licence') or None,
                'cycle': request.GET.get('cycle', '')})

    # Aide : recherche d'un étudiant pour remplir le formulaire
    q_etudiant = (request.GET.get('etudiant') or '').strip()
    candidats = []
    if disponible and q_etudiant and code == Attestation.SOUTENANCE:
        from .models import Soutenance
        qs = Soutenance.objects.select_related('etudiant', 'session', 'filiere')
        for terme in q_etudiant.split()[:4]:
            qs = qs.filter(Q(etudiant__matricule__icontains=terme) | Q(etudiant__nom__icontains=terme)
                           | Q(etudiant__prenom__icontains=terme) | Q(etudiant__ine__icontains=terme))
        candidats = list(qs.order_by('etudiant__nom', 'etudiant__prenom', '-session__annee')[:30])
        existantes = {(a.etudiant_id, a.cycle): a for a in Attestation.objects.filter(
            type=code, etudiant__in={c.etudiant_id for c in candidats})}
        for c in candidats:
            c.attestation = existantes.get((c.etudiant_id, c.cycle))
    elif disponible and q_etudiant:
        qs = Inscription.objects.select_related('etudiant', 'licence', 'session', 'filiere')
        for terme in q_etudiant.split()[:4]:
            qs = qs.filter(Q(etudiant__matricule__icontains=terme) | Q(etudiant__nom__icontains=terme)
                           | Q(etudiant__prenom__icontains=terme) | Q(etudiant__ine__icontains=terme))
        candidats = list(qs.order_by('etudiant__nom', 'etudiant__prenom', '-session__annee')[:30])
        deja = Attestation.objects.filter(type=code, etudiant__in={c.etudiant_id for c in candidats})
        if code == Attestation.FIN_ETUDES:
            existantes = {(a.etudiant_id, a.cycle): a for a in deja}
            for c in candidats:
                c.attestation = existantes.get((c.etudiant_id, c.cycle))
        elif code == Attestation.ORIENTATION:
            # une seule orientation par étudiant : on affiche une ligne par étudiant
            existantes = {a.etudiant_id: a for a in deja}
            uniques = {}
            for c in candidats:
                uniques.setdefault(c.etudiant_id, c)
            candidats = list(uniques.values())
            for c in candidats:
                c.attestation = existantes.get(c.etudiant_id)
        else:
            existantes = {(a.etudiant_id, a.session_id, a.licence_id): a for a in deja}
            for c in candidats:
                c.attestation = existantes.get((c.etudiant_id, c.session_id, c.licence_id))

    # Attestations enregistrées
    q = (request.GET.get('q') or '').strip()
    enregistrees = Attestation.objects.filter(type=code).select_related('etudiant', 'licence', 'session', 'filiere')
    if q:
        enregistrees = enregistrees.filter(Q(num_ref__icontains=q) | Q(etudiant__matricule__icontains=q)
                                           | Q(etudiant__nom__icontains=q) | Q(etudiant__prenom__icontains=q))
    page = Paginator(enregistrees.order_by('-date_delivrance', '-id'), 25).get_page(request.GET.get('page'))

    imprimer = None
    if request.GET.get('imprimer', '').isdigit():
        imprimer = Attestation.objects.filter(pk=request.GET['imprimer'], type=code).first()

    return render(request, 'pfe/attestations.html', {
        'types': TYPES_ATTESTATION, 'type_url': type_url, 'libelle': libelle, 'disponible': disponible,
        'complement': COMPLEMENT_ATTESTATION[type_url],
        'form': form, 'instance': instance, 'imprimer': imprimer,
        'q_etudiant': q_etudiant, 'candidats': candidats, 'q': q, 'page': page,
        'prochaine_ref': Attestation.prochaine_reference(code),
        'actif': 'attestations',
    })


@login_required
@scolarite_requise
@require_POST
def attestation_supprimer(request, pk):
    """Supprime une attestation. Son numéro n'est jamais réattribué (numérotation continue)."""
    attestation = get_object_or_404(Attestation, pk=pk)
    type_url = next((t[0] for t in TYPES_ATTESTATION if t[1] == attestation.type), 'inscription')
    num = attestation.num_ref
    attestation.delete()
    messages.success(request, f"Attestation n° {num} supprimée. Ce numéro ne sera pas réutilisé.")
    return redirect('attestations_type', type_url=type_url)


@login_required
@scolarite_requise
def attestation_pdf(request, pk):
    attestation = get_object_or_404(
        Attestation.objects.select_related('etudiant', 'licence', 'session', 'filiere'), pk=pk)
    if attestation.type not in CONTENUS_ATTESTATION:
        raise Http404("Ce type d'attestation n'est pas encore disponible.")
    # Fin d'études : cachet et signature numérisés UNIQUEMENT pour un super utilisateur qui le demande.
    # Pour tout autre compte, « ?signature=1 » est ignoré : la version imprimée est sans cachet ni signature.
    signee = signature_autorisee(request, attestation) and request.GET.get('signature') == '1'
    reponse = HttpResponse(pdf_attestation(attestation, url_verification(request, attestation), cachet_signature=signee),
                           content_type='application/pdf')
    nom = f"attestation_{attestation.get_type_display().lower()}_{attestation.etudiant.matricule or attestation.pk}"
    if signee:
        nom += '_signee'
    reponse['Content-Disposition'] = f'inline; filename="{slugify(nom)}.pdf"'
    # Jamais en cache (navigateur, proxy) : une version signée ne doit pas être reprise par un autre compte
    reponse['Cache-Control'] = 'no-store, private'
    return reponse


def signature_autorisee(request, attestation):
    """Option « avec cachet et signature » : attestation de fin d'études et super utilisateur seulement."""
    return attestation.type == Attestation.FIN_ETUDES and request.user.is_superuser


# ------------------------------------------------------------------ vérification publique (QR code)

def url_verification(request, attestation):
    """Adresse encodée dans le QR code : GPED_URL_PUBLIQUE en production, sinon l'adresse de la requête."""
    chemin = reverse('verifier_document', args=[attestation.code_verification])
    base = getattr(settings, 'GPED_URL_PUBLIQUE', '')
    return f"{base}{chemin}" if base else request.build_absolute_uri(chemin)


def verifier_document(request, code):
    """
    Page PUBLIQUE ouverte en scannant le QR code d'un document : elle confirme que le document
    a bien été délivré par l'IST de Mamou et affiche les informations qu'il doit contenir.
    Le code est aléatoire (impossible à deviner) ; un document supprimé n'est plus reconnu.
    """
    # Un vrai code fait 22 caractères [A-Za-z0-9_-] : tout le reste est refusé sans interroger la base
    attestation = None
    if re.fullmatch(r'[A-Za-z0-9_-]{16,32}', code or ''):
        attestation = (Attestation.objects.select_related('etudiant', 'licence', 'session', 'filiere')
                       .filter(code_verification=code).first())
    contexte = {'attestation': attestation, 'contact': settings.GPED_ATTESTATION}
    releve = None
    if attestation is None and re.fullmatch(r'[A-Za-z0-9_-]{16,32}', code or ''):
        from .models import ReleveNote
        releve = (ReleveNote.objects.select_related('etudiant', 'filiere', 'session', 'licence', 'semestre')
                  .filter(code_verification=code).first())
    if releve is not None:
        from .releves import infos_publiques_releve
        titre, infos = infos_publiques_releve(releve)
        contexte.update({'attestation': releve, 'titre_document': titre, 'infos': infos})
    elif attestation is not None and attestation.etudiant is not None:
        contexte.update({'titre_document': titre_document(attestation), 'infos': infos_publiques(attestation)})
    else:
        contexte['attestation'] = None
    reponse = render(request, 'pfe/verification.html', contexte, status=200 if contexte['attestation'] else 404)
    reponse['X-Robots-Tag'] = 'noindex, nofollow'
    return reponse
