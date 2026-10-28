import io
import os
import re
from pathlib import Path as pathlib_path
from urllib.parse import quote

import openpyxl
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from .importer import importer_etudiants
from .models import Etudiant, Inscription, Licence, Personnel, Promotion, Promotion_Licence_Session, Service, Session


def fichier_xlsx(lignes, nom='import.xlsx'):
    wb = openpyxl.Workbook()
    for l in lignes:
        wb.active.append(l)
    flux = io.BytesIO()
    wb.save(flux)
    return SimpleUploadedFile(nom, flux.getvalue())


def photo_jpeg(nom='moi.jpg'):
    """Photo d'identité JPEG valide (espace étudiant : la photo est obligatoire)."""
    from PIL import Image
    flux = io.BytesIO()
    Image.new('RGB', (600, 800), (180, 150, 120)).save(flux, 'JPEG')
    return SimpleUploadedFile(nom, flux.getvalue(), 'image/jpeg')


def signer(vide=False):
    """Certification + signature dessinée (data URL PNG, comme le cadre de signature de l'espace)."""
    import base64
    from PIL import Image, ImageDraw
    img = Image.new('RGBA', (600, 180), (255, 255, 255, 0))
    if not vide:
        d = ImageDraw.Draw(img)
        d.line([(40, 120), (120, 40), (200, 130), (300, 50), (420, 120), (540, 60)], fill=(17, 24, 39, 255), width=5)
    flux = io.BytesIO()
    img.save(flux, 'PNG')
    return {'certifie': 'on', 'signature': 'data:image/png;base64,' + base64.b64encode(flux.getvalue()).decode()}


class GpedTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.session = Session.objects.create(annee='2025-2026', en_cours=True)
        cls.licence = Licence.objects.create(nom='Licence 1')
        cls.scol = Service.objects.create(nom='Scolarité', code='SCOL')
        cls.dep = Service.objects.create(nom='Département', code='DEP')
        cls.info = Service.objects.create(nom='Informatique', code='INFO', programme=True)
        cls.gc = Service.objects.create(nom='Génie civil', code='GC', programme=True)
        for i, (login, service) in enumerate([('scol', cls.scol), ('dep', cls.dep), ('info', cls.info)]):
            u = User.objects.create_user(login, password='x')
            p = Personnel.objects.create(utilisateur=u, prenom=login, nom='T', telephone=str(i),
                                         email=f'{login}@t.local', fonc='Agent')
            p.service.add(service)
        User.objects.create_user('sans_service', password='x')
        for k, prog in enumerate([cls.info, cls.info, cls.gc]):
            e = Etudiant.objects.create(matricule=f'M{k}', prenom='P', nom=f'N{k}', email=f'e{k}@t.local')
            Inscription.objects.create(etudiant=e, session=cls.session, licence=cls.licence, filiere=prog)

        # Parcours des promotions (Promotion_Licence_Session)
        cls.s2324 = Session.objects.create(annee='2023-2024')
        cls.s2425 = Session.objects.create(annee='2024-2025')
        cls.l2 = Licence.objects.create(nom='Licence 2')
        cls.l3 = Licence.objects.create(nom='Licence 3')
        cls.p18 = Promotion.objects.create(nom='P18')      # 3 niveaux
        cls.p20 = Promotion.objects.create(nom='P20')      # 1 niveau
        Promotion.objects.create(nom='P21')                # aucun niveau défini
        for promo, lic, ses in [(cls.p18, cls.licence, cls.s2324), (cls.p18, cls.l2, cls.s2425),
                                (cls.p18, cls.l3, cls.session), (cls.p20, cls.licence, cls.session)]:
            Promotion_Licence_Session.objects.create(promotion=promo, licence=lic, session=ses)

    def connecter(self, login):
        self.client.login(username=login, password='x')

    def donnees(self, **params):
        return self.client.get(reverse('etudiants_donnees'), params).json()

    def test_perimetre_par_role(self):
        self.connecter('scol')
        self.assertEqual(self.donnees()['total'], 3)
        self.connecter('dep')
        self.assertEqual(self.donnees()['total'], 3)
        self.connecter('info')
        self.assertEqual(self.donnees()['total'], 2)
        # un programme ne peut pas élargir son périmètre via les filtres
        self.assertEqual(self.donnees(programme=self.gc.id)['total'], 0)

    def test_scolarite_reconnue_par_son_nom(self):
        """« Service Scolarité » (code 0000, comme dans la base de l'IST) donne le rôle scolarité."""
        svc = Service.objects.create(nom='Service Scolarité', code='0000')
        u = User.objects.create_user('sc_nom', password='x')
        Personnel.objects.create(utilisateur=u, prenom='S', nom='C', telephone='99', email='sc@t.local',
                                 fonc='Agent').service.add(svc)
        self.connecter('sc_nom')
        self.assertEqual(self.client.get(reverse('importer')).status_code, 200)
        # un autre service au code 0000 reste un simple service
        dg = Service.objects.create(nom='Direction Générale', code='0000')
        u2 = User.objects.create_user('dg', password='x')
        Personnel.objects.create(utilisateur=u2, prenom='D', nom='G', telephone='98', email='dg@t.local',
                                 fonc='Agent').service.add(dg)
        self.connecter('dg')
        self.assertEqual(self.client.get(reverse('importer')).status_code, 403)
        self.assertEqual(self.client.get(reverse('etudiants_liste')).status_code, 200)

    def test_attestations_sans_image_de_signature(self):
        """Serveur où l'image de signature (hors Git) n'a pas encore été copiée : les PDF s'impriment quand même."""
        import datetime
        import shutil
        import tempfile
        from django.conf import settings
        from django.test import override_settings
        from .attestations import pdf_attestation, signature_disponible
        from .models import Attestation
        dossier = pathlib_path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, dossier, True)
        for image in settings.GPED_ATTESTATION['dossier_images'].glob('*.png'):
            if image.name != 'signature_scolarite.png':
                shutil.copy(image, dossier / image.name)
        e = Etudiant.objects.get(matricule='M0')
        with override_settings(GPED_ATTESTATION={**settings.GPED_ATTESTATION, 'dossier_images': dossier}):
            self.assertFalse(signature_disponible())
            for type_ in (Attestation.ORIENTATION, Attestation.INSCRIPTION, Attestation.NIVEAU, Attestation.FIN_ETUDES):
                a = Attestation(num_ref='0001/ISTM/SC/26', type=type_, etudiant=e, licence=self.licence,
                                session=self.session, filiere=self.info, cycle='Licence')
                a.date_delivrance = datetime.date(2026, 10, 22)
                pdf = pdf_attestation(a, 'https://exemple/v/abc', cachet_signature=True)
                self.assertEqual(pdf[:4], b'%PDF', type_)

    def test_attestation_inscription_titre_numero_et_pdf(self):
        from pfe.attestations import contenu_inscription, titre_inscription
        from .models import Attestation, Semestre
        Semestre.objects.create(nom='Semestre 3', abrege='S3', licence=self.l2)
        Semestre.objects.create(nom='Semestre 4', abrege='S4', licence=self.l2)
        # Titre selon le niveau
        self.assertEqual(titre_inscription(self.licence), "ATTESTATION D'INSCRIPTION")          # Licence 1
        self.assertEqual(titre_inscription(Licence(nom='1ère Année')), "ATTESTATION D'INSCRIPTION")
        self.assertEqual(titre_inscription(Licence(nom='M1')), "ATTESTATION D'INSCRIPTION")
        self.assertEqual(titre_inscription(self.l2), "ATTESTATION DE REINSCRIPTION")

        e = Etudiant.objects.get(matricule='M0')
        insc_l2 = Inscription.objects.create(etudiant=e, session=self.s2425, licence=self.l2, filiere=self.info)
        insc_l1 = Inscription.objects.get(etudiant=e, licence=self.licence)
        self.connecter('scol')

        url = reverse('attestations_type', args=['inscription'])

        # Formulaire « Enregistrer » : matricule + licence ; session et programme viennent de l'inscription
        r = self.client.get(url, {'matricule': 'M0', 'licence': self.l2.pk})
        self.assertContains(r, 'Enregistrer')
        self.assertContains(r, 'value="M0"')
        r = self.client.post(url, {'matricule': 'm0', 'licence': self.l2.pk})
        a1 = Attestation.objects.get()
        self.assertRedirects(r, f'{url}?imprimer={a1.pk}', fetch_redirect_response=False)
        self.assertEqual((a1.type, a1.etudiant, a1.session, a1.filiere, a1.licence),
                         (Attestation.INSCRIPTION, e, self.s2425, self.info, self.l2))
        self.assertEqual(a1.session_actuelle, self.session)
        self.assertTrue(a1.num_ref.startswith('0001/ISTM/SC/'))
        self.assertContains(self.client.get(f'{url}?imprimer={a1.pk}'), 'prête')
        r = self.client.get(reverse('attestation_pdf', args=[a1.pk]))
        self.assertEqual((r.status_code, r['Content-Type'], r.content[:4]), (200, 'application/pdf', b'%PDF'))

        # Doublon refusé (même étudiant, licence, session) ; erreurs claires
        r = self.client.post(url, {'matricule': 'M0', 'licence': self.l2.pk})
        self.assertContains(r, 'existe déjà')
        self.assertContains(self.client.post(url, {'matricule': 'XXX', 'licence': self.l2.pk}), 'Aucun étudiant')
        self.assertContains(self.client.post(url, {'matricule': 'M0', 'licence': self.l3.pk}), "aucune inscription")
        self.assertEqual(Attestation.objects.count(), 1)

        # Deuxième attestation : n° suivant
        self.client.post(url, {'matricule': 'M0', 'licence': self.licence.pk})
        a2 = Attestation.objects.order_by('id').last()
        self.assertEqual(a2.num_ref[:4], '0002')

        # Modification : le n° ne change pas, session/programme recalculés
        r = self.client.post(f'{url}?modifier={a2.pk}', {'matricule': 'M1', 'licence': self.licence.pk})
        a2.refresh_from_db()
        self.assertEqual((a2.num_ref[:4], a2.etudiant.matricule), ('0002', 'M1'))

        # Suppression de la DERNIÈRE attestation : son n° n'est pas réattribué
        self.client.post(reverse('attestation_supprimer', args=[a2.pk]))
        self.assertEqual(Attestation.objects.count(), 1)
        self.assertEqual(Attestation.prochaine_reference()[:4], '0003')

        # Le n° suit le plus grand numéro de l'année, tous types confondus (jamais réutilisé)
        Attestation.objects.create(type=Attestation.NIVEAU, etudiant=e)                    # 0003
        Attestation.objects.create(type=Attestation.NIVEAU, etudiant=e)                    # 0004
        Attestation.objects.filter(num_ref__startswith='0004').delete()
        self.assertEqual(Attestation.prochaine_reference()[:4], '0005')

        # Contenu : semestres et année, programme, session
        titre, (texte, _) = contenu_inscription(a1)
        self.assertEqual(titre, 'ATTESTATION DE REINSCRIPTION')
        for attendu in ('<b>Semestres S3 et S4</b> (Licence 2)', 'Matricule : <b>M0</b>',
                        'Département <b>Informatique</b>', '<b>2024-2025</b>'):
            self.assertIn(attendu, texte)
        # Niveau réel, tel qu'enregistré (jamais converti en « 4ème Année »)
        from .attestations import niveau_reel
        self.assertEqual(niveau_reel(Licence(nom='LICENCE 4')), 'LICENCE 4')
        self.assertEqual(niveau_reel(Licence(nom=' 1ère  Année ')), '1ère Année')

        # Réservé à la scolarité ; liste et recherche
        r = self.client.get(reverse('attestations'), {'etudiant': 'M0'})
        self.assertContains(r, a1.num_ref)
        self.connecter('info')
        self.assertEqual(self.client.get(reverse('attestations')).status_code, 403)
        self.assertEqual(self.client.get(reverse('attestation_pdf', args=[a1.pk])).status_code, 403)

    def test_attestation_orientation(self):
        from pfe.attestations import contenu_orientation
        from .models import Attestation
        e = Etudiant.objects.get(matricule='M0')
        Etudiant.objects.filter(pk=e.pk).update(ine='INE1', pvbac='303462', option_de_bac='Sciences mathématiques',
                                                centre_de_bac='GSP LE MERITE', annee_orientation='2023')
        Inscription.objects.create(etudiant=e, session=self.s2324, licence=self.l2, filiere=self.gc)  # plus ancienne
        self.connecter('scol')
        url = reverse('attestations_type', args=['orientation'])

        # Une inscription d'abord : la numérotation est continue entre les types
        self.client.post(reverse('attestations_type', args=['inscription']), {'matricule': 'M1', 'licence': self.licence.pk})

        r = self.client.get(url)
        self.assertContains(r, 'Enregistrer')
        self.assertContains(r, '0002/ISTM/DG/')                       # n° suivant, suffixe DG
        r = self.client.post(url, {'matricule': 'M0'})
        a = Attestation.objects.get(type=Attestation.ORIENTATION)
        self.assertRedirects(r, f'{url}?imprimer={a.pk}', fetch_redirect_response=False)
        self.assertTrue(a.num_ref.startswith('0002/ISTM/DG/'))
        # Département = programme de la DERNIÈRE inscription (2025-2026, Informatique)
        self.assertEqual((a.filiere, a.session, a.licence), (self.info, self.session, self.licence))

        titre, (texte, _), _style = contenu_orientation(a)
        self.assertEqual(titre[0][:30], 'ATTESTATION DE REUSSITE A LA S')
        for attendu in ('INE : <b>INE1</b>', 'PV Bac Unique : <b>303462</b>', 'Profil : <b>Sciences mathématiques</b>',
                        'Centre : <b>GSP LE MERITE</b>', 'Session : <b>2023</b>', 'Département : <b>Informatique</b>'):
            self.assertIn(attendu, texte)
        r = self.client.get(reverse('attestation_pdf', args=[a.pk]))
        self.assertEqual((r.status_code, r.content[:4]), (200, b'%PDF'))

        # Une seule par étudiant ; fiche incomplète -> avertissement, « — » à l'impression
        self.assertContains(self.client.post(url, {'matricule': 'M0'}), "déjà une attestation d&#x27;orientation")
        r = self.client.post(url, {'matricule': 'M2'}, follow=True)
        self.assertContains(r, 'Informations absentes')
        self.assertEqual(Attestation.objects.filter(type=Attestation.ORIENTATION).count(), 2)

    def test_attestation_niveau(self):
        from pfe.attestations import contenu_niveau
        from .models import Attestation, Semestre
        Semestre.objects.create(nom='Semestre 3', abrege='S3', licence=self.l2)
        Semestre.objects.create(nom='Semestre 4', abrege='S4', licence=self.l2)
        e = Etudiant.objects.get(matricule='M0')
        Inscription.objects.create(etudiant=e, session=self.s2425, licence=self.l2, filiere=self.info)
        self.connecter('scol')
        url = reverse('attestations_type', args=['niveau'])

        # Inscription puis niveau pour la même année : autorisé (types différents), n° continus
        self.client.post(reverse('attestations_type', args=['inscription']), {'matricule': 'M0', 'licence': self.l2.pk})
        self.assertContains(self.client.get(url), 'Enregistrer')
        r = self.client.post(url, {'matricule': 'M0', 'licence': self.l2.pk})
        a = Attestation.objects.get(type=Attestation.NIVEAU)
        self.assertRedirects(r, f'{url}?imprimer={a.pk}', fetch_redirect_response=False)
        self.assertTrue(a.num_ref.startswith('0002/ISTM/SC/'))
        self.assertEqual((a.session, a.filiere, a.licence), (self.s2425, self.info, self.l2))

        titre, (texte, _) = contenu_niveau(a)
        self.assertEqual(titre, 'ATTESTATION DE NIVEAU')
        for attendu in ('Je soussigné, atteste', '<b>P N0</b>', 'Matricule : <b>M0</b>', 'Département : <b>Informatique</b>',
                        'cours théoriques et pratiques des Semestres <b>S3 et S4</b> (Licence 2)', '<b>2024-2025</b>'):
            self.assertIn(attendu, texte)
        self.assertEqual(self.client.get(reverse('attestation_pdf', args=[a.pk])).content[:4], b'%PDF')

        # Doublon de niveau refusé ; page publique de vérification
        self.assertContains(self.client.post(url, {'matricule': 'M0', 'licence': self.l2.pk}), 'attestation de niveau existe')
        self.client.logout()
        r = self.client.get(reverse('verifier_document', args=[a.code_verification]))
        self.assertContains(r, 'ATTESTATION DE NIVEAU')
        self.assertContains(r, 'Semestres S3 et S4 (Licence 2)')

    def test_attestation_fin_etudes_avec_cycle(self):
        from pfe.attestations import contenu_fin_etudes, placer_texte, titre_document
        from .models import Attestation
        e = Etudiant.objects.get(matricule='M0')
        Etudiant.objects.filter(pk=e.pk).update(ine='DIHO0807698329')
        Inscription.objects.filter(etudiant=e).update(cycle='Licence')
        Inscription.objects.create(etudiant=e, session=self.s2324, licence=self.l2, filiere=self.gc, cycle='Licence')
        self.connecter('scol')
        url = reverse('attestations_type', args=['fin-etudes'])
        self.assertContains(self.client.get(url), 'Cycle terminé')

        # Sans choix de cycle : celui de la dernière inscription (Licence 2025-2026, Informatique)
        r = self.client.post(url, {'matricule': 'M0', 'cycle': ''})
        a = Attestation.objects.get(type=Attestation.FIN_ETUDES)
        self.assertRedirects(r, f'{url}?imprimer={a.pk}', fetch_redirect_response=False)
        self.assertEqual((a.cycle, a.filiere, a.session), ('Licence', self.info, self.session))
        self.assertTrue(a.num_ref.startswith('0001/ISTM/SC/'))

        titre, (texte, fin) = contenu_fin_etudes(a)
        self.assertEqual(titre, "ATTESTATION DE FIN D'ETUDES SUPERIEURES")
        for attendu in ('M ou Mme', '<b>P N0</b>', 'INE : <b>DIHO0807698329</b>', 'Département <b>Informatique</b>',
                        'pratiques du <b>Premier cycle LMD (Licence)</b> de notre institution'):
            self.assertIn(attendu, texte)
        self.assertEqual(titre_document(a), "ATTESTATION DE FIN D'ETUDES SUPERIEURES (Premier cycle LMD)")
        self.assertEqual(self.client.get(reverse('attestation_pdf', args=[a.pk])).content[:4], b'%PDF')

        # Un seul par cycle ; un autre cycle sans inscription est refusé
        self.assertContains(self.client.post(url, {'matricule': 'M0', 'cycle': 'Licence'}), 'déjà une attestation de fin')
        self.assertContains(self.client.post(url, {'matricule': 'M0', 'cycle': 'Master'}), 'aucune inscription en Master')
        # Cycle absent de l'inscription et non choisi : on demande de le préciser
        self.assertContains(self.client.post(url, {'matricule': 'M1', 'cycle': ''}), "Le cycle n&#x27;est pas renseigné")

        # Page publique : cycle validé
        self.client.logout()
        r = self.client.get(reverse('verifier_document', args=[a.code_verification]))
        self.assertContains(r, 'Premier cycle LMD (Licence)')
        self.assertContains(r, 'DIHO0807698329')

    def test_fin_etudes_cachet_signature_reserves_au_super_utilisateur(self):
        from .models import Attestation
        e = Etudiant.objects.get(matricule='M0')
        Inscription.objects.filter(etudiant=e).update(cycle='Licence')
        a = Attestation.objects.create(type=Attestation.FIN_ETUDES, etudiant=e, licence=self.licence,
                                       session=self.session, filiere=self.info, cycle='Licence',
                                       num_ref=Attestation.prochaine_reference(Attestation.FIN_ETUDES))
        pdf = reverse('attestation_pdf', args=[a.pk])
        liste = reverse('attestations_type', args=['fin-etudes'])
        images = lambda r: r.content.count(b'/Subtype /Image')            # noqa: E731

        # Scolarité : sans cachet ni signature, même en ajoutant « ?signature=1 » à la main
        self.connecter('scol')
        sans = self.client.get(pdf)
        force = self.client.get(pdf + '?signature=1')
        self.assertEqual(images(force), images(sans))
        self.assertNotIn('signee', force['Content-Disposition'])
        self.assertEqual(sans['Cache-Control'], 'no-store, private')
        self.assertNotContains(self.client.get(liste), '?signature=1')

        # Super utilisateur : bouton et version avec cachet et signature (image en plus)
        User.objects.create_superuser('admin', password='x')
        self.client.login(username='admin', password='x')
        self.assertContains(self.client.get(liste), f'{pdf}?signature=1')
        avec = self.client.get(pdf + '?signature=1')
        self.assertGreater(images(avec), images(sans))
        self.assertIn('_signee.pdf', avec['Content-Disposition'])
        self.assertEqual(images(self.client.get(pdf)), images(sans))     # « Imprimer » : sans, pour lui aussi

        # Les autres attestations gardent leur signature, et l'option ne s'y applique pas
        insc = Attestation.objects.create(type=Attestation.INSCRIPTION, etudiant=e, licence=self.licence,
                                          session=self.session, filiere=self.info,
                                          num_ref=Attestation.prochaine_reference(Attestation.INSCRIPTION))
        r_insc = reverse('attestation_pdf', args=[insc.pk])
        self.assertEqual(images(self.client.get(r_insc)), images(self.client.get(r_insc + '?signature=1')))
        self.assertGreater(images(self.client.get(r_insc)), images(sans))

    def test_texte_trop_long_jamais_perdu(self):
        """Un texte qui déborde est réduit, jamais tronqué (tous les paragraphes sont dessinés)."""
        import io as _io
        from reportlab.pdfgen.canvas import Canvas
        from pfe.attestations import STYLE_TEXTE, placer_texte
        c = Canvas(_io.BytesIO(), pageCompression=0)
        long = 'Département ' + 'TRES LONG NOM ' * 60
        placer_texte(c, [long, 'En foi de quoi, dernier paragraphe.'], STYLE_TEXTE, 60, 240, 700, 405)
        c.showPage()
        self.assertIn(b'En foi de quoi', c.getpdfdata().replace(b'\\', b''))

    def test_verification_publique_par_qr_code(self):
        from .models import Attestation
        a = Attestation.objects.create(type=Attestation.INSCRIPTION, etudiant=Etudiant.objects.get(matricule='M0'),
                                       licence=self.licence, filiere=self.info, session=self.session)
        b = Attestation.objects.create(type=Attestation.INSCRIPTION, etudiant=Etudiant.objects.get(matricule='M1'),
                                       licence=self.licence, filiere=self.info, session=self.session)
        self.assertEqual(len(a.code_verification), 22)
        self.assertNotEqual(a.code_verification, b.code_verification)     # aléatoire, non déductible

        # Page publique : sans connexion, elle confirme et affiche le contenu du document
        self.client.logout()
        r = self.client.get(reverse('verifier_document', args=[a.code_verification]))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['X-Robots-Tag'], 'noindex, nofollow')
        for attendu in ('Document authentique', a.num_ref, "ATTESTATION D&#x27;INSCRIPTION", 'P N0', 'M0',
                        'Informatique', '2025-2026'):
            self.assertContains(r, attendu)
        self.assertNotContains(r, 'e0@t.local')                           # rien d'autre sur l'étudiant

        # Code absent de la base (inventé, trafiqué, trop long, injection…) : page gePed « non reconnu »
        for faux in ['CodeInventeABCDEFGHIJ12', 'faux-code', 'x', 'A' * 300, '<script>alert(1)</script>',
                     "1' OR '1'='1", 'a/b/c', a.code_verification[:-1] + ('A' if a.code_verification[-1] != 'A' else 'B')]:
            r = self.client.get('/verifier/' + quote(faux, safe='') + '/')
            self.assertContains(r, 'Document non reconnu', status_code=404, msg_prefix=faux)
            self.assertNotContains(r, '<script>alert', status_code=404)
        code = b.code_verification
        b.delete()
        self.assertEqual(self.client.get(reverse('verifier_document', args=[code])).status_code, 404)

        # Le PDF contient le QR code : on vérifie que l'adresse publique est bien utilisée
        from pfe.views import url_verification
        from django.test import RequestFactory, override_settings
        req = RequestFactory().get('/')
        self.assertEqual(url_verification(req, a), f'http://testserver/verifier/{a.code_verification}/')
        with override_settings(GPED_URL_PUBLIQUE='https://geped.istmamou.org'):
            self.assertEqual(url_verification(req, a), f'https://geped.istmamou.org/verifier/{a.code_verification}/')
        self.connecter('scol')
        r = self.client.get(reverse('attestation_pdf', args=[a.pk]))
        self.assertEqual((r.status_code, r.content[:4]), (200, b'%PDF'))

    def test_sans_service_refuse(self):
        self.connecter('sans_service')
        self.assertEqual(self.client.get(reverse('etudiants_liste')).status_code, 403)

    def test_actions_reservees_scolarite(self):
        insc = Inscription.objects.filter(filiere=self.info).first()
        for login in ('dep', 'info'):
            self.connecter(login)
            self.assertEqual(self.client.get(reverse('importer')).status_code, 403)
            self.assertEqual(self.client.get(reverse('inscription_creer')).status_code, 403)
            self.assertEqual(self.client.get(reverse('inscription_modifier', args=[insc.pk])).status_code, 403)
        self.connecter('scol')
        self.assertEqual(self.client.get(reverse('importer')).status_code, 200)

    def test_fiche_hors_perimetre_404(self):
        autre = Inscription.objects.get(filiere=self.gc)
        self.connecter('info')
        self.assertEqual(self.client.get(reverse('etudiant_fiche', args=[autre.pk])).status_code, 404)

    def test_export_respecte_perimetre(self):
        self.connecter('info')
        r = self.client.get(reverse('etudiants_export'), {'format': 'csv'})
        contenu = b''.join(r.streaming_content).decode('utf-8-sig')
        self.assertEqual(len(contenu.strip().splitlines()), 3)   # en-tête + 2
        r = self.client.get(reverse('etudiants_export'), {'format': 'xlsx'})
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.max_row, 3)

    def test_recherche_et_tri(self):
        self.connecter('scol')
        self.assertEqual(self.donnees(q='N2')['total'], 1)
        d = self.donnees(tri='nom', ordre='desc')
        self.assertEqual(d['lignes'][0][3], 'N2')

    def test_import_une_ligne_une_inscription(self):
        entetes = ['matricule', 'prenom', 'nom', 'sexe', 'date_de_naissance', 'email',
                   'programme', 'promotion', 'session', 'licence']
        f = fichier_xlsx([
            entetes,
            ['2601', 'Awa', 'NOUVELLE', 'Femme', None, None, 'Génie civil', 'p 18', '2023 - 2024', 'L1'],   # 2 crée
            [2601, 'Awa', 'NOUVELLE', 'F', None, None, 'Génie civil', 'P18', '2024-2025', 'Licence 2'],     # 3 crée
            ['2601', 'Awa', 'NOUVELLE', 'F', None, None, 'INFO', 'P18', '2025/2026', 'LICENCE 3'],          # 4 crée
            ['M0', 'Paul', 'CHANGE', 'M', None, None, 'INFO', 'P20', '2025-2026', 'Licence 1'],             # 5 déjà en base
            ['2601', 'Awa', 'NOUVELLE', 'F', None, None, 'INFO', 'P18', '2023-2024', 'Licence 1'],          # 6 doublon fichier
            [None] * 10,                                                                                   # 7 vide
            ['M9', 'X', 'Y', 'M', None, None, 'INCONNU', 'P20', '2025-2026', 'L1'],
            ['M8', 'X', 'Y', 'M', None, None, 'INFO', 'P99', '2025-2026', 'L1'],
            ['M7', 'X', 'Y', 'M', None, None, 'INFO', 'P20', '2030-2031', 'L1'],
            ['M6', 'X', 'Y', 'M', None, None, 'INFO', 'P20', '2025-2026', 'Licence 9'],
            ['', 'X', 'Y', 'M', None, None, 'INFO', None, '2025-2026', 'L1'],
            ['M5', 'X', 'Y', 'M', 'pas une date', None, 'INFO', 'P20', '2025-2026', 'L1'],
            ['M4', 'Z', 'Z', 'M', None, 'e0@t.local', 'INFO', 'P20', '2025-2026', 'L1'],                     # e-mail pris
        ])
        total = Inscription.objects.count()
        rapport = importer_etudiants(f, simulation=True)
        self.assertEqual((rapport.lignes, rapport.inscriptions_creees, rapport.etudiants_crees,
                          rapport.etudiants_existants, rapport.nb_doublons, rapport.nb_erreurs), (12, 3, 1, 1, 2, 7))
        messages = ' '.join(m for _, m in rapport.erreurs)
        for attendu in ('programme inconnu « INCONNU »',
                        'promotion inconnue « P99 »', 'session inconnue « 2030-2031 »', 'licence inconnue « Licence 9 »',
                        'valeur manquante : matricule, promotion', 'date invalide', 'e-mail e0@t.local déjà utilisé'):
            self.assertIn(attendu, messages)
        self.assertEqual([n for n, _ in rapport.erreurs], [8, 9, 10, 11, 12, 13, 14])
        # Doublons (en base ou répétés dans le fichier) : ignorés, ce ne sont pas des erreurs
        self.assertEqual(rapport.doublons, [(5, 'M0 · Licence 1 · 2025-2026 : déjà inscrit(e) dans gePed'),
                                            (6, '2601 · Licence 1 · 2023-2024 : déjà présente à la ligne 2 du fichier')])
        self.assertFalse(Etudiant.objects.filter(matricule='2601').exists())       # simulation : rien d'écrit

        f.seek(0)
        rapport = importer_etudiants(f)
        self.assertEqual((rapport.inscriptions_creees, rapport.nb_doublons), (3, 2))
        # Chaque ligne telle quelle : session, licence, programme et promotion du fichier
        lignes = Inscription.objects.filter(etudiant__matricule='2601').order_by('session__annee')
        self.assertEqual([(i.session.annee, i.licence.nom, i.filiere.nom, i.promotion.nom) for i in lignes],
                         [('2023-2024', 'Licence 1', 'Génie civil', 'P18'),
                          ('2024-2025', 'Licence 2', 'Génie civil', 'P18'),
                          ('2025-2026', 'Licence 3', 'Informatique', 'P18')])
        nouvelle = Etudiant.objects.get(matricule='2601')
        self.assertEqual((nouvelle.sexe, nouvelle.email), ('F', '2601@etudiant.gped'))
        # L'étudiant existant n'est pas modifié, son inscription n'est pas dupliquée
        m0 = Etudiant.objects.get(matricule='M0')
        self.assertEqual((m0.nom, Inscription.objects.filter(etudiant=m0).count()), ('N0', 1))
        self.assertEqual(Inscription.objects.count(), total + 3)

        # Réimport : tout est déjà là
        f.seek(0)
        rapport = importer_etudiants(f)
        self.assertEqual((rapport.inscriptions_creees, rapport.etudiants_crees, rapport.nb_doublons), (0, 0, 5))
        self.assertEqual(Inscription.objects.count(), total + 3)

    def test_import_doublons_vus_en_simulation_ignores_a_l_import(self):
        csv = ('matricule;prenom;nom;programme;promotion;session;licence\n'
               'M0;Paul;N0;Informatique;P20;2025-2026;Licence 1\n'           # 2 déjà en base
               'D1;Ana;BAH;Informatique;P20;2025-2026;Licence 1\n'           # 3 nouvelle
               'D1;Ana;BAH;Informatique;P20;2025-2026;Licence 1\n').encode('utf-8')   # 4 répétée
        fichier = lambda: SimpleUploadedFile('l.csv', csv)                 # noqa: E731
        self.connecter('scol')
        # Simulation : la liste des doublons est affichée
        r = self.client.post(reverse('importer'), {'fichier': fichier(), 'simulation': 'on'})
        self.assertContains(r, 'Inscription déjà existante')
        self.assertContains(r, 'déjà présente à la ligne 3 du fichier')
        self.assertEqual(r.context['rapport'].nb_erreurs, 0)
        # Import réel : les doublons sont ignorés (compteur seulement), les autres lignes enregistrées
        r = self.client.post(reverse('importer'), {'fichier': fichier()})
        self.assertNotContains(r, 'Inscription déjà existante')
        self.assertContains(r, '2 doublons ignorés')
        self.assertEqual(Inscription.objects.filter(etudiant__matricule='D1').count(), 1)
        self.assertEqual(Inscription.objects.filter(etudiant__matricule='M0').count(), 1)

    def test_import_preinscription_genere_les_matricules(self):
        from .importer import code_lettre
        self.assertEqual((code_lettre('Ana'), code_lettre('éric'), code_lettre('  moussa')), ('35', '39', '47'))
        tc = Service.objects.create(nom='Tronc commun', code='1', programme=True)      # code « 1 » → « 01 »
        Etudiant.objects.create(matricule='2601000135', prenom='Déjà', nom='PRIS', email='pris@t.local')
        entete = 'prenom;nom;date_de_naissance;email;programme;promotion;session;licence\n'
        lignes = ('Ana;BAH;01/01/2005;;TRONC COMMUN;P20;2025-2026;Licence 1\n'       # 2
                  'Éric;SOW;;Eric@T.local;Tronc commun;P20;2025-2026;Licence 1\n'     # 3
                  'Ana;BAH;01/01/2005;;Tronc commun;P20;2025-2026;Licence 1\n'        # 4 doublon fichier
                  'Moussa;DIALLO;;;INFO;P20;2025 - 2026;L1\n')                          # 5
        f = lambda texte: SimpleUploadedFile('p.csv', (entete + texte).encode('utf-8'))  # noqa: E731

        # Inscrits : le matricule reste obligatoire
        with self.assertRaisesMessage(ValueError, '« matricule »'):
            importer_etudiants(f(lignes))

        r = importer_etudiants(f(lignes), preinscription=True)
        self.assertEqual((r.etudiants_crees, r.inscriptions_creees, r.nb_erreurs, r.nb_doublons), (3, 3, 0, 1))
        self.assertIn('déjà présente à la ligne 2 du fichier', r.doublons[0][1])
        # session 26 + code programme + compteur (inscriptions du programme + 1) + lettre du prénom ;
        # 2601000135 est déjà pris : on passe au compteur suivant
        self.assertEqual([(n, m) for n, _, _, m in r.matricules_generes],
                         [(2, '2601000235'), (3, '2601000339'), (5, '26INFO000347')])
        ana = Etudiant.objects.get(matricule='2601000235')
        self.assertEqual((ana.prenom, ana.preinscrit, ana.email), ('Ana', True, '2601000235@etudiant.gped'))
        self.assertEqual(Inscription.objects.get(etudiant=ana).filiere, tc)

        # Réimport : reconnus par date de naissance + nom + prénom, ou par e-mail → pas de doublon
        r = importer_etudiants(f(lignes.split('Ana;BAH;01/01/2005;;Tronc commun')[0]), preinscription=True)
        self.assertEqual((r.etudiants_crees, r.inscriptions_creees, r.nb_doublons, r.matricules_generes), (0, 0, 2, []))

    def test_import_preinscrit_reconnu_par_ine(self):
        csv = ('prenom;nom;ine;email;programme;promotion;session;licence\n'
               'Awa;DIALLO;diho0807698329;awa@t.local;Informatique;P20;2025-2026;Licence 1\n').encode('utf-8')
        importer_etudiants(SimpleUploadedFile('p.csv', csv), preinscription=True)
        awa = Etudiant.objects.get(nom='DIALLO')
        self.assertEqual(awa.ine, 'DIHO0807698329')                      # INE enregistré (en majuscules)
        # Même INE, e-mail différent, autre année : même étudiant, nouvelle inscription, pas de 2e matricule
        csv = ('prenom;nom;ine;email;programme;promotion;session;licence\n'
               'Awa;DIALLO;DIHO0807698329;autre@t.local;Informatique;P20;2024-2025;Licence 2\n').encode('utf-8')
        r = importer_etudiants(SimpleUploadedFile('p.csv', csv), preinscription=True)
        self.assertEqual((r.etudiants_crees, r.inscriptions_creees, r.matricules_generes), (0, 1, []))
        self.assertEqual(Etudiant.objects.filter(nom='DIALLO').count(), 1)
        self.assertEqual(awa.etudiant_inscription.count(), 2)
        awa.refresh_from_db()
        self.assertEqual(awa.email, 'autre@t.local')                       # fiche mise à jour

    def test_import_preinscription_ine_connu_met_a_jour_sans_toucher_au_matricule(self):
        # Étudiant déjà inscrit (matricule M0), fiche en partie remplie ; son INE revient en préinscription
        Etudiant.objects.filter(matricule='M0').update(ine='INE000', telephone='600', adresse='Mamou', preinscrit=False)
        Etudiant.objects.create(matricule='X9', prenom='Autre', nom='ETU', email='pris@t.local')
        entete = 'matricule;prenom;nom;ine;email;telephone;adresse;ville;programme;promotion;session;licence\n'
        # Autre matricule dans le fichier, même INE ; adresse vide ; e-mail d'un autre étudiant
        csv = entete + 'NOUVEAU1;P;N0;ine000;pris@t.local;622 00 00 00;;Labé;Informatique;P20;2024-2025;Licence 2\n'
        r = importer_etudiants(SimpleUploadedFile('p.csv', csv.encode('utf-8')), preinscription=True)
        self.assertEqual((r.etudiants_crees, r.etudiants_existants, r.etudiants_mis_a_jour, r.inscriptions_creees),
                         (0, 1, 1, 1))
        self.assertFalse(Etudiant.objects.filter(matricule='NOUVEAU1').exists())
        e = Etudiant.objects.get(ine='INE000')
        self.assertEqual((e.matricule, e.telephone, e.ville, e.adresse, e.email, e.preinscrit),
                         ('M0', '622 00 00 00', 'Labé', 'Mamou', 'e0@t.local', False))   # statut inchangé
        self.assertEqual(e.etudiant_inscription.count(), 2)
        # Simulation : rien n'est modifié
        csv = entete + ';P;N0;INE000;;699;;;Informatique;P20;2023-2024;Licence 1\n'
        r = importer_etudiants(SimpleUploadedFile('p.csv', csv.encode('utf-8')), simulation=True, preinscription=True)
        self.assertEqual(r.etudiants_mis_a_jour, 1)
        self.assertEqual(Etudiant.objects.get(matricule='M0').telephone, '622 00 00 00')
        # Import « inscription » : l'INE ne sert pas, la fiche n'est pas modifiée (comportement inchangé)
        csv = entete + 'M0;P;N0;INE000;;611;;;Informatique;P20;2023-2024;Licence 1\n'
        importer_etudiants(SimpleUploadedFile('p.csv', csv.encode('utf-8')))
        self.assertEqual(Etudiant.objects.get(matricule='M0').telephone, '622 00 00 00')

    def test_import_colonnes_obligatoires(self):
        csv = 'matricule;prenom;nom;promotion\nC1;Ana;BAH;P20\n'.encode('utf-8')
        with self.assertRaisesMessage(ValueError, '« programme », « session », « licence »'):
            importer_etudiants(SimpleUploadedFile('l.csv', csv))

    def test_import_csv_preinscription_et_cycle(self):
        csv = ('matricule;prenom;nom;programme;promotion;session;licence;cycle\n'
               'P1;Ana;BAH;Informatique;P20;2025-2026;Licence 1;\n'
               'P2;Ali;SOW;GENIE CIVIL;P20;2025-2026;Licence 1;Ingénieur\n'
               'P3;Awa;BARRY;Informatique;P20;2025-2026;Licence 1;Inconnu\n').encode('utf-8-sig')
        rapport = importer_etudiants(SimpleUploadedFile('l.csv', csv), cycle='DUT', preinscription=True)
        self.assertEqual((rapport.etudiants_crees, rapport.nb_erreurs), (2, 1))
        self.assertTrue(rapport.preinscription)
        p1 = Inscription.objects.get(etudiant__matricule='P1')
        self.assertEqual((p1.cycle, p1.filiere, p1.etudiant.preinscrit), ('DUT', self.info, True))  # cycle du formulaire
        p2 = Inscription.objects.get(etudiant__matricule='P2')
        self.assertEqual((p2.cycle, p2.filiere), ('Ingénieur', self.gc))      # colonne du fichier, programme sans accents

        # Import « inscription » d'une autre année : nouvelle inscription, l'étudiant n'est plus préinscrit,
        # le reste de sa fiche ne change pas
        csv = ('matricule;prenom;nom;programme;promotion;session;licence\n'
               'P1;Autre;NOM;Informatique;P20;2024-2025;Licence 2\n').encode('utf-8')
        importer_etudiants(SimpleUploadedFile('l.csv', csv))
        e = Etudiant.objects.get(matricule='P1')
        self.assertEqual((e.prenom, e.nom, e.preinscrit, e.etudiant_inscription.count()), ('Ana', 'BAH', False, 2))

    def test_inscription_formulaire(self):
        self.connecter('scol')
        r = self.client.post(reverse('inscription_creer'), {
            'etu-prenom': 'Ana', 'etu-nom': 'SOW', 'etu-sexe': 'F', 'etu-email': 'ana@t.local',
            'insc-session': self.session.id, 'insc-licence': self.licence.id, 'insc-filiere': self.info.id,
        })
        self.assertEqual(r.status_code, 302)
        self.assertFalse(Inscription.objects.get(etudiant__email='ana@t.local').redouble)

    def _media_temporaire(self):
        import shutil, tempfile
        from django.test import override_settings
        dossier = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, dossier, True)
        reglage = override_settings(MEDIA_ROOT=dossier)
        reglage.enable()
        self.addCleanup(reglage.disable)
        return dossier

    @staticmethod
    def _image(format_='PNG', taille=(60, 80), mode='RGB', nom=None, **options):
        from PIL import Image
        flux = io.BytesIO()
        couleur = (65, 98, 124, 120) if mode == 'RGBA' else (65, 98, 124)
        Image.new(mode, taille, couleur).save(flux, format_, **options)
        return SimpleUploadedFile(nom or f'photo.{format_.lower()}', flux.getvalue(), content_type=f'image/{format_.lower()}')

    def test_compression_photo(self):
        from django.core.exceptions import ValidationError
        from PIL import Image
        from .photos import POIDS_CIBLE, compresser

        # Grande photo de téléphone « couchée » (EXIF orientation 6) et très chargée (bruit) avec EXIF
        img = Image.effect_noise((2400, 1800), 90).convert('RGB')
        exif = Image.Exif()
        exif[0x0112] = 6                      # orientation : à tourner de 90°
        exif[0x010F] = 'Telephone test'       # marque de l'appareil
        flux = io.BytesIO()
        img.save(flux, 'JPEG', quality=95, exif=exif)
        resultat = compresser(SimpleUploadedFile('p.jpg', flux.getvalue()))
        sortie = Image.open(io.BytesIO(resultat))
        self.assertEqual((sortie.format, sortie.size), ('JPEG', (600, 800)))      # redressée puis réduite
        self.assertFalse(sortie.info.get('exif'))                                # métadonnées supprimées
        self.assertLessEqual(len(resultat), POIDS_CIBLE)
        # Au-delà de 10 Mo : refusé
        with self.assertRaisesMessage(ValidationError, '10 Mo maximum'):
            compresser(SimpleUploadedFile('lourde.jpg', b'0' * (10 * 1024 * 1024 + 1)))

        # PNG transparent → JPEG sur fond blanc
        sortie = Image.open(io.BytesIO(compresser(self._image('PNG', mode='RGBA'))))
        self.assertEqual((sortie.format, sortie.mode), ('JPEG', 'RGB'))
        # Déjà conforme (JPEG léger, sans EXIF, ≤ 600 × 800) : gardé tel quel
        deja = self._image('JPEG', (300, 400))
        brut = deja.read()
        self.assertEqual(compresser(deja), brut)
        # Pas une image
        with self.assertRaisesMessage(ValidationError, "pas une image lisible"):
            compresser(SimpleUploadedFile('cv.pdf', b'%PDF-1.4 pas une image'))

    def test_photo_formulaire_inscription(self):
        from .photos import POIDS_CIBLE
        self._media_temporaire()
        base = {'etu-prenom': 'Ana', 'etu-nom': 'SOW', 'etu-sexe': 'F', 'etu-email': 'photo@t.local',
                'insc-session': self.session.id, 'insc-licence': self.licence.id, 'insc-filiere': self.info.id}
        self.connecter('scol')
        html = self.client.get(reverse('inscription_creer')).content.decode()
        self.assertIn('Prendre ou charger une photo', html)
        self.assertIn('cropper.min.js', html)
        # Fichier qui n'est pas une image : refusé
        r = self.client.post(reverse('inscription_creer'),
                             {**base, 'etu-photo': SimpleUploadedFile('cv.pdf', b'%PDF', content_type='application/pdf')})
        self.assertContains(r, 'image lisible : choisissez une photo JPG, PNG ou WEBP')
        # Grande photo : compressée en JPEG à l'enregistrement, nommée d'après l'étudiant
        r = self.client.post(reverse('inscription_creer'), {**base, 'etu-photo': self._image('PNG', (1800, 2400))})
        insc = Inscription.objects.get(etudiant__email='photo@t.local')
        self.assertRedirects(r, reverse('etudiant_fiche', args=[insc.pk]) + '?imprimer=1', fetch_redirect_response=False)
        photo = insc.etudiant.photo
        self.assertRegex(photo.name, r'^etudiant/photo/etudiant-\d+_\d{14}\.jpg$')
        self.assertLessEqual(photo.size, POIDS_CIBLE)
        self.assertContains(self.client.get(reverse('etudiant_fiche', args=[insc.pk])), photo.url)
        # Après l'inscription : la fiche propose d'imprimer la fiche d'inscription (PDF léger, avec la photo)
        self.assertContains(self.client.get(reverse('etudiant_fiche', args=[insc.pk]), {'imprimer': 1}),
                            'Imprimer la fiche')
        r = self.client.get(reverse('inscription_fiche_pdf', args=[insc.pk]))
        self.assertEqual((r.status_code, r['Content-Type'], r.content[:4]), (200, 'application/pdf', b'%PDF'))
        self.assertIn('fiche_inscription_', r['Content-Disposition'])                 # Licence 1 : inscription
        self.assertIn(b'/Subtype /Image', r.content)
        self.assertLess(len(r.content), 120 * 1024)
        # Bourse et type de formation : cochés sur la fiche (formulaire → fiche de l'étudiant → PDF)
        from . import fiche_inscription
        cochees = []
        original = fiche_inscription._case
        fiche_inscription._case = lambda c, x, y, libelle, l, cochee=False: (cochees.append(libelle) if cochee else None,
                                                                             original(c, x, y, libelle, l, cochee))
        try:
            Etudiant.objects.filter(pk=insc.etudiant_id).update(boursier='payante', type_formation='initiale')
            self.client.get(reverse('inscription_fiche_pdf', args=[insc.pk]))
        finally:
            fiche_inscription._case = original
        self.assertEqual(cochees, ['Formation payante', 'Initiale'])
        # Autre niveau : fiche de réinscription
        from .fiche_inscription import titre_fiche
        self.assertEqual(titre_fiche(self.l2), "FICHE DE RÉINSCRIPTION")
        self.assertEqual(titre_fiche(Licence(nom='1ère Année')), "FICHE D'INSCRIPTION")
        self.connecter('dep')                                                     # réservé à la scolarité
        self.assertEqual(self.client.get(reverse('inscription_fiche_pdf', args=[insc.pk])).status_code, 403)
        self.connecter('scol')
        # Modification : « Retirer la photo » supprime aussi le fichier
        chemin = photo.path
        self.client.post(reverse('inscription_modifier', args=[insc.pk]), {**base, 'etu-retirer_photo': 'on'})
        insc.etudiant.refresh_from_db()
        self.assertFalse(insc.etudiant.photo)
        self.assertFalse(os.path.exists(chemin))

        # Étudiant existant : la page de modification enregistre (ou retire) la photo IMMÉDIATEMENT…
        modifier = reverse('inscription_modifier', args=[insc.pk])
        self.assertContains(self.client.get(modifier), f'data-photo-url="{reverse("etudiant_photo", args=[insc.etudiant_id])}"')
        self.client.post(reverse('etudiant_photo', args=[insc.etudiant_id]), {'photo': self._image()})
        insc.etudiant.refresh_from_db()
        self.assertTrue(insc.etudiant.photo)
        # … et enregistrer ensuite le formulaire (ouvert avant) ne remet jamais une ancienne photo
        self.client.post(reverse('etudiant_photo', args=[insc.etudiant_id]), {'retirer': '1'})
        self.client.post(modifier, base)
        insc.etudiant.refresh_from_db()
        self.assertFalse(insc.etudiant.photo)

    def test_photo_directe_et_droits(self):
        import os as _os
        self._media_temporaire()
        e = Etudiant.objects.get(matricule='M0')
        url = reverse('etudiant_photo', args=[e.pk])
        fiche = reverse('etudiant_fiche', args=[Inscription.objects.get(etudiant=e).pk])

        # Département : voit la fiche mais ne peut pas changer la photo
        self.connecter('dep')
        self.assertNotContains(self.client.get(fiche), 'data-photo-ouvrir')
        self.assertEqual(self.client.post(url, {'photo': self._image()}).status_code, 403)

        # Scolarité : enregistre ; une nouvelle photo remplace (et supprime) l'ancienne
        self.connecter('scol')
        self.assertContains(self.client.get(fiche), 'data-photo-ouvrir')
        d = self.client.post(url, {'photo': self._image()}).json()
        e.refresh_from_db()
        self.assertEqual(d['url'], e.photo.url)
        ancienne = e.photo.path
        d = self.client.post(url, {'photo': self._image('JPEG')}).json()
        e.refresh_from_db()
        self.assertFalse(_os.path.exists(ancienne))
        self.assertTrue(_os.path.exists(e.photo.path))
        r = self.client.post(url, {'photo': SimpleUploadedFile('x.txt', b'texte')})
        self.assertEqual((r.status_code, 'image lisible' in r.json()['erreur']), (400, True))

        # L'étudiant lui-même (compte lié à sa fiche) peut changer sa photo, pas celle d'un autre
        u = User.objects.create_user('etu_m0', password='x')
        Etudiant.objects.filter(pk=e.pk).update(utilisateur=u)
        self.client.login(username='etu_m0', password='x')
        self.assertEqual(self.client.post(url, {'retirer': '1'}).json(), {'url': None})
        autre = Etudiant.objects.get(matricule='M1')
        self.assertEqual(self.client.post(reverse('etudiant_photo', args=[autre.pk]),
                                          {'photo': self._image()}).status_code, 403)

    def test_photo_mon_compte(self):
        self._media_temporaire()
        self.connecter('info')
        self.assertContains(self.client.get(reverse('mon_compte')), 'data-photo-ouvrir')
        d = self.client.post(reverse('mon_compte_photo'), {'photo': self._image()}).json()
        p = Personnel.objects.get(utilisateur__username='info')
        self.assertEqual(d['url'], p.photo.url)
        self.assertTrue(p.photo.name.startswith('personnel/photo/info_'))
        # La photo s'affiche dans l'en-tête et le menu de toutes les pages
        self.assertContains(self.client.get(reverse('tableau_de_bord')), p.photo.url, count=2)
        # Compte sans fiche (ni personnel ni étudiant) : refus explicite
        self.client.login(username='sans_service', password='x')
        r = self.client.post(reverse('mon_compte_photo'), {'photo': self._image()})
        self.assertEqual(r.status_code, 400)

    def test_case_redoublant_decochee_par_defaut(self):
        self.connecter('scol')
        html = self.client.get(reverse('inscription_creer')).content.decode()
        champ = html.split('id="id_insc-redouble"')[0].rsplit('<input', 1)[1]
        self.assertNotIn('checked', champ)

    def test_filtres_cycle_et_statut(self):
        Inscription.objects.filter(etudiant__matricule='M0').update(cycle='Master')
        Etudiant.objects.filter(matricule='M1').update(preinscrit=True)
        self.connecter('scol')
        self.assertEqual(self.donnees(cycle='Master')['total'], 1)
        self.assertEqual(self.donnees(statut='preinscrit')['total'], 1)
        self.assertEqual(self.donnees(statut='inscrit')['total'], 2)
        # Liste des étudiants : jamais les préinscrits (ils ont leur page), export compris
        self.assertEqual(self.donnees()['total'], 2)
        self.assertEqual(self.donnees(statut='redoublant')['total'] + self.donnees(statut='passant')['total'], 2)
        r = self.client.get(reverse('etudiants_export'), {'format': 'csv', 'cycle': 'Master'})
        contenu = b''.join(r.streaming_content).decode('utf-8-sig')
        self.assertIn(';Master;', contenu)
        r = self.client.get(reverse('etudiants_export'), {'format': 'csv'})
        self.assertNotIn('M1;', b''.join(r.streaming_content).decode('utf-8-sig'))
        # Session : aucune par défaut (toutes) ; dès qu'une session est choisie, la liste s'y limite
        Inscription.objects.create(etudiant=Etudiant.objects.get(matricule='M0'), session=self.s2425,
                                   licence=self.l2, filiere=self.info)
        self.assertEqual(self.donnees()['total'], 3)
        self.assertEqual(self.donnees(session=self.session.pk)['total'], 2)
        self.assertEqual(self.donnees(session=self.s2425.pk)['total'], 1)
        self.assertEqual(self.donnees(session='toutes')['total'], 3)          # ancien lien
        html = self.client.get(reverse('etudiants_liste')).content.decode()
        filtre_session = html.split('name="session"', 1)[1].split('</select>', 1)[0]
        self.assertNotIn('selected', filtre_session)
        self.assertIn('<option value="">Toutes les sessions</option>', filtre_session)

    def test_mon_compte_coordonnees_seulement(self):
        self.connecter('info')
        self.assertEqual(self.client.get(reverse('mon_compte')).status_code, 200)
        r = self.client.post(reverse('mon_compte'), {
            'coordonnees': '1', 'email': 'Nouveau@T.local', 'telephone': '620 11 22 33',
            # champs non autorisés : doivent être ignorés
            'nom': 'PIRATE', 'fonc': 'Directeur', 'username': 'autre', 'service': self.scol.id,
        })
        self.assertEqual(r.status_code, 302)
        p = Personnel.objects.get(utilisateur__username='info')
        self.assertEqual((p.email, p.telephone, p.nom, p.fonc), ('nouveau@t.local', '620 11 22 33', 'T', 'Agent'))
        self.assertEqual(list(p.service.all()), [self.info])
        self.assertEqual(p.utilisateur.email, 'nouveau@t.local')
        self.assertEqual(p.utilisateur.username, 'info')

    def test_mon_compte_email_ou_telephone_deja_pris(self):
        self.connecter('info')
        r = self.client.post(reverse('mon_compte'), {'coordonnees': '1', 'email': 'scol@t.local', 'telephone': '0'})
        self.assertEqual(r.status_code, 200)
        self.assertIn('email', r.context['form_coord'].errors)
        self.assertIn('telephone', r.context['form_coord'].errors)
        self.assertEqual(Personnel.objects.get(utilisateur__username='info').email, 'info@t.local')

    def test_mon_compte_mot_de_passe(self):
        self.connecter('info')
        r = self.client.post(reverse('mon_compte'), {
            'mot_de_passe': '1', 'old_password': 'mauvais', 'new_password1': 'Nv@2026xyz', 'new_password2': 'Nv@2026xyz'})
        self.assertEqual(r.status_code, 200)               # ancien mot de passe faux
        r = self.client.post(reverse('mon_compte'), {
            'mot_de_passe': '1', 'old_password': 'x', 'new_password1': 'Nv@2026xyz', 'new_password2': 'Nv@2026xyz'})
        self.assertEqual(r.status_code, 302)
        self.assertTrue(User.objects.get(username='info').check_password('Nv@2026xyz'))
        # toujours connecté après le changement
        self.assertEqual(self.client.get(reverse('etudiants_liste')).status_code, 200)

    def test_mot_de_passe_oublie_parcours_complet(self):
        from django.core import mail
        # L'e-mail n'est que sur la fiche personnel (User.email vide)
        self.assertEqual(User.objects.get(username='info').email, '')
        r = self.client.post(reverse('password_reset'), {'email': 'INFO@t.local'})
        self.assertRedirects(r, reverse('password_reset_done'))
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ['info@t.local'])
        lien = re.search(r'https?://[^/]+(/reinitialiser/\S+/)', mail.outbox[0].body).group(1)

        r = self.client.get(lien, follow=True)                 # redirige vers .../set-password/
        self.assertTrue(r.context['validlink'])
        r = self.client.post(r.redirect_chain[-1][0], {'new_password1': 'Oubli@2026x', 'new_password2': 'Oubli@2026x'})
        self.assertRedirects(r, reverse('password_reset_complete'))
        self.assertTrue(User.objects.get(username='info').check_password('Oubli@2026x'))
        # Le lien ne peut pas resservir
        self.assertFalse(self.client.get(lien, follow=True).context['validlink'])

    def test_mot_de_passe_oublie_adresse_inconnue(self):
        from django.core import mail
        r = self.client.post(reverse('password_reset'), {'email': 'personne@nulle.part'})
        self.assertRedirects(r, reverse('password_reset_done'))   # même réponse : ne révèle rien
        self.assertEqual(len(mail.outbox), 0)

    def test_tableau_de_bord_par_role(self):
        self.connecter('scol')
        r = self.client.get(reverse('tableau_de_bord'))
        self.assertEqual(r.context['total'], 3)
        self.assertEqual({g['nom']: g['n'] for g in r.context['par_programme']}, {'Informatique': 2, 'Génie civil': 1})
        self.connecter('info')
        r = self.client.get(reverse('tableau_de_bord'))
        self.assertEqual(r.context['total'], 2)
        self.assertEqual([g['nom'] for g in r.context['par_programme']], ['Informatique'])
        self.assertNotContains(r, 'Nouvelle inscription')

    def test_preinscriptions_liste_et_inscription_directe(self):
        Etudiant.objects.filter(matricule__in=['M0', 'M2']).update(preinscrit=True)
        i0 = Inscription.objects.get(etudiant__matricule='M0')
        i2 = Inscription.objects.get(etudiant__matricule='M2')    # Génie civil

        # Un programme voit la page mais seulement ses préinscrits, et ne peut pas inscrire
        self.connecter('info')
        r = self.client.get(reverse('preinscriptions_liste'))
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.context['peut_inscrire'])
        self.assertEqual(r.context['nb_preinscrits'], 1)
        self.assertEqual(self.donnees(statut='preinscrit')['total'], 1)
        r = self.client.post(reverse('valider_preinscriptions'), {'ids': [i0.pk]})
        self.assertEqual(r.status_code, 403)

        # La scolarité inscrit directement (AJAX) ; un id déjà inscrit est ignoré
        self.connecter('scol')
        self.assertTrue(self.client.get(reverse('preinscriptions_liste')).context['peut_inscrire'])
        i1 = Inscription.objects.get(etudiant__matricule='M1')
        r = self.client.post(reverse('valider_preinscriptions'), {'ids': [i0.pk, i1.pk]},
                             HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(r.json(), {'inscrits': 1})
        self.assertFalse(Etudiant.objects.get(matricule='M0').preinscrit)

        # Depuis la fiche : formulaire classique avec retour sur la fiche
        fiche = reverse('etudiant_fiche', args=[i2.pk])
        r = self.client.post(reverse('valider_preinscriptions'), {'ids': [i2.pk], 'suivant': fiche})
        self.assertRedirects(r, fiche)
        self.assertFalse(Etudiant.objects.get(matricule='M2').preinscrit)
        # Redirection externe refusée
        r = self.client.post(reverse('valider_preinscriptions'), {'ids': [i2.pk], 'suivant': '//evil.example'})
        self.assertRedirects(r, reverse('preinscriptions_liste'))

    # ---------------------------------------------------------------- matières

    def _semestres(self):
        from .models import Semestre
        Service.objects.filter(pk__in=[self.info.pk, self.gc.pk]).update(cycle='Licence')   # cycle des programmes
        s1 = Semestre.objects.create(nom='SEMESTRE 1', abrege='S1', licence=self.licence)
        s3 = Semestre.objects.create(nom='SEMESTRE 3', abrege='S3', licence=self.l2)
        return s1, s3

    def _matiere(self, nom, programme, semestre, **extra):
        from .models import Matiere
        return Matiere.objects.create(nom=nom, filiere=programme, semestre=semestre, licence=semestre.licence,
                                      session=self.session, credit=extra.pop('credit', 6), **extra)

    def test_matieres_droits_par_role(self):
        s1, _ = self._semestres()
        m_info = self._matiere('Algorithmique', self.info, s1)
        m_gc = self._matiere('Résistance des matériaux', self.gc, s1)
        liste = reverse('matieres_liste')

        # Scolarité : tout voir, tout gérer
        self.connecter('scol')
        r = self.client.get(liste)
        self.assertContains(r, 'Algorithmique')
        self.assertContains(r, 'Résistance des matériaux')
        self.assertContains(r, reverse('matiere_creer'))

        # Autre service (département) : voit tout, ne modifie rien
        self.connecter('dep')
        r = self.client.get(liste)
        self.assertContains(r, 'Résistance des matériaux')
        self.assertNotContains(r, reverse('matiere_creer'))
        self.assertEqual(self.client.get(reverse('matiere_creer')).status_code, 403)
        self.assertEqual(self.client.get(reverse('matiere_modifier', args=[m_info.pk])).status_code, 403)

        # Chef de programme (Informatique) : uniquement ses matières ; celles des autres « n'existent pas »
        self.connecter('info')
        r = self.client.get(liste)
        self.assertContains(r, 'Algorithmique')
        self.assertNotContains(r, 'Résistance des matériaux')
        self.assertEqual(self.client.get(reverse('matiere_modifier', args=[m_gc.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse('matiere_supprimer', args=[m_gc.pk])).status_code, 404)
        # Liste déroulante : son seul programme, choisi d'office ; un autre programme est refusé
        r = self.client.get(reverse('matiere_creer'))
        self.assertEqual(list(r.context['form'].fields['filiere'].queryset), [self.info])
        self.assertEqual(r.context['form'].initial['filiere'], self.info.pk)
        r = self.client.post(reverse('matiere_creer'), {'nom': 'Pirate', 'filiere': self.gc.pk, 'session': self.session.pk,
                                                        'licence': self.licence.pk, 'semestre': s1.pk})
        self.assertIn('filiere', r.context['form'].errors)

    def test_description_repliee_dans_la_liste(self):
        s1, _ = self._semestres()
        self._matiere('Chimie', self.info, s1)
        self.connecter('scol')
        r = self.client.get(reverse('matieres_liste'))
        self.assertNotContains(r, 'id="tout-deplier"')                        # aucune description : pas de bouton
        m = self._matiere('Biologie', self.info, s1, description='Cours magistral\net travaux dirigés')
        r = self.client.get(reverse('matieres_liste'))
        self.assertContains(r, 'id="tout-deplier"')
        self.assertContains(r, f'id="description-{m.pk}"')
        self.assertContains(r, 'class="collapse description-matiere"')          # repliée par défaut
        self.assertContains(r, 'Cours magistral<br>et travaux dirigés')
        self.assertContains(r, '<div class="collapse" id="detail-credits">')        # crédits par semestre repliés
        self.assertContains(r, 'cr. (2)')

    def test_tri_des_matieres(self):
        s1, s3 = self._semestres()
        self._matiere('Chimie', self.info, s1, credit=2, code='C3')
        self._matiere('Algèbre', self.gc, s3, credit=6, code='A1')
        self._matiere('Biologie', self.info, s1, credit=4, code='B2')
        self.connecter('scol')
        noms = lambda **p: [m.nom for m in self.client.get(reverse('matieres_liste'), p).context['page']]  # noqa: E731
        self.assertEqual(noms(tri='nom'), ['Algèbre', 'Biologie', 'Chimie'])
        self.assertEqual(noms(tri='nom', ordre='desc'), ['Chimie', 'Biologie', 'Algèbre'])
        self.assertEqual(noms(tri='credit', ordre='desc'), ['Algèbre', 'Biologie', 'Chimie'])
        self.assertEqual(noms(tri='code'), ['Algèbre', 'Biologie', 'Chimie'])
        self.assertEqual(noms(tri='semestre', ordre='desc')[0], 'Algèbre')
        # Par défaut (ou tri inconnu) : programme, licence, semestre, nom
        self.assertEqual(noms(), ['Algèbre', 'Biologie', 'Chimie'])            # Génie civil avant Informatique
        self.assertEqual(noms(tri='pirate'), noms())
        # Le lien du titre actif inverse l'ordre, et le tri suit les filtres
        r = self.client.get(reverse('matieres_liste'), {'tri': 'credit', 'programme': self.info.pk})
        credit = next(c for c in r.context['colonnes'] if c['cle'] == 'credit')
        self.assertTrue(credit['actif'])
        self.assertIn('ordre=desc', credit['url'])
        self.assertIn(f'programme={self.info.pk}', credit['url'])
        self.assertEqual([m.nom for m in r.context['page']], ['Chimie', 'Biologie'])

    def test_matiere_formulaire_et_suppression_protegee(self):
        from .models import Matiere, Note
        s1, s3 = self._semestres()
        self.connecter('scol')
        base = {'nom': 'Bases de données', 'filiere': self.info.pk, 'session': self.session.pk,
                'licence': self.licence.pk, 'semestre': s1.pk, 'credit': ''}
        # Semestre d'une autre licence : refusé
        r = self.client.post(reverse('matiere_creer'), {**base, 'semestre': s3.pk})
        self.assertIn("SEMESTRE 3 n'appartient pas à Licence 1", r.context['form'].errors['semestre'][0])
        # Crédit vide : 6 ; cycle vide : celui du programme
        self.client.post(reverse('matiere_creer'), base)
        m = Matiere.objects.get(nom='Bases de données')
        self.assertEqual((m.credit, m.filiere, m.semestre, m.cycle), (6, self.info, s1, 'Licence'))
        # Cycle choisi : gardé
        self.client.post(reverse('matiere_creer'), {**base, 'nom': 'Réseaux', 'cycle': 'Ingénieur'})
        self.assertEqual(Matiere.objects.get(nom='Réseaux').cycle, 'Ingénieur')
        # Programme sans cycle et cycle vide : refusé
        Service.objects.filter(pk=self.gc.pk).update(cycle=None)
        r = self.client.post(reverse('matiere_creer'), {**base, 'nom': 'Topographie', 'filiere': self.gc.pk})
        self.assertIn("Précisez le cycle", r.context['form'].errors['cycle'][0])
        # Même nom (autre casse), même programme, même session : refusé
        r = self.client.post(reverse('matiere_creer'), {**base, 'nom': '  BASES  de donnees '})
        self.assertIn('existe déjà', r.context['form'].errors['nom'][0])
        # Modification
        self.client.post(reverse('matiere_modifier', args=[m.pk]), {**base, 'credit': 4, 'code': 'INF102'})
        m.refresh_from_db()
        self.assertEqual((m.credit, m.code), (4, 'INF102'))
        # Suppression refusée si la matière a des notes, possible sinon
        e = Etudiant.objects.get(matricule='M0')
        Note.objects.create(etudiant=e, matiere=m, controle=6, examen=7)
        r = self.client.post(reverse('matiere_supprimer', args=[m.pk]), follow=True)
        self.assertContains(r, 'Suppression impossible')
        self.assertTrue(Matiere.objects.filter(pk=m.pk).exists())
        Note.objects.all().delete()
        self.client.post(reverse('matiere_supprimer', args=[m.pk]))
        self.assertFalse(Matiere.objects.filter(pk=m.pk).exists())

    def test_import_matieres(self):
        from .import_matieres import ImportRefuse, importer_matieres
        from .models import Matiere
        from .roles import get_profil
        s1, s3 = self._semestres()
        self._matiere('Algorithmique', self.info, s1, credit=4, code='ALG')
        entetes = ['matiere', 'programme', 'licence', 'semestre', 'session', 'code', 'credits', 'description']
        fichier = lambda *lignes: fichier_xlsx([entetes, *lignes])           # noqa: E731
        lignes = [
            ['ALGORITHMIQUE', 'INFO', 'Licence 1', 'S1', '2025 - 2026', None, 5, None],            # 2 mise à jour
            ['Réseaux', 'Informatique', 'L1', 'SEMESTRE 1', '2025-2026', 'RES1', None, 'Intro'],   # 3 crédit vide → 6
            [None] * 8,                                                                            # 4 vide
            ['Béton armé', 'Génie civil', 'Licence 2', 'S3', '2025-2026', None, '3', None],       # 5
            ['Réseaux', 'Informatique', 'Licence 1', 'S1', '2025-2026', None, 6, None],            # 6 doublon fichier
            ['Analyse', 'Informatique', 'Licence 1', 'S3', '2025-2026', None, None, None],         # 7 S3 ≠ L1
            ['Physique', 'MEDECINE', 'Licence 1', 'S1', '2025-2026', None, None, None],            # 8 programme inconnu
            ['Chimie', 'Informatique', 'Licence 1', 'S1', '2025-2026', None, 'beaucoup', None],    # 9 crédit invalide
        ]
        self.connecter('scol')
        profil_scol = get_profil(self.client.get(reverse('matieres_liste')).wsgi_request)
        r = importer_matieres(fichier(*lignes), profil_scol, simulation=True)
        self.assertEqual((r.lignes, r.creees, r.mises_a_jour, r.nb_erreurs), (7, 2, 1, 4))
        self.assertEqual([n for n, _ in r.erreurs], [6, 7, 8, 9])
        messages = ' '.join(m for _, m in r.erreurs)
        for attendu in ('doublon dans le fichier', "SEMESTRE 3 n'appartient pas à Licence 1", 'programme inconnu',
                        'crédit invalide'):
            self.assertIn(attendu, messages)
        self.assertFalse(Matiere.objects.filter(nom='Réseaux').exists())                  # simulation

        r = importer_matieres(fichier(*lignes), profil_scol)
        algo = Matiere.objects.get(nom='Algorithmique')
        self.assertEqual((algo.credit, algo.code), (5, 'ALG'))          # mis à jour ; code vide : conservé
        reseaux = Matiere.objects.get(nom='Réseaux')
        self.assertEqual((reseaux.credit, reseaux.code, reseaux.description, reseaux.cycle), (6, 'RES1', 'Intro', 'Licence'))
        self.assertEqual(Matiere.objects.get(nom='Béton armé').filiere, self.gc)
        self.assertEqual(algo.cycle, 'Licence')          # matière existante sans cycle : celui du programme

        # Colonne « cycle » : prise telle quelle ; programme sans cycle et colonne vide : ligne rejetée
        Service.objects.filter(pk=self.gc.pk).update(cycle=None)
        avec_cycle = fichier_xlsx([entetes + ['cycle'],
                                   ['Mécanique', 'GENIE CIVIL', 'Licence 1', 'S1', '2025-2026', None, None, None, 'ING'],
                                   ['Hydraulique', 'GENIE CIVIL', 'Licence 1', 'S1', '2025-2026', None, None, None, None]])
        r = importer_matieres(avec_cycle, profil_scol)
        self.assertEqual(Matiere.objects.get(nom='Mécanique').cycle, 'Ingénieur')
        self.assertEqual(r.erreurs, [(3, "cycle manquant : le programme Génie civil n'en a pas, "
                                         "précisez-le dans la colonne « cycle »")])
        Service.objects.filter(pk=self.gc.pk).update(cycle='Licence')
        # Réimport : rien ne change
        r = importer_matieres(fichier(*lignes), profil_scol)
        self.assertEqual((r.creees, r.mises_a_jour, r.inchangees), (0, 0, 3))

        # Chef de programme : une ligne d'un autre programme → tout l'import est refusé
        self.connecter('info')
        profil_info = get_profil(self.client.get(reverse('matieres_liste')).wsgi_request)
        mix = fichier(['Compilation', 'Informatique', 'Licence 1', 'S1', '2025-2026', None, None, None],
                      ['Topographie', 'GENIE CIVIL', 'Licence 1', 'S1', '2025-2026', None, None, None])
        with self.assertRaisesMessage(ImportRefuse, 'ligne 3 (GENIE CIVIL)'):
            importer_matieres(mix, profil_info)
        self.assertFalse(Matiere.objects.filter(nom__in=['Compilation', 'Topographie']).exists())
        r = self.client.post(reverse('matieres_importer'), {'fichier': fichier(
            ['Compilation', 'Informatique', 'Licence 1', 'S1', '2025-2026', None, None, None],
            ['Topographie', 'MEDECINE', 'Licence 1', 'S1', '2025-2026', None, None, None])})
        self.assertContains(r, 'Import refusé')
        self.assertFalse(Matiere.objects.filter(nom='Compilation').exists())
        # Son programme uniquement : accepté
        r = importer_matieres(fichier(['Compilation', 'INFO', 'Licence 1', 'S1', '2025-2026', None, None, None]),
                              profil_info)
        self.assertEqual(r.creees, 1)
        self.assertEqual(self.client.get(reverse('matieres_modele')).status_code, 200)

    # ---------------------------------------------------------------- notes

    def test_moyenne_d_une_matiere_et_mentions(self):
        from .notes import mention, moyenne_des_notes
        self.assertEqual(moyenne_des_notes(8, 6, 7), 6.9)          # 25 % + 35 % + 40 %
        self.assertEqual(moyenne_des_notes(8, None, 7), 7.4)       # sans TP : 40 % + 60 %
        self.assertEqual(moyenne_des_notes(None, None, 6), 6)      # une seule note : elle-même
        self.assertEqual(moyenne_des_notes(None, 6, 7), 6.53)      # poids ramenés aux notes données
        self.assertIsNone(moyenne_des_notes(None, None, None))
        for moyenne, attendu in ((4.99, ''), (5, 'Passable'), (6, 'Assez bien'), (7, 'Bien'),
                                 (7.99, 'Bien'), (8, 'Très bien'), (10, 'Très bien'), (None, '')):
            self.assertEqual(mention(moyenne), attendu, moyenne)

    def test_notes_tronquees_jamais_arrondies(self):
        """Saisie, import et relevés : deux décimales TRONQUÉES (6,257 → 6,25), jamais arrondies."""
        from .notes import en_note
        from .pdf_releves import fmt
        self.assertEqual([en_note(v) for v in ('6,257', '6.999', '5,005', '7,2', 'ABS')], [6.25, 6.99, 5.0, 7.2, 0.0])
        self.assertEqual([fmt(v) for v in (6.257, 6.999, 7.19999999, 10, None)], ['6,25', '6,99', '7,20', '10,00', '—'])

    def test_note_moyenne_deposee_rattrapage_absent(self):
        from django.core.exceptions import ValidationError
        from .models import Note
        s1, _ = self._semestres()
        m = self._matiere('Électronique', self.info, s1)
        e0, e1, e2 = (Etudiant.objects.get(matricule=f'M{k}') for k in range(3))
        # Moyenne déposée par le programme : gardée telle quelle, même si des notes sont données
        n = Note.objects.create(etudiant=e0, matiere=m, controle=2, examen=2, moyenne=6.4, moyenne_deposee=True)
        self.assertEqual((n.moyenne, n.moyenne_finale, n.validee, n.mention), (6.4, 6.4, True, 'Assez bien'))
        # Rattrapage meilleur : retenu, mention « Session »
        n = Note.objects.create(etudiant=e1, matiere=m, controle=4, examen=4, rattrapage=7)
        self.assertEqual((n.moyenne, n.moyenne_finale, n.par_session, n.validee), (4.0, 7, True, True))
        # Rattrapage moins bon : on garde la moyenne
        n.rattrapage = 3
        n.save()
        self.assertEqual((n.moyenne_finale, n.par_session), (4.0, False))
        # Absent = 0 : matière non validée
        n = Note.objects.create(etudiant=e2, matiere=m, examen=0)
        self.assertEqual((n.moyenne_finale, n.validee, n.mention), (0, False, ''))
        # Barème sur 10
        with self.assertRaises(ValidationError):
            Note(etudiant=e2, matiere=m, examen=11).full_clean()

    def test_resultats_semestre_annee_cycle_et_soutenance(self):
        from .models import Deliberation, Note, Resultat, Semestre, Soutenance
        from .notes import recalculer
        s1, _ = self._semestres()
        s2 = Semestre.objects.create(nom='SEMESTRE 2', abrege='S2', licence=self.licence)
        Inscription.objects.filter(filiere=self.info).update(cycle='Licence')
        e0, e1 = Etudiant.objects.get(matricule='M0'), Etudiant.objects.get(matricule='M1')
        m1 = self._matiere('Algorithmique', self.info, s1, credit=6)
        m2 = self._matiere('Anglais', self.info, s1, credit=4)
        m3 = self._matiere('Réseaux', self.info, s2, credit=6)
        Note.objects.create(etudiant=e0, matiere=m1, moyenne=8, moyenne_deposee=True)
        Note.objects.create(etudiant=e0, matiere=m2, controle=4, examen=4)
        Note.objects.create(etudiant=e1, matiere=m1, moyenne=9, moyenne_deposee=True)
        Note.objects.create(etudiant=e1, matiere=m2, moyenne=9, moyenne_deposee=True)
        recalculer([e0.pk, e1.pk])
        res = lambda etu, niveau, **f: Resultat.objects.get(etudiant=etu, niveau=niveau, **f)   # noqa: E731

        # Semestre 1 : pondéré par les crédits (8×6 + 4×4) / 10 = 6,4 ; seule la matière ≥ 5 est validée
        s1_e0 = res(e0, Resultat.SEMESTRE, semestre=s1)
        self.assertEqual((s1_e0.moyenne, s1_e0.credits_valides, s1_e0.credits_total, s1_e0.mention, s1_e0.rang),
                         (6.4, 6, 10, 'Assez bien', 2))
        self.assertEqual(res(e1, Resultat.SEMESTRE, semestre=s1).rang, 1)
        # Semestre 2 sans aucune note : pas de moyenne, note manquante signalée
        s2_e0 = res(e0, Resultat.SEMESTRE, semestre=s2)
        self.assertEqual((s2_e0.moyenne, s2_e0.notes_manquantes), (None, 1))
        # Année : la matière sans note compte 0 → (48 + 16 + 0) / 16 = 4,0
        a = res(e0, Resultat.ANNEE)
        self.assertEqual((a.moyenne, a.credits_valides, a.credits_total, a.notes_manquantes, a.mention), (4.0, 6, 16, 1, ''))

        # Note ajoutée en S2, soutenance : année 6,25 ; générale = 6,25 × 0,87 + 9 × 0,13 = 6,61
        Note.objects.create(etudiant=e0, matiere=m3, moyenne=6, moyenne_deposee=True)
        Soutenance.objects.create(etudiant=e0, cycle='Licence', filiere=self.info, session=self.session,
                                  promotion=self.p18, note=9)
        recalculer([e0.pk])
        self.assertEqual(res(e0, Resultat.ANNEE).moyenne, 6.25)
        c = res(e0, Resultat.CYCLE)
        # 6,25 × 0,87 + 9 × 0,13 = 6,6075 → tronqué à 6,60 (jamais arrondi au-dessus)
        self.assertEqual((c.moyenne, c.soutenance, c.moyenne_generale, c.mention, c.session), (6.25, 9, 6.6, 'Assez bien', self.session))

        # Licence redoublée : le cycle prend la dernière année de chaque licence
        Inscription.objects.create(etudiant=e0, session=self.s2425, licence=self.licence, filiere=self.info, cycle='Licence')
        from .models import Matiere
        ancienne = Matiere.objects.create(nom='Algorithmique', filiere=self.info, semestre=s1, licence=self.licence,
                                          session=self.s2425, credit=6)
        Note.objects.create(etudiant=e0, matiere=ancienne, moyenne=2, moyenne_deposee=True)
        recalculer([e0.pk])
        self.assertEqual(res(e0, Resultat.ANNEE, session=self.s2425).moyenne, 2.0)
        self.assertEqual(res(e0, Resultat.CYCLE).moyenne, 6.25)

        # Délibération : la note est verrouillée
        n = Note.objects.get(etudiant=e0, matiere=m1)
        self.assertFalse(n.verrouillee)
        Deliberation.objects.create(filiere=self.info, session=self.session, semestre=s1)
        self.assertTrue(n.verrouillee)
        self.assertFalse(Note.objects.get(etudiant=e0, matiere=m3).verrouillee)       # S2 pas délibéré

    # ---------------------------------------------------------------- écrans des notes

    def _preparer_notes(self):
        s1, _ = self._semestres()
        Inscription.objects.update(cycle='Licence')
        algo = self._matiere('Algorithmique', self.info, s1, credit=6, code='ALG1')
        anglais = self._matiere('Anglais', self.info, s1, credit=4)
        return s1, algo, anglais

    def test_saisie_des_notes_par_matiere(self):
        import shutil, tempfile
        from django.test import override_settings
        from .models import Deliberation, Note, NoteHistorique, Resultat
        s1, algo, _ = self._preparer_notes()
        e0, e1 = Etudiant.objects.get(matricule='M0'), Etudiant.objects.get(matricule='M1')
        url = reverse('notes_saisie', args=[algo.pk])

        self.connecter('dep')                                     # autre service : pas d'accès aux notes
        self.assertEqual(self.client.get(url).status_code, 403)
        self.connecter('info')                                    # chef de programme : ses matières
        r = self.client.get(url)
        self.assertContains(r, 'N0')
        self.assertContains(r, 'N1')
        self.assertNotContains(r, '>N2</b>')                      # inscrit en Génie civil (pas « N2 » seul :
                                                                  # le jeton CSRF aléatoire peut le contenir)
        gc = self._matiere('Topographie', self.gc, s1)
        self.assertEqual(self.client.get(reverse('notes_saisie', args=[gc.pk])).status_code, 404)

        # Une note invalide : rien n'est enregistré
        r = self.client.post(url, {f'controle_{e0.pk}': '6', f'examen_{e0.pk}': 'x'})
        self.assertContains(r, 'note invalide')
        self.assertFalse(Note.objects.exists())
        # Sans TP : 40 % + 60 % → 7,2 ; moyenne déposée 6,5 + rattrapage 7 → 7 « Session »
        r = self.client.post(url, {f'controle_{e0.pk}': '6', f'examen_{e0.pk}': '8',
                                   f'moyenne_{e1.pk}': '6,5', f'rattrapage_{e1.pk}': '7'})
        self.assertRedirects(r, url)
        n0, n1 = Note.objects.get(etudiant=e0, matiere=algo), Note.objects.get(etudiant=e1, matiere=algo)
        self.assertEqual((n0.moyenne, n0.moyenne_finale, n0.moyenne_deposee), (7.2, 7.2, False))
        self.assertEqual((n1.moyenne, n1.moyenne_finale, n1.par_session), (6.5, 7, True))
        self.assertTrue(Resultat.objects.filter(etudiant=e0, niveau=Resultat.SEMESTRE, semestre=s1).exists())
        # Modifier une note déjà saisie : motif ET justificatif PDF obligatoires, puis tracée
        dossier = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, dossier, True)
        modif = {f'controle_{e0.pk}': '6', f'examen_{e0.pk}': '9', f'moyenne_{e1.pk}': '6,5', f'rattrapage_{e1.pk}': '7'}
        self.assertContains(self.client.post(url, modif), 'indiquez le motif et joignez le justificatif PDF')
        r = self.client.post(url, {**modif, 'motif': 'Erreur de report',
                                   'justificatif': SimpleUploadedFile('pv.txt', b'pas un pdf')})
        self.assertContains(r, 'doit être un fichier PDF')
        self.assertFalse(NoteHistorique.objects.exists())
        with override_settings(BASE_DIR=pathlib_path(dossier)):
            self.client.post(url, {**modif, 'motif': 'Erreur de report',
                                   'justificatif': SimpleUploadedFile('pv.pdf', b'%PDF-1.4 pv')})
        h = NoteHistorique.objects.get()
        self.assertEqual((h.avant['examen'], h.apres['examen'], h.apres_deliberation), (8, 9, False))
        self.assertEqual((h.utilisateur.username, h.motif), ('info', 'Erreur de report'))
        self.assertTrue((pathlib_path(dossier) / 'prive' / h.justificatif.name).exists())

        # Délibéré, année en cours : toujours modifiable par le chef de programme (motif + PDF)
        Deliberation.objects.create(filiere=self.info, session=self.session, semestre=s1)
        self.assertNotContains(self.client.get(url), 'disabled')
        with override_settings(BASE_DIR=pathlib_path(dossier)):
            self.client.post(url, {**modif, f'examen_{e0.pk}': '10', 'motif': 'PV du jury',
                                   'justificatif': SimpleUploadedFile('pv.pdf', b'%PDF-1.4 pv')})
        h = NoteHistorique.objects.filter(apres_deliberation=True).get()
        self.assertEqual((h.motif, h.apres['examen']), ('PV du jury', 10))

    def test_admin_niveaux_des_promotions(self):
        from django.contrib.auth.models import User
        User.objects.create_superuser('admin', password='x')
        self.client.login(username='admin', password='x')
        r = self.client.get(reverse('admin:pfe_promotion_licence_session_changelist'))
        self.assertContains(r, 'P18 · Licence 2 · 2024-2025')
        self.assertEqual(self.client.get(reverse('admin:pfe_promotion_licence_session_add')).status_code, 200)
        # Ajout direct : promotion, licence, session (un doublon est refusé)
        donnees = {'promotion': self.p20.pk, 'licence': self.l2.pk, 'session': self.s2425.pk}
        self.client.post(reverse('admin:pfe_promotion_licence_session_add'), donnees)
        self.assertTrue(Promotion_Licence_Session.objects.filter(promotion=self.p20, licence=self.l2).exists())
        r = self.client.post(reverse('admin:pfe_promotion_licence_session_add'), donnees)
        self.assertContains(r, 'existe déjà')
        # Depuis la promotion : ses niveaux par session, modifiables sur la même page
        r = self.client.get(reverse('admin:pfe_promotion_change', args=[self.p18.pk]))
        self.assertContains(r, 'Niveaux par session')

    def test_releves_delivres_nombre_affiche_regle_dans_l_admin(self):
        from django.contrib.auth.models import User
        from .models import ReglageAffichage, ReleveNote
        e = Etudiant.objects.get(matricule='M0')
        rel = [ReleveNote.objects.create(type='E', etudiant=e, filiere=self.info) for _ in range(3)]
        self.connecter('scol')
        # Sans réglage : 50 par défaut, tout est affiché
        r = self.client.get(reverse('releves'))
        self.assertEqual(len(r.context['page']), 3)
        # Réglage à 2 dans l'admin : les 2 derniers seulement, pour tout le monde
        User.objects.create_superuser('admin', password='x')
        self.client.login(username='admin', password='x')
        url_admin = reverse('admin:pfe_reglageaffichage_add')
        self.client.post(url_admin, {'nb_releves': '2'})
        self.assertEqual(ReglageAffichage.actuel().nb_releves, 2)
        self.assertEqual(self.client.get(url_admin).status_code, 403)              # une seule ligne
        reglage = ReglageAffichage.objects.get()
        self.assertEqual(self.client.get(reverse('admin:pfe_reglageaffichage_delete', args=[reglage.pk])).status_code,
                         403)
        self.connecter('scol')
        r = self.client.get(reverse('releves'))
        self.assertEqual([x.pk for x in r.context['page']], [rel[2].pk, rel[1].pk])
        self.assertNotContains(r, 'sur 3')
        # La recherche trouve aussi celui qui n'est pas affiché
        r = self.client.get(reverse('releves'), {'r': rel[0].num_ref})
        self.assertEqual([x.pk for x in r.context['page']], [rel[0].pk])
        # Un relevé délivré après le réglage s'ajoute : 2 + 1 = 3 (le plus ancien reste masqué)
        nouveau = ReleveNote.objects.create(type='E', etudiant=e, filiere=self.gc)
        r = self.client.get(reverse('releves'))
        self.assertEqual([x.pk for x in r.context['page']], [nouveau.pk, rel[2].pk, rel[1].pk])
        self.assertNotContains(r, '3 relevés')                           # nombre masqué pour l'instant
        # Réenregistrer le même nombre ne change rien ; un nouveau nombre fait repartir le décompte
        reglage.save()
        self.assertEqual(len(self.client.get(reverse('releves')).context['page']), 3)
        reglage.nb_releves = 1
        reglage.save()
        self.assertEqual([x.pk for x in self.client.get(reverse('releves')).context['page']], [nouveau.pk])
        # Hors de l'espace admin : la scolarité n'a pas accès au réglage
        self.assertNotEqual(self.client.get(reverse('admin:pfe_reglageaffichage_change', args=[reglage.pk])).status_code,
                            200)

    def test_super_utilisateur_modifie_les_notes_sans_justificatif(self):
        from django.contrib.auth.models import User
        from .forms import ImportNotesForm
        from .models import Note, NoteHistorique
        from .notes import MOTIF_SUPER_UTILISATEUR
        from .roles import Profil, SCOLARITE
        s1, algo, _ = self._preparer_notes()
        e0 = Etudiant.objects.get(matricule='M0')
        Note.objects.create(etudiant=e0, matiere=algo, examen=6)
        url = reverse('notes_saisie', args=[algo.pk])
        # Scolarité : motif et PDF toujours obligatoires
        self.connecter('scol')
        self.assertContains(self.client.post(url, {f'examen_{e0.pk}': '8'}), 'indiquez le motif')
        # Super utilisateur : ni motif ni PDF, modification tracée avec la mention
        User.objects.create_superuser('admin', password='x')
        self.client.login(username='admin', password='x')
        self.assertContains(self.client.get(url), 'sont <strong>facultatifs</strong> pour vous')
        self.assertRedirects(self.client.post(url, {f'examen_{e0.pk}': '8'}), url)
        h = NoteHistorique.objects.get()
        self.assertEqual((h.avant['examen'], h.apres['examen'], h.motif, h.justificatif.name),
                         (6, 8, MOTIF_SUPER_UTILISATEUR, ''))
        # Un fichier joint doit quand même être un vrai PDF
        r = self.client.post(url, {f'examen_{e0.pk}': '9', 'justificatif': SimpleUploadedFile('pv.txt', b'texte')})
        self.assertContains(r, 'doit être un fichier PDF')
        # Import : PDF de l'enseignant facultatif pour lui seul, même pour modifier une note enregistrée
        from .import_notes import importer_notes
        self.assertFalse(ImportNotesForm(profil=Profil(SCOLARITE, superuser=True)).fields['pdf_enseignant'].required)
        modif = fichier_xlsx([['matricule', 'matiere', 'session', 'controle', 'tp', 'examen', 'moyenne', 'rattrapage'],
                              ['M0', 'Algorithmique', '2025-2026', None, None, 7, None, None]])
        r = importer_notes(modif, Profil(SCOLARITE, superuser=True), User.objects.get(username='admin'))
        self.assertEqual((r.mises_a_jour, r.erreurs), (1, []))
        self.assertEqual(Note.objects.get(etudiant=e0, matiere=algo).examen, 7)

    def test_saisie_cascade_et_un_seul_etudiant(self):
        from urllib.parse import quote
        from .models import Note
        s1, algo, anglais = self._preparer_notes()
        e0, e1 = Etudiant.objects.get(matricule='M0'), Etudiant.objects.get(matricule='M1')
        self.connecter('info')
        # Cascade : licence proposée une fois le programme choisi, matières une fois le semestre choisi
        base = reverse('notes_matieres')
        r = self.client.get(base, {'session': self.session.pk, 'programme': self.info.pk})
        self.assertEqual([l.pk for l in r.context['c_licences']], [self.licence.pk])
        self.assertEqual(r.context['c_matieres'], [])
        r = self.client.get(base, {'session': self.session.pk, 'programme': self.info.pk, 'licence': self.licence.pk,
                                   'semestre': s1.pk})
        self.assertEqual([m.nom for m in r.context['c_matieres']], ['Algorithmique', 'Anglais'])
        self.assertContains(r, f'data-url="{reverse("notes_saisie", args=[algo.pk])}"')
        # Programme d'un autre service ou licence sans matière : choix vidés
        r = self.client.get(base, {'programme': self.gc.pk, 'licence': self.l2.pk})
        self.assertEqual((r.context['cascade']['programme'], r.context['cascade']['licence']), (self.info.pk, None))
        # Sans JavaScript : la matière choisie ouvre la saisie (seulement une matière visible)
        self.assertRedirects(self.client.get(base, {'matiere': algo.pk}), reverse('notes_saisie', args=[algo.pk]))
        gc = self._matiere('Topographie', self.gc, s1)
        self.assertEqual(self.client.get(base, {'matiere': gc.pk}).status_code, 200)
        # La grille montre la cascade préremplie
        r = self.client.get(reverse('notes_saisie', args=[algo.pk]))
        self.assertEqual(r.context['cascade']['matiere'], algo.pk)
        self.assertContains(r, 'id="cascade-matiere"')

        # Crayon de la liste des notes : seulement cet étudiant, puis retour à la liste
        Note.objects.create(etudiant=e0, matiere=algo, examen=6)
        Note.objects.create(etudiant=e1, matiere=algo, examen=7)
        liste = reverse('notes_liste') + '?session=' + str(self.session.pk) + '&page=1'
        r = self.client.get(liste)
        lien = f'{reverse("notes_saisie", args=[algo.pk])}?etudiant={e0.pk}&amp;retour={quote(liste, safe="")}'
        self.assertContains(r, lien)
        url = f'{reverse("notes_saisie", args=[algo.pk])}?etudiant={e0.pk}&retour={quote(liste, safe="")}'
        r = self.client.get(url)
        self.assertEqual([l['etudiant'].pk for l in r.context['lignes']], [e0.pk])
        self.assertContains(r, 'Modification de la note de')
        self.assertNotContains(r, f'examen_{e1.pk}')
        # Saisie d'une nouvelle note (Anglais, pas de justificatif) : retour à la liste
        r = self.client.post(f'{reverse("notes_saisie", args=[anglais.pk])}?etudiant={e1.pk}&retour={quote(liste, safe="")}',
                             {f'examen_{e1.pk}': '8', f'examen_{e0.pk}': '2'})
        self.assertRedirects(r, liste, fetch_redirect_response=False)
        self.assertEqual(list(Note.objects.filter(matiere=anglais).values_list('etudiant_id', flat=True)), [e1.pk])
        # Adresse de retour extérieure ignorée : on reste sur la saisie
        r = self.client.post(f'{reverse("notes_saisie", args=[anglais.pk])}?etudiant={e1.pk}&retour=https://pirate.example/',
                             {f'examen_{e1.pk}': '8'})
        self.assertRedirects(r, f'{reverse("notes_saisie", args=[anglais.pk])}?etudiant={e1.pk}&retour=https://pirate.example/',
                             fetch_redirect_response=False)

    def test_notes_annee_ecoulee_reservees_aux_chefs(self):
        from .models import Matiere, Note, Personnel
        s1, _ = self._semestres()
        e0 = Etudiant.objects.get(matricule='M0')
        Inscription.objects.create(etudiant=e0, session=self.s2425, licence=self.l2, filiere=self.info)
        ancienne = Matiere.objects.create(nom='Réseaux', filiere=self.info, semestre=s1, licence=self.l2,
                                          session=self.s2425, credit=6)
        Note.objects.create(etudiant=e0, matiere=ancienne, moyenne=6, moyenne_deposee=True)
        url = reverse('notes_saisie', args=[ancienne.pk])
        post = {f'moyenne_{e0.pk}': '9', 'motif': 'x', 'justificatif': SimpleUploadedFile('pv.pdf', b'%PDF-1.4 x')}
        # Ni le chef de programme ni un agent de la scolarité (non chefs) ne la modifient
        for login in ('info', 'scol'):
            self.connecter(login)
            r = self.client.get(url)
            self.assertContains(r, 'Année écoulée')
            self.assertContains(r, 'disabled')
            self.assertEqual(self.client.post(url, post).status_code, 403)
        # Chef de programme coché « Chef » mais Informatique n'est pas un département : toujours refusé
        Personnel.objects.filter(utilisateur__username='info').update(chef=True)
        self.connecter('info')
        self.assertEqual(self.client.post(url, post).status_code, 403)
        # Informatique devient son propre département : son chef peut modifier
        Service.objects.filter(pk=self.info.pk).update(departement=self.info)
        self.assertContains(self.client.get(url), 'En tant que chef')
        # Chef de la scolarité : peut modifier
        Personnel.objects.filter(utilisateur__username='scol').update(chef=True)
        self.connecter('scol')
        self.assertNotContains(self.client.get(url), 'disabled')

    def test_pieces_du_dossier(self):
        """Scan (photos) ou PDF → un PDF compressé dans le stockage privé ; scolarité seulement."""
        import shutil, tempfile
        from django.test import override_settings
        from PIL import Image
        from .models import PieceDossier
        dossier = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, dossier, True)
        e0 = Etudiant.objects.get(matricule='M0')
        insc = Inscription.objects.get(etudiant=e0)
        url = reverse('piece_ajouter', args=[e0.pk])

        def photo(nom='page.jpg', taille=(2400, 3200)):
            b = io.BytesIO()
            Image.new('RGB', taille, (240, 240, 235)).save(b, 'JPEG', quality=95)
            return SimpleUploadedFile(nom, b.getvalue(), 'image/jpeg')

        self.connecter('info')                                         # chef de programme : pas d'accès
        self.assertEqual(self.client.post(url, {'type': 'acte_naissance'}).status_code, 403)
        self.assertNotContains(self.client.get(reverse('etudiant_fiche', args=[insc.pk])), 'Dossier de l')
        self.connecter('scol')
        with override_settings(BASE_DIR=pathlib_path(dossier)):
            # Deux photos de pages → un PDF de 2 pages, petit
            r = self.client.post(url, {'type': 'acte_naissance', 'pages': [photo('p1.jpg'), photo('p2.jpg')]}, follow=True)
            self.assertContains(r, 'Acte de naissance ajouté au dossier (2 pages')
            p = PieceDossier.objects.get()
            self.assertEqual((p.pages, p.par_etudiant, p.utilisateur.username), (2, False, 'scol'))
            chemin = pathlib_path(dossier) / 'prive' / p.fichier.name
            self.assertTrue(chemin.read_bytes().startswith(b'%PDF-'))
            self.assertLess(p.taille, 300 * 1024)
            # Un nouveau dépôt du même type remplace l'ancien (fichier supprimé)
            self.client.post(url, {'type': 'acte_naissance', 'pages': [SimpleUploadedFile('acte.pdf', b'%PDF-1.4 acte')]})
            self.assertEqual(PieceDossier.objects.count(), 1)
            self.assertFalse(chemin.exists())
            # Refusés : PDF trop lourd, PDF + photo, fichier qui n'est pas une image
            r = self.client.post(url, {'type': 'diplome_bac', 'pages': [SimpleUploadedFile('x.pdf', b'%PDF-' + b'0' * (3 * 1024 * 1024))]}, follow=True)
            self.assertContains(r, 'PDF trop lourd')
            r = self.client.post(url, {'type': 'diplome_bac', 'pages': [SimpleUploadedFile('x.pdf', b'%PDF-1.4'), photo()]}, follow=True)
            self.assertContains(r, "s&#x27;envoie seul")
            r = self.client.post(url, {'type': 'diplome_bac', 'pages': [SimpleUploadedFile('x.txt', b'texte')]}, follow=True)
            self.assertContains(r, 'seuls les PDF et les images')
            # Consultation (scolarité) et suppression
            p = PieceDossier.objects.get()
            self.assertContains(self.client.get(reverse('etudiant_fiche', args=[insc.pk])), 'Acte de naissance')
            r = self.client.get(reverse('piece_voir', args=[p.pk]))
            self.assertEqual((r['Content-Type'], b''.join(r.streaming_content)), ('application/pdf', b'%PDF-1.4 acte'))
            self.connecter('info')
            self.assertEqual(self.client.get(reverse('piece_voir', args=[p.pk])).status_code, 403)
            self.connecter('scol')
            self.client.post(reverse('piece_supprimer', args=[p.pk]))
            self.assertFalse(PieceDossier.objects.exists())

    def _espace(self):
        """Espace étudiant : cache local (essais limités) et dossier privé temporaires."""
        import shutil, tempfile
        from django.core.cache import cache
        from django.test import override_settings
        dossier = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, dossier, True)
        reglages = override_settings(BASE_DIR=pathlib_path(dossier), CACHES={
            'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache'}})
        reglages.enable()
        self.addCleanup(reglages.disable)
        cache.clear()
        return dossier

    def test_attestations_directeur_de_master(self):
        from .models import Attestation
        master = Service.objects.create(nom='MASTER EN ENERGIE', code='M03', programme=True, cycle='Master')
        m1 = Licence.objects.create(nom='MASTER 1')
        # Directeur du Master : chef rattaché au programme de Master
        u = User.objects.create_user('directeur', password='x')
        p = Personnel.objects.create(utilisateur=u, prenom='Awa', nom='Bah', telephone='901', email='dir@t.local',
                                     fonc='Directeur du Master', chef=True)
        p.service.add(master)
        etu = Etudiant.objects.create(matricule='MM1', prenom='Ali', nom='SOW', email='mm1@t.local')
        Inscription.objects.create(etudiant=etu, session=self.session, licence=m1, filiere=master, cycle='Master')
        url = reverse('attestations_type', args=['inscription'])
        # Chef de département (programme qui est son propre département) : toujours refusé
        Service.objects.filter(pk=self.info.pk).update(departement=self.info)
        Personnel.objects.filter(utilisateur__username='info').update(chef=True)
        self.connecter('info')
        self.assertEqual(self.client.get(url).status_code, 403)
        self.assertNotContains(self.client.get(reverse('tableau_de_bord')), 'bi-file-earmark-text')
        # Scolarité : une attestation d'un étudiant de Licence
        self.connecter('scol')
        self.client.post(url, {'matricule': 'M0', 'licence': self.licence.pk})
        autre = Attestation.objects.get(etudiant__matricule='M0')
        # Directeur : le menu, ses étudiants seulement
        self.client.login(username='directeur', password='x')
        self.assertContains(self.client.get(reverse('tableau_de_bord')), 'bi-file-earmark-text')
        r = self.client.get(url, {'etudiant': 'SOW'})
        self.assertEqual([c.etudiant.matricule for c in r.context['candidats']], ['MM1'])
        self.assertEqual(self.client.get(url, {'etudiant': 'N0'}).context['candidats'], [])
        r = self.client.post(url, {'matricule': 'M1', 'licence': self.licence.pk})
        self.assertContains(r, "pas inscrit dans un Master que vous dirigez")
        self.assertFalse(Attestation.objects.filter(etudiant__matricule='M1').exists())
        r = self.client.post(url, {'matricule': 'MM1', 'licence': m1.pk})
        a = Attestation.objects.get(etudiant=etu)
        self.assertEqual(a.filiere, master)
        r = self.client.get(url)
        self.assertEqual([x.pk for x in r.context['page']], [a.pk])                   # pas celle de la scolarité
        self.assertEqual(self.client.get(reverse('attestation_pdf', args=[a.pk]))['Content-Type'], 'application/pdf')
        self.assertEqual(self.client.get(reverse('attestation_pdf', args=[autre.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse('attestation_supprimer', args=[autre.pk])).status_code, 404)
        self.assertEqual(self.client.get(url, {'modifier': autre.pk}).status_code, 404)
        self.assertEqual(self.client.get(reverse('releves')).status_code, 403)      # relevés : scolarité seule
        self.client.post(reverse('attestation_supprimer', args=[a.pk]))
        self.assertFalse(Attestation.objects.filter(pk=a.pk).exists())

    def test_commande_demo_master(self):
        import re
        from io import StringIO
        from django.contrib.auth import authenticate
        from django.core.management import call_command
        sortie = StringIO()
        call_command('gped_demo_master', '--email-etudiant', 'Essai.Master@Exemple.com', stdout=sortie)
        mdp = re.search(r'notez-le\) : (\S+)', sortie.getvalue()).group(1)
        self.assertIsNotNone(authenticate(username='demo.master', password=mdp))
        programme = Service.objects.get(nom='MASTER EN ENERGIE')
        self.assertEqual((programme.programme, programme.cycle), (True, 'Master'))
        perso = Personnel.objects.get(utilisateur__username='demo.master')
        self.assertEqual((perso.chef, perso.nom_signature, list(perso.service.all())),
                         (True, 'Dr Démo MASTER ÉNERGIE', [programme]))
        insc = Inscription.objects.get(etudiant__matricule='26M03DEMO01')
        self.assertEqual((insc.etudiant.email, insc.licence.nom, insc.cycle, insc.etudiant.preinscrit),
                         ('essai.master@exemple.com', 'MASTER 1', 'Master', True))
        # Relancée : rien n'est dupliqué, le mot de passe ne change pas
        sortie = StringIO()
        call_command('gped_demo_master', stdout=sortie)
        self.assertIn('Mot de passe inchangé', sortie.getvalue())
        self.assertIsNotNone(authenticate(username='demo.master', password=mdp))
        self.assertEqual(Etudiant.objects.filter(matricule='26M03DEMO01').count(), 1)
        # Suppression : le compte et l'étudiante partent, le programme reste
        call_command('gped_demo_master', '--supprimer', stdout=StringIO())
        self.assertFalse(User.objects.filter(username='demo.master').exists())
        self.assertFalse(Etudiant.objects.filter(matricule='26M03DEMO01').exists())
        self.assertTrue(Service.objects.filter(nom='MASTER EN ENERGIE').exists())

    def test_espace_master_lien_par_email(self):
        import re
        from django.core import mail
        from .models import DemandeEtudiant, PieceDossier
        self._espace()
        Session.objects.filter(pk=self.session.pk).update(inscriptions_ouvertes=True, reinscriptions_ouvertes=True)
        m1, m2 = Licence.objects.create(nom='MASTER 1'), Licence.objects.create(nom='MASTER 2')
        pm1, pm0 = Promotion.objects.create(nom='PM1'), Promotion.objects.create(nom='PM0')
        for promo, lic, ses in [(pm1, m1, self.session), (pm0, m1, self.s2425), (pm0, m2, self.session)]:
            Promotion_Licence_Session.objects.create(promotion=promo, licence=lic, session=ses)
        # Préinscrit en Master (inscription) et ancien étudiant de Master (réinscription)
        awa = Etudiant.objects.create(matricule='MA1', prenom='Awa', nom='BAH', email='awa.master@exemple.com',
                                      ine='INEMA1', pvbac='111', preinscrit=True)
        Inscription.objects.create(etudiant=awa, session=self.session, licence=m1, filiere=self.info, promotion=pm1,
                                   cycle='Master')
        ali = Etudiant.objects.create(matricule='MA2', prenom='Ali', nom='SOW', email='ali.master@exemple.com',
                                      ine='INEMA2')
        Inscription.objects.create(etudiant=ali, session=self.s2425, licence=m1, filiere=self.info, promotion=pm0,
                                   cycle='Master')

        def lien(adresse):
            mail.outbox.clear()
            r = self.client.post(reverse('espace_accueil'), {'mode': 'master', 'm-email': adresse}, follow=True)
            self.assertContains(r, 'Si cette adresse est celle')                   # même réponse dans tous les cas
            return mail.outbox
        # Adresse directe : l'espace sur l'onglet Master
        self.assertRedirects(self.client.get(reverse('espace_master')), reverse('espace_accueil') + '?onglet=master')
        self.assertContains(self.client.get(reverse('espace_accueil') + '?onglet=master'), 'Recevoir mon lien')
        # Adresse inconnue : rien n'est envoyé ; un Master ne passe pas par l'INE
        self.assertEqual(lien('personne@exemple.com'), [])
        r = self.client.post(reverse('espace_accueil'), {'mode': 'inscription', 'i-ine': 'INEMA1', 'i-pvbac': '111'})
        self.assertContains(r, 'utilisez l&#x27;onglet « Master »')
        # Adresse connue mais pas en Master : la raison part par e-mail, sans lien
        envoyes = lien('e0@t.local')
        self.assertEqual(len(envoyes), 1)
        self.assertIn("pas étudiant(e) en Master", envoyes[0].body)
        self.assertNotIn('/etudiant/lien/', envoyes[0].body)
        # Inscription en Master : lien personnel ; l'ouvrir (GET) ne le consomme pas
        envoyes = lien('AWA.master@exemple.com')
        self.assertEqual(envoyes[0].to, ['awa.master@exemple.com'])
        url = re.search(r'https?://[^/\s]+(/etudiant/lien/[^/\s]+/)', envoyes[0].body).group(1)
        self.assertContains(self.client.get(url), 'Ouvrir mon dossier')
        self.assertContains(self.client.get(url), 'Ouvrir mon dossier')
        self.assertRedirects(self.client.post(url), reverse('espace_dossier'))
        r = self.client.get(reverse('espace_dossier'))
        self.assertContains(r, 'Inscription en Master')
        self.assertContains(r, 'value="awa.master@exemple.com"')                  # adresse déjà remplie
        self.assertContains(r, 'Diplôme de licence')
        # Une seule fois : le lien a servi
        self.assertRedirects(self.client.post(url), reverse('espace_accueil') + '?onglet=master',
                             fetch_redirect_response=False)
        # Pièces : bac et acte obligatoires ; diplôme de licence et relevé de l'année dernière facultatifs
        base = {'telephone': '622 11 22 33', 'email': 'awa.master@exemple.com', 'adresse': 'Petel/Mamou',
                'tuteur': 'Mariam CAMARA', 'contact_de_tuteur': '628 99 88 77', **signer()}
        r = self.client.post(reverse('espace_dossier'), {**base, 'photo': photo_jpeg()})
        self.assertContains(r, 'Obligatoire : scannez ou envoyez cette pièce.', count=3)
        r = self.client.post(reverse('espace_dossier'), {
            **base, 'photo': photo_jpeg(), 'piece_acte_naissance': SimpleUploadedFile('a.pdf', b'%PDF-1.4 a'),
            'piece_diplome_bac': SimpleUploadedFile('d.pdf', b'%PDF-1.4 d'),
            'piece_releve_bac': SimpleUploadedFile('r.pdf', b'%PDF-1.4 r')})
        self.assertRedirects(r, reverse('espace_confirmation'))
        d = DemandeEtudiant.objects.get(etudiant=awa)
        self.assertEqual((d.type, d.statut), (DemandeEtudiant.INSCRIPTION, DemandeEtudiant.EN_ATTENTE))
        self.client.post(reverse('espace_quitter'))
        # Réinscription en Master : acte et pièces du bac déjà dans le dossier → pas redemandés
        for t in (PieceDossier.ACTE_NAISSANCE, PieceDossier.DIPLOME_BAC, PieceDossier.RELEVE_BAC):
            PieceDossier.objects.create(etudiant=ali, type=t, fichier=f'pieces/{t}.pdf', taille=10)
        url = re.search(r'(/etudiant/lien/[^/\s]+/)', lien('ali.master@exemple.com')[0].body).group(1)
        self.assertRedirects(self.client.post(url), reverse('espace_dossier'))
        r = self.client.get(reverse('espace_dossier'))
        self.assertContains(r, 'Réinscription en Master')
        self.assertContains(r, f'<option value="{m2.pk}" selected')                 # niveau prévu : MASTER 2
        r = self.client.post(reverse('espace_dossier'), {
            **base, 'email': 'ali.master@exemple.com', 'photo': photo_jpeg(), 'licence': m2.pk,
            'boursier': 'boursier', 'type_formation': 'initiale'})
        self.assertRedirects(r, reverse('espace_confirmation'))                   # relevé facultatif
        d = DemandeEtudiant.objects.get(etudiant=ali)
        self.assertEqual((d.type, d.licence, d.derniere_inscription.licence), (DemandeEtudiant.REINSCRIPTION, m2, m1))

    def test_espace_etudiant_inscription_en_ligne(self):
        from .models import DemandeEtudiant, PieceDossier
        self._espace()
        p = Etudiant.objects.create(matricule='2601000135', prenom='Awa', nom='BAH', email='awa@etudiant.gped',
                                    ine='INEP1', pvbac='123 456', preinscrit=True, telephone='600',
                                    boursier='boursier', type_formation='initiale')
        insc = Inscription.objects.create(etudiant=p, session=self.session, licence=self.licence, filiere=self.info,
                                          promotion=self.p20)
        accueil = reverse('espace_accueil')
        identifiants = {'mode': 'inscription', 'i-ine': 'inep1', 'i-pvbac': '123456'}
        # Fermé par défaut
        r = self.client.post(accueil, identifiants)
        self.assertContains(r, 'inscriptions en ligne</strong> sont fermées')
        Session.objects.filter(pk=self.session.pk).update(inscriptions_ouvertes=True)
        # Mauvais PV : refus sans dire quel champ est faux
        r = self.client.post(accueil, {**identifiants, 'i-pvbac': '999'})
        self.assertContains(r, 'Aucun dossier ne correspond')
        # Bon couple (espaces et casse ignorés) → dossier
        r = self.client.post(accueil, identifiants)
        self.assertRedirects(r, reverse('espace_dossier'))
        r = self.client.get(reverse('espace_dossier'))
        self.assertContains(r, 'Awa BAH')
        self.assertContains(r, 'class="bloc-photo')                             # même bloc photo que la scolarité
        self.assertContains(r, '<div class="lib">Département</div>')               # département et programme
        self.assertContains(r, '<div class="lib">Programme</div>')
        self.assertContains(r, 'Prendre ou charger une photo')
        self.assertNotContains(r, 'name="licence"')                       # pas de choix de niveau à l'inscription
        # Envoi : coordonnées + acte de naissance (PDF) → demande en attente, fiche inchangée
        # Contacts : vides au départ (même si la fiche en a), sans ville, adresse « Quartier/Mamou »
        r = self.client.get(reverse('espace_dossier'))
        self.assertNotContains(r, 'value="600"')                                   # téléphone de la fiche non repris
        self.assertNotContains(r, 'name="ville"')
        self.assertContains(r, 'placeholder="Quartier/Mamou"')
        # Tous obligatoires ; numéros d'au moins 8 chiffres ; contact du tuteur différent du sien
        r = self.client.post(reverse('espace_dossier'), {'telephone': '622 11 22 33'})
        self.assertContains(r, 'Ce champ est obligatoire', count=4)                # e-mail, adresse, tuteur, contact
        r = self.client.post(reverse('espace_dossier'), {
            'telephone': '6221', 'email': 'awa.bah@exemple.com', 'adresse': 'Petel/Mamou', 'tuteur': 'Mariam CAMARA',
            'contact_de_tuteur': '+224 622 11 22 33'})
        self.assertContains(r, 'Numéro incomplet')
        r = self.client.post(reverse('espace_dossier'), {
            'telephone': '622 11 22 33', 'email': 'awa.bah@exemple.com', 'adresse': 'Petel/Mamou',
            'tuteur': 'Mariam CAMARA', 'contact_de_tuteur': '+224 622 11 22 33'})
        self.assertContains(r, 'différent de votre propre numéro')
        complet = {'telephone': '622 11 22 33', 'email': 'awa.bah@exemple.com', 'adresse': '  Petel /Mamou ',
                   'tuteur': 'Mariam CAMARA', 'contact_de_tuteur': '628 99 88 77', **signer()}
        # Photo obligatoire (l'étudiant n'en a pas encore) : refusé sans photo
        r = self.client.post(reverse('espace_dossier'), complet)
        self.assertContains(r, 'elle est obligatoire')
        self.assertFalse(DemandeEtudiant.objects.exists())
        # Toutes les pièces sont obligatoires : il manque le diplôme et le relevé du bac
        r = self.client.post(reverse('espace_dossier'), {
            **complet, 'photo': photo_jpeg(), 'piece_acte_naissance': SimpleUploadedFile('acte.pdf', b'%PDF-1.4 acte')})
        self.assertContains(r, 'Obligatoire : scannez ou envoyez cette pièce.', count=2)
        self.assertFalse(DemandeEtudiant.objects.exists())
        r = self.client.post(reverse('espace_dossier'), {
            **complet, 'photo': photo_jpeg(), 'piece_acte_naissance': SimpleUploadedFile('acte.pdf', b'%PDF-1.4 acte'),
            'piece_diplome_bac': SimpleUploadedFile('diplome.pdf', b'%PDF-1.4 diplome'),
            'piece_releve_bac': SimpleUploadedFile('releve_bac.pdf', b'%PDF-1.4 releve')})
        self.assertRedirects(r, reverse('espace_confirmation'))
        # Retour sur la demande : l'étudiant retrouve ce qu'il a saisi
        self.assertContains(self.client.get(reverse('espace_dossier')), 'value="628 99 88 77"')
        d = DemandeEtudiant.objects.get()
        self.assertEqual((d.type, d.statut, d.contacts['telephone']), ('inscription', 'en_attente', '622 11 22 33'))
        self.assertEqual((d.boursier, d.type_formation), ('boursier', 'initiale'))      # cochés par défaut
        p.refresh_from_db()
        self.assertEqual((p.telephone, p.preinscrit), ('600', True))      # rien d'appliqué avant validation
        self.assertTrue(PieceDossier.objects.get(etudiant=p, type='acte_naissance').par_etudiant)
        # Photo : l'étudiant (sans compte) voit celle qu'il a envoyée, en attente de validation
        from PIL import Image
        img = io.BytesIO()
        Image.new('RGB', (900, 1200), (180, 150, 120)).save(img, 'JPEG')
        self.client.post(reverse('espace_dossier'), {'telephone': '622 11 22 33', 'email': 'awa.bah@exemple.com',
                                                     'adresse': 'Petel/Mamou', 'tuteur': 'Mariam CAMARA',
                                                     'contact_de_tuteur': '628 99 88 77', **signer(),
                                                     'photo': SimpleUploadedFile('moi.jpg', img.getvalue(), 'image/jpeg')})
        r = self.client.get(reverse('espace_dossier'))
        self.assertContains(r, reverse('espace_photo'))
        self.assertContains(r, 'en attente de validation par la scolarité')
        r = self.client.get(reverse('espace_photo'))
        self.assertEqual((r.status_code, r['Content-Type']), (200, 'image/jpeg'))
        self.assertTrue(b''.join(r.streaming_content).startswith(bytes([0xFF, 0xD8])))   # JPEG compressé
        # Hors de l'espace (autre navigateur) : photo inaccessible
        from django.test import Client
        self.assertEqual(Client().get(reverse('espace_photo')).status_code, 404)
        # « Inscrire » (page Préinscription) applique la demande
        self.connecter('scol')
        self.client.post(reverse('valider_preinscriptions'), {'ids': [insc.pk]})
        p.refresh_from_db()
        d.refresh_from_db()
        self.assertEqual((p.telephone, p.email, p.tuteur, p.preinscrit, d.statut),
                         ('622 11 22 33', 'awa.bah@exemple.com', 'Mariam CAMARA', False, 'validee'))
        # Inscription en ligne validée : la case « Préinscrit » reste modifiable (la validation est signalée)
        from django.forms.models import model_to_dict
        from .forms import EtudiantForm
        donnees = {k: v for k, v in model_to_dict(p, exclude=['photo', 'utilisateur']).items() if v not in (None, False)}
        f = EtudiantForm({**donnees, 'preinscrit': 'on'}, instance=p)
        self.assertFalse(f.fields['preinscrit'].disabled)
        self.assertTrue(f.is_valid(), f.errors)
        self.assertTrue(f.cleaned_data['preinscrit'])
        self.assertContains(self.client.get(reverse('inscription_modifier', args=[insc.pk])), 'Inscription en ligne validée le')

    def test_espace_etudiant_essais_limites_et_expiration(self):
        import time
        self._espace()
        Session.objects.filter(pk=self.session.pk).update(inscriptions_ouvertes=True)
        accueil = reverse('espace_accueil')
        for _ in range(5):
            self.client.post(accueil, {'mode': 'inscription', 'i-ine': 'X', 'i-pvbac': 'Y'})
        r = self.client.post(accueil, {'mode': 'inscription', 'i-ine': 'X', 'i-pvbac': 'Y'})
        self.assertContains(r, "Trop d&#x27;essais")
        # Sans identification (ou accès expiré) : retour à l'accueil
        self.assertRedirects(self.client.get(reverse('espace_dossier')), accueil)
        s = self.client.session
        s['espace_etudiant'] = {'etudiant': 1, 'type': 'inscription', 'inscription': 1, 'expire': time.time() - 1}
        s.save()
        self.assertRedirects(self.client.get(reverse('espace_dossier')), accueil)

    def test_espace_etudiant_reinscription_et_validation(self):
        from .models import DemandeEtudiant, Matiere, Note
        self._espace()
        s1, _ = self._semestres()
        Session.objects.filter(pk=self.session.pk).update(reinscriptions_ouvertes=True)
        e = Etudiant.objects.create(matricule='R1', prenom='Ali', nom='SOW', email='ali@t.local', ine='INER1',
                                    telephone='611')
        derniere = Inscription.objects.create(etudiant=e, session=self.s2425, licence=self.l2, filiere=self.info,
                                              promotion=self.p18, cycle='Licence')
        algo = Matiere.objects.create(nom='Algorithmique', filiere=self.info, semestre=s1, licence=self.l2,
                                      session=self.s2425, credit=6)
        Note.objects.create(etudiant=e, matiere=algo, moyenne=3.5, moyenne_deposee=True)
        # Préinscrit ou mauvais matricule : refusé
        r = self.client.post(reverse('espace_accueil'), {'mode': 'reinscription', 'r-ine': 'INER1', 'r-matricule': 'R2'})
        self.assertContains(r, 'Aucun dossier ne correspond')
        r = self.client.post(reverse('espace_accueil'), {'mode': 'reinscription', 'r-ine': 'iner1', 'r-matricule': 'r1'})
        self.assertRedirects(r, reverse('espace_dossier'))
        r = self.client.get(reverse('espace_dossier'))
        self.assertContains(r, 'Dernier niveau')
        self.assertContains(r, f'<option value="{self.l3.pk}" selected>')        # niveau prévu pour P18
        self.assertNotContains(r, f'<option value="{self.l2.pk}"')                # pas d'autre niveau
        self.assertContains(r, 'value="boursier" class="btn-check" id="id_boursier_0" required checked')   # par défaut
        self.assertContains(r, 'value="initiale" class="btn-check" id="id_type_formation_0" required checked')
        # Relevé de notes de la dernière année : obligatoire pour une réinscription
        sans_releve = {'telephone': '622 00 00 01', 'email': 'ali.sow@exemple.com', 'adresse': 'Horé Fello/Mamou',
                       'tuteur': 'Mamadou SOW', 'contact_de_tuteur': '664 00 00 02', 'licence': self.l3.pk,
                       'boursier': 'payante', 'type_formation': 'continue', 'photo': photo_jpeg(), **signer()}
        r = self.client.post(reverse('espace_dossier'), sans_releve)
        self.assertContains(r, 'Obligatoire : scannez ou envoyez cette pièce.')
        self.assertFalse(DemandeEtudiant.objects.filter(type='reinscription').exists())
        r = self.client.post(reverse('espace_dossier'), {
            'telephone': '622 00 00 01', 'email': 'ali.sow@exemple.com', 'adresse': 'Horé Fello/Mamou',
            'tuteur': 'Mamadou SOW', 'contact_de_tuteur': '664 00 00 02', 'licence': self.l3.pk,
            'boursier': 'payante', 'type_formation': 'continue', **signer(),
            'photo': photo_jpeg(), 'piece_releve_notes': SimpleUploadedFile('releve.pdf', b'%PDF-1.4 releve')})
        self.assertRedirects(r, reverse('espace_confirmation'))
        d = DemandeEtudiant.objects.get(type='reinscription')
        self.assertEqual((d.licence, d.derniere_inscription, d.boursier), (self.l3, derniere, 'payante'))
        self.assertFalse(Inscription.objects.filter(etudiant=e, session=self.session).exists())
        # Scolarité : liste, fiche avec les notes empilées (note < 5 en rouge), refus sans motif impossible
        self.connecter('info')
        self.assertEqual(self.client.get(reverse('demandes_liste')).status_code, 403)
        self.connecter('scol')
        r = self.client.get(reverse('demandes_liste'))
        self.assertContains(r, 'SOW')
        self.assertContains(r, '<span class="ini">AS</span>')                   # initiales, pas de photo
        self.assertNotContains(r, 'class="photo-demande"')
        r = self.client.get(reverse('demande_detail', args=[d.pk]))
        self.assertContains(r, 'class="photos-demande"')                         # photo de l'étudiant
        self.assertContains(r, 'Algorithmique')
        self.assertContains(r, 'class="note-rouge"')
        self.assertContains(r, '622 00 00 01')
        self.assertContains(r, 'value="payante" checked')                      # choix de l'étudiant, coché
        r = self.client.post(reverse('demande_detail', args=[d.pk]), {'action': 'refuser'}, follow=True)
        self.assertContains(r, 'Indiquez le motif')
        # Fiche de réinscription imprimable avant la validation : niveau prévu, données de la demande, rien d'enregistré
        from .views_demandes import inscription_de_la_fiche
        self.assertContains(self.client.get(reverse('demandes_liste')), reverse('demande_fiche_pdf', args=[d.pk]))
        r = self.client.get(reverse('demande_fiche_pdf', args=[d.pk]))
        self.assertEqual((r.status_code, r['Content-Type']), (200, 'application/pdf'))
        self.assertIn('fiche_reinscription_R1', r['Content-Disposition'])
        apercu, photo = inscription_de_la_fiche(DemandeEtudiant.objects.get(pk=d.pk))
        self.assertEqual((apercu.pk, apercu.licence, apercu.promotion, apercu.etudiant.telephone, apercu.etudiant.boursier),
                         (None, self.l3, self.p18, '622 00 00 01', 'payante'))
        self.assertTrue(photo)                                                    # photo proposée, en attente
        self.assertFalse(Inscription.objects.filter(etudiant=e, session=self.session).exists())
        self.assertNotEqual(Etudiant.objects.get(pk=e.pk).telephone, '622 00 00 01')
        # Validation : la scolarité change la formation ; inscription créée, coordonnées et choix appliqués
        # Décision préremplie : niveau prévu pour P18 (Licence 3), sa promotion, son programme et son cycle
        r = self.client.get(reverse('demande_detail', args=[d.pk]))
        for attendu in (f'<option value="{self.l3.pk}" selected>', f'<option value="{self.p18.pk}" selected>',
                        f'<option value="{self.info.pk}" selected>', '<option value="Licence" selected>'):
            self.assertContains(r, attendu)
        self.assertContains(r, 'id="promotions-prevues"')                         # niveau → promotion prévue
        # Choix incomplet : refusé, rien n'est créé
        r = self.client.post(reverse('demande_detail', args=[d.pk]), {'action': 'valider', 'licence': self.l3.pk},
                             follow=True)
        self.assertContains(r, 'Choisissez le niveau, la promotion et le programme')
        self.assertFalse(Inscription.objects.filter(etudiant=e, session=self.session).exists())
        self.client.post(reverse('demande_detail', args=[d.pk]), {
            'action': 'valider', 'boursier': 'payante', 'type_formation': 'initiale', 'licence': self.l3.pk,
            'promotion': self.p18.pk, 'filiere': self.info.pk, 'cycle': 'Licence'})
        insc = Inscription.objects.get(etudiant=e, session=self.session)
        self.assertEqual((insc.licence, insc.filiere, insc.promotion, insc.cycle, insc.redouble),
                         (self.l3, self.info, self.p18, 'Licence', False))
        e.refresh_from_db()
        self.assertEqual((e.telephone, e.boursier, e.type_formation), ('622 00 00 01', 'payante', 'initiale'))
        # Validée : la fiche est celle de l'inscription créée ; refusée : pas de fiche
        self.assertEqual(inscription_de_la_fiche(DemandeEtudiant.objects.get(pk=d.pk))[0], insc)
        self.assertEqual(self.client.get(reverse('demande_fiche_pdf', args=[d.pk])).status_code, 200)
        DemandeEtudiant.objects.filter(pk=d.pk).update(statut=DemandeEtudiant.REFUSEE)
        self.assertRedirects(self.client.get(reverse('demande_fiche_pdf', args=[d.pk])),
                             reverse('demande_detail', args=[d.pk]))
        # Déjà inscrit pour la session : l'espace le dit
        self.client.logout()
        r = self.client.post(reverse('espace_accueil'), {'mode': 'reinscription', 'r-ine': 'INER1', 'r-matricule': 'R1'})
        self.assertContains(r, 'déjà inscrit(e)')

    def test_demandes_reglees_ailleurs_quittent_la_liste(self):
        from .models import DemandeEtudiant
        self._espace()
        p = Etudiant.objects.create(matricule='P7', prenom='Awa', nom='KABA', email='kaba@t.local', preinscrit=True)
        Inscription.objects.create(etudiant=p, session=self.session, licence=self.licence, filiere=self.info)
        r1 = Etudiant.objects.create(matricule='R7', prenom='Ali', nom='DIALLO', email='diallo@t.local')
        Inscription.objects.create(etudiant=r1, session=self.s2425, licence=self.licence, filiere=self.info)
        d_insc = DemandeEtudiant.objects.create(etudiant=p, session=self.session, type='inscription',
                                                contacts={'telephone': '600'})
        d_reinsc = DemandeEtudiant.objects.create(etudiant=r1, session=self.session, type='reinscription',
                                                  licence=self.l2)
        self.connecter('scol')
        url = reverse('demandes_liste')
        self.assertContains(self.client.get(url, {'type': 'inscription'}), 'KABA')
        self.assertEqual(self.client.get(url).context['nb_demandes'], 2)
        # Réglées autrement : la case « Préinscrit » décochée, l'inscription de la session saisie à la main
        Etudiant.objects.filter(pk=p.pk).update(preinscrit=False)
        Inscription.objects.create(etudiant=r1, session=self.session, licence=self.l2, filiere=self.info)
        r = self.client.get(url, {'type': 'inscription'})
        self.assertNotContains(r, 'KABA')
        self.assertEqual(r.context['nb_demandes'], 0)
        self.assertNotContains(self.client.get(url, {'type': 'reinscription'}), 'DIALLO')
        for d in (d_insc, d_reinsc):
            d.refresh_from_db()
            self.assertEqual((d.statut, d.traitee_par), ('validee', None))
        p.refresh_from_db()
        self.assertNotEqual(p.telephone, '600')                                   # changements non appliqués
        self.assertContains(self.client.get(url, {'type': 'inscription', 'statut': 'validee'}), 'validée ailleurs')
        self.assertContains(self.client.get(reverse('demande_detail', args=[d_insc.pk])), 'hors des demandes en ligne')

    def test_espace_etudiant_promotions_prevues_seulement(self):
        """Seules les promotions prévues pour la session (Promotion_Licence_Session) passent en ligne."""
        self._espace()
        Session.objects.filter(pk=self.session.pk).update(inscriptions_ouvertes=True, reinscriptions_ouvertes=True)
        accueil = reverse('espace_accueil')
        # Préinscrit sans promotion, puis dans une promotion non prévue à ce niveau : scolarité
        p = Etudiant.objects.create(matricule='Q1', prenom='Ina', nom='SY', email='sy@t.local', ine='INEQ1',
                                    pvbac='111', preinscrit=True)
        insc = Inscription.objects.create(etudiant=p, session=self.session, licence=self.licence, filiere=self.info)
        hors = 'présentez-vous au Service de la Scolarité'
        self.assertContains(self.client.post(accueil, {'mode': 'inscription', 'i-ine': 'INEQ1', 'i-pvbac': '111'}), hors)
        Inscription.objects.filter(pk=insc.pk).update(promotion=self.p18)          # P18 n'est pas en Licence 1
        self.assertContains(self.client.post(accueil, {'mode': 'inscription', 'i-ine': 'INEQ1', 'i-pvbac': '111'}), hors)
        Inscription.objects.filter(pk=insc.pk).update(promotion=self.p20)          # P20 : Licence 1 en 2025-2026
        self.assertRedirects(self.client.post(accueil, {'mode': 'inscription', 'i-ine': 'INEQ1', 'i-pvbac': '111'}),
                             reverse('espace_dossier'))
        # Ancien étudiant d'une promotion sans niveau prévu (P21) : scolarité
        self.client.post(reverse('espace_quitter'))
        a = Etudiant.objects.create(matricule='Q2', prenom='Bah', nom='LY', email='ly@t.local', ine='INEQ2')
        Inscription.objects.create(etudiant=a, session=self.s2425, licence=self.licence, filiere=self.info,
                                   promotion=Promotion.objects.get(nom='P21'))
        self.assertContains(self.client.post(accueil, {'mode': 'reinscription', 'r-ine': 'INEQ2', 'r-matricule': 'Q2'}),
                            hors)

    def test_espace_etudiant_securite(self):
        """En-têtes stricts, essais limités par INE (même en changeant d'adresse), envois trop lourds refusés."""
        self._espace()
        Session.objects.filter(pk=self.session.pk).update(inscriptions_ouvertes=True)
        accueil = reverse('espace_accueil')
        r = self.client.get(accueil)
        self.assertIn("frame-ancestors 'none'", r['Content-Security-Policy'])
        self.assertIn("script-src 'self';", r['Content-Security-Policy'])
        self.assertEqual((r['X-Frame-Options'], r['X-Content-Type-Options']), ('DENY', 'nosniff'))
        self.assertIn('no-cache', r['Cache-Control'])
        self.assertNotIn('<script>', r.content.decode())                        # aucun script en ligne
        # Deviner le PV d'un INE en changeant d'adresse IP (ou de X-Forwarded-For) : bloqué au 5e échec
        p = Etudiant.objects.create(matricule='S1', prenom='Ina', nom='SY', email='sy@t.local', ine='INES1',
                                    pvbac='424242', preinscrit=True)
        Inscription.objects.create(etudiant=p, session=self.session, licence=self.licence, filiere=self.info,
                                   promotion=self.p20)
        for i in range(5):
            self.client.post(accueil, {'mode': 'inscription', 'i-ine': 'INES1', 'i-pvbac': str(100000 + i)},
                             REMOTE_ADDR=f'10.0.0.{i}', HTTP_X_FORWARDED_FOR=f'1.2.3.{i}')
        r = self.client.post(accueil, {'mode': 'inscription', 'i-ine': 'INES1', 'i-pvbac': '424242'},
                             REMOTE_ADDR='10.0.0.99')
        self.assertContains(r, "Trop d&#x27;essais")                            # même avec le bon PV
        # Un autre INE, depuis une autre adresse, n'est pas bloqué
        q = Etudiant.objects.create(matricule='S2', prenom='Ali', nom='BA', email='ba@t.local', ine='INES2',
                                    pvbac='777', preinscrit=True)
        Inscription.objects.create(etudiant=q, session=self.session, licence=self.licence, filiere=self.info,
                                   promotion=self.p20)
        r = self.client.post(accueil, {'mode': 'inscription', 'i-ine': 'INES2', 'i-pvbac': '777'}, REMOTE_ADDR='10.0.1.1')
        self.assertRedirects(r, reverse('espace_dossier'))
        # Envoi trop lourd : refusé avant tout traitement
        r = self.client.post(reverse('espace_dossier'), {'telephone': '6'}, CONTENT_LENGTH=str(50 * 1024 * 1024))
        self.assertRedirects(r, reverse('espace_dossier'), fetch_redirect_response=False)
        self.assertEqual(self.client.get(reverse('espace_dossier'))['Content-Security-Policy'][:19], "default-src 'self';")
        # Photo d'un autre étudiant : inaccessible (seulement celle de l'étudiant identifié, sinon 404)
        self.assertEqual(self.client.get(reverse('espace_photo')).status_code, 404)

    def test_reinscrire_un_etudiant_prerempli_avec_ses_notes(self):
        """Page « Réinscrire un étudiant » : inscription de la session préremplie et toutes ses notes."""
        from .models import Matiere, Note
        s1, _ = self._semestres()
        e = Etudiant.objects.create(matricule='R9', prenom='Ali', nom='SOW', email='sow9@t.local')
        Inscription.objects.create(etudiant=e, session=self.s2425, licence=self.l2, filiere=self.info,
                                   promotion=self.p18, cycle='Licence')
        m = Matiere.objects.create(nom='Réseaux', filiere=self.info, semestre=s1, licence=self.l2, session=self.s2425,
                                   credit=6)
        Note.objects.create(etudiant=e, matiere=m, moyenne=4, moyenne_deposee=True)
        self.connecter('scol')
        r = self.client.get(reverse('inscription_creer'), {'etudiant': e.pk})
        for attendu in (f'<option value="{self.session.pk}" selected>', f'<option value="{self.l3.pk}" selected>',
                        f'<option value="{self.p18.pk}" selected>', f'<option value="{self.info.pk}" selected>',
                        '<option value="Licence" selected>', 'Inscription préremplie', 'Réseaux', 'class="note-rouge"',
                        'id="promotions-prevues"'):
            self.assertContains(r, attendu)
        self.assertNotContains(r, 'id="id_insc-redouble" checked')
        # Déjà inscrit pour la session : signalé, rien de prérempli
        Inscription.objects.create(etudiant=e, session=self.session, licence=self.l3, filiere=self.info)
        r = self.client.get(reverse('inscription_creer'), {'etudiant': e.pk})
        self.assertContains(r, 'Déjà inscrit(e) pour')
        # Nouvel étudiant : pas de notes ni de préremplissage
        r = self.client.get(reverse('inscription_creer'))
        self.assertNotContains(r, 'Notes de l')

    def test_import_notes_inscrit_automatiquement(self):
        """Notes d'une session où l'étudiant n'est pas inscrit : la scolarité l'y inscrit automatiquement."""
        from .import_notes import importer_notes
        from .models import Matiere, Note
        from .roles import get_profil
        s1, _ = self._semestres()
        topo = Matiere.objects.create(nom='Topographie', filiere=self.gc, semestre=s1, licence=self.l2,
                                      session=self.s2425, credit=6)
        for prog in (self.info, self.gc):                                      # « Anglais » existe deux fois
            Matiere.objects.create(nom='Anglais', filiere=prog, semestre=s1, licence=self.l2, session=self.s2425,
                                   credit=2)
        Matiere.objects.create(nom='Chimie', filiere=self.gc, semestre=s1, licence=self.l3, session=self.s2425,
                               credit=2)
        entetes = ['matricule', 'matiere', 'session', 'programme', 'moyenne']
        fichier = lambda *l: fichier_xlsx([entetes, *l])           # noqa: E731
        lignes = [['M1', 'Topographie', '2024-2025', '', 6],           # 2 inscrit en Génie civil L2 (unique)
                  ['M1', 'Anglais', '2024-2025', 'Génie civil', 7],    # 3 levé par la colonne programme
                  ['M0', 'Anglais', '2024-2025', '', 5],               # 4 ambigu (deux programmes)
                  ['M2', 'Topographie', '2024-2025', '', 6],           # 5 M2 : deux licences (L2 et L3)…
                  ['M2', 'Chimie', '2024-2025', '', 6]]                # 6 … → pas d'inscription
        self.connecter('scol')
        scol = get_profil(self.client.get(reverse('notes_importer')).wsgi_request)
        r = importer_notes(fichier(*lignes), scol, simulation=True)
        self.assertEqual([(l[1], l[4], l[5], l[6]) for l in r.inscriptions_creees], [('M1', self.l2, self.gc, self.p18)])
        self.assertFalse(Inscription.objects.filter(session=self.s2425).exists())   # simulation : rien de créé
        self.assertIn('plusieurs programmes ou licences', dict(r.erreurs)[4])
        self.assertIn('plusieurs programmes ou licences', dict(r.erreurs)[5])
        r = importer_notes(fichier(*lignes), scol)
        insc = Inscription.objects.get(etudiant__matricule='M1', session=self.s2425)
        self.assertEqual((insc.licence, insc.filiere, insc.promotion, insc.redouble), (self.l2, self.gc, self.p18, False))
        self.assertEqual(Note.objects.get(etudiant__matricule='M1', matiere=topo).moyenne_finale, 6)
        self.assertEqual(r.creees, 2)
        self.assertFalse(Inscription.objects.filter(etudiant__matricule__in=['M0', 'M2'], session=self.s2425).exists())
        # Préinscrit : jamais inscrit automatiquement
        Etudiant.objects.filter(matricule='M0').update(preinscrit=True)
        r = importer_notes(fichier(['M0', 'Anglais', '2024-2025', 'Informatique', 5]), scol)
        self.assertIn('préinscrit', r.erreurs[0][1])
        # Chef de programme : pas d'inscription automatique
        Etudiant.objects.filter(matricule='M0').update(preinscrit=False)
        self.connecter('info')
        info = get_profil(self.client.get(reverse('notes_importer')).wsgi_request)
        r = importer_notes(fichier(['M0', 'Anglais', '2024-2025', 'Informatique', 5]), info,
                           pdf_enseignant=SimpleUploadedFile('f.pdf', b'%PDF-1.4 f'), simulation=True)
        self.assertIn("faites inscrire l'étudiant par la scolarité", r.erreurs[0][1])
        self.assertEqual(r.inscriptions_creees, [])

    def test_espace_etudiant_signature_obligatoire(self):
        from .models import DemandeEtudiant
        self._espace()
        Session.objects.filter(pk=self.session.pk).update(inscriptions_ouvertes=True)
        p = Etudiant.objects.create(matricule='SG1', prenom='Awa', nom='BAH', email='sg1@etudiant.gped', ine='INESG1',
                                    pvbac='321', preinscrit=True)
        insc = Inscription.objects.create(etudiant=p, session=self.session, licence=self.licence, filiere=self.info,
                                          promotion=self.p20)
        self.client.post(reverse('espace_accueil'), {'mode': 'inscription', 'i-ine': 'INESG1', 'i-pvbac': '321'})
        # Programme : son vrai nom, jamais la description (qui porte souvent le nom du département)
        Service.objects.filter(pk=self.info.pk).update(description='Génie Informatique', departement=self.dep)
        r = self.client.get(reverse('espace_dossier'))
        self.assertContains(r, '<div class="lib">Programme</div><div class="val">Informatique</div>', html=False)
        self.assertNotContains(r, 'Génie Informatique')
        self.assertContains(r, 'Certification et signature')
        self.assertContains(r, 'class="toile-signature"')
        base = {'telephone': '622 11 22 33', 'email': 'awa.sg@exemple.com', 'adresse': 'Petel/Mamou',
                'tuteur': 'Mariam CAMARA', 'contact_de_tuteur': '628 99 88 77'}
        pieces = lambda: {'photo': photo_jpeg(), 'piece_acte_naissance': SimpleUploadedFile('a.pdf', b'%PDF-1.4 a'),   # noqa: E731
                          'piece_diplome_bac': SimpleUploadedFile('d.pdf', b'%PDF-1.4 d'),
                          'piece_releve_bac': SimpleUploadedFile('r.pdf', b'%PDF-1.4 r')}
        # Sans certification ni signature, puis signature vide ou falsifiée : refusé
        r = self.client.post(reverse('espace_dossier'), {**base, **pieces()})
        self.assertContains(r, 'Cochez la case pour certifier')
        self.assertContains(r, 'Signez dans le cadre')
        r = self.client.post(reverse('espace_dossier'), {**base, **pieces(), **signer(vide=True)})
        self.assertContains(r, 'Signature trop courte ou vide')
        r = self.client.post(reverse('espace_dossier'), {**base, **pieces(), 'certifie': 'on',
                                                         'signature': 'data:image/png;base64,PHN2Zz4='})
        self.assertContains(r, 'Signature illisible')
        self.assertFalse(DemandeEtudiant.objects.exists())
        # Signée : enregistrée (image, date, appareil, empreinte)
        r = self.client.post(reverse('espace_dossier'), {**base, **pieces(), **signer()}, HTTP_USER_AGENT='Téléphone essai')
        self.assertRedirects(r, reverse('espace_confirmation'))
        d = DemandeEtudiant.objects.get()
        self.assertTrue(d.signature.name.startswith('demandes/signatures/'))
        self.assertEqual((d.signature_appareil, len(d.empreinte)), ('Téléphone essai', 64))
        self.assertIsNotNone(d.signee_le)
        # Retour sur la page : informations modifiables, récapitulatif, signature enregistrée affichée
        r = self.client.get(reverse('espace_dossier'))
        self.assertContains(r, 'value="628 99 88 77"')
        self.assertContains(r, 'class="recap-demande')
        self.assertContains(r, 'Signature enregistrée le')
        self.assertEqual(self.client.get(reverse('espace_signature'))['Content-Type'], 'image/png')
        # Demande envoyée avant la signature : « Signature : à ajouter », puis signée en la renvoyant
        DemandeEtudiant.objects.update(signature='', signee_le=None)
        r = self.client.get(reverse('espace_dossier'))
        self.assertContains(r, 'Signature : à ajouter')
        self.assertContains(r, "pas encore signée")
        r = self.client.post(reverse('espace_dossier'), {**base, 'tuteur': 'Mamadou BAH', **signer()},
                             HTTP_USER_AGENT='Téléphone essai')
        self.assertRedirects(r, reverse('espace_confirmation'))           # photo et pièces déjà fournies
        d.refresh_from_db()
        self.assertTrue(d.signature)
        self.assertEqual(d.contacts['tuteur'], 'Mamadou BAH')
        # Scolarité : signature visible sur la demande ; fiche d'inscription PDF avec la signature
        self.connecter('scol')
        self.assertContains(self.client.get(reverse('demande_detail', args=[d.pk])), 'Signée le')
        # Fiche d'inscription imprimable directement, avant la validation (préinscription de la session)
        r = self.client.get(reverse('demande_fiche_pdf', args=[d.pk]))
        self.assertEqual((r.status_code, r['Content-Type']), (200, 'application/pdf'))
        self.assertIn('fiche_inscription_SG1', r['Content-Disposition'])
        self.assertGreaterEqual(r.content.count(b'/Subtype /Image'), 2)              # photo et signature
        r = self.client.get(reverse('demande_signature', args=[d.pk]))
        self.assertEqual((r.status_code, r['Content-Type']), (200, 'image/png'))
        self.client.post(reverse('valider_preinscriptions'), {'ids': [insc.pk]})
        from . import fiche_inscription
        dessins = []
        original = fiche_inscription._signature_etudiant
        fiche_inscription._signature_etudiant = lambda c, i, x, y: (dessins.append(i.pk), original(c, i, x, y))
        try:
            r = self.client.get(reverse('inscription_fiche_pdf', args=[insc.pk]))
        finally:
            fiche_inscription._signature_etudiant = original
        self.assertEqual((r.status_code, dessins), (200, [insc.pk]))
        self.assertGreaterEqual(r.content.count(b'/Subtype /Image'), 1)

    def test_signataire_chef_scolarite_ou_chef_du_programme_en_master(self):
        from django.conf import settings
        from .models import ReleveNote
        from .pdf_releves import pdf_releve
        from .releves import contenu
        from .signataires import signataire, signataire_inscription
        # Aucun chef enregistré : le réglage actuel
        self.assertEqual(signataire().nom, settings.GPED_ATTESTATION['signataire'])
        # Chef de la scolarité : civilité, prénom, NOM ; sa fonction ; ses coordonnées dans l'en-tête
        Personnel.objects.filter(utilisateur__username='scol').update(
            chef=True, civilite='M.', prenom='Lanciné 2', nom='Conde', fonc='Le Chef Service de la Scolarité',
            email='scol@ist.gn', telephone='+224 600')
        s = signataire()
        self.assertEqual((s.nom, s.fonction, s.service, s.email, s.telephone, s.scolarite),
                         ('M. Lanciné 2 CONDE', 'Le Chef Service de la Scolarité', 'SERVICE SCOLARITÉ', 'scol@ist.gn',
                          '+224 600', True))
        # Master : le chef du programme (son programme dans l'en-tête, pas d'image de cachet de la scolarité)
        Personnel.objects.filter(utilisateur__username='info').update(
            chef=True, civilite='Dr', prenom='Awa', nom='Bah', fonc='Chef du programme', email='master@ist.gn',
            telephone='+224 611')
        insc = Inscription.objects.get(etudiant__matricule='M0')
        self.assertEqual(signataire_inscription(insc).nom, 'M. Lanciné 2 CONDE')             # Licence
        Inscription.objects.filter(pk=insc.pk).update(cycle='Master')
        insc.refresh_from_db()
        m = signataire_inscription(insc)
        self.assertEqual((m.nom, m.fonction, m.service, m.email, m.scolarite),
                         ('Dr Awa BAH', 'Chef du programme', 'INFORMATIQUE', 'master@ist.gn', False))
        self.assertEqual(signataire(self.gc, cycle='Master').nom, 'M. Lanciné 2 CONDE')       # pas de chef : scolarité
        self.assertEqual(signataire(self.info, licence=Licence(nom='MASTER 1')).nom, 'Dr Awa BAH')
        # Documents : fiche d'inscription et relevé de l'étudiant en Master (textes dessinés dans le PDF)
        from unittest import mock
        from reportlab.pdfgen.canvas import Canvas

        def textes_dessines(produire):
            vus = []
            originaux = {n: getattr(Canvas, n) for n in ('drawString', 'drawCentredString', 'drawRightString')}

            def espion(nom):
                return lambda c, x, y, texte, *a, **k: (vus.append(str(texte)), originaux[nom](c, x, y, texte, *a, **k))[1]
            with mock.patch.multiple(Canvas, **{n: espion(n) for n in originaux}):
                produire()
            return ' | '.join(vus)
        self.connecter('scol')
        textes = textes_dessines(lambda: self.client.get(reverse('inscription_fiche_pdf', args=[insc.pk])))
        for attendu in ('Dr Awa BAH', 'Email : master@ist.gn', 'Tél : +224 611', 'INFORMATIQUE', 'CHEF DU PROGRAMME'):
            self.assertIn(attendu, textes)
        self.assertNotIn('CONDE', textes)
        import datetime
        releve = ReleveNote(type='E', etudiant=insc.etudiant, filiere=self.info, num_ref='ESSAI',
                            date_delivrance=datetime.date.today())
        textes = textes_dessines(lambda: pdf_releve(releve, contenu(releve), cachet_signature=True))
        self.assertIn('Dr Awa BAH', textes)
        # Licence : signée par le chef de la scolarité, ses coordonnées dans l'en-tête
        Inscription.objects.filter(pk=insc.pk).update(cycle='Licence')
        textes = textes_dessines(lambda: self.client.get(reverse('inscription_fiche_pdf', args=[insc.pk])))
        for attendu in ('M. Lanciné 2 CONDE', 'LE CHEF SERVICE DE LA SCOLARITÉ', 'SERVICE SCOLARITÉ', 'scol@ist.gn'):
            self.assertIn(attendu, textes)

    def test_personnel_chef_seulement_scolarite_ou_programme(self):
        from .forms import PersonnelForm
        admin = User.objects.create_superuser('admin_chef', 'a@t.local', 'x')
        donnees = {'civilite': 'Mme', 'prenom': 'Awa', 'nom': 'Bah', 'fonc': 'Chef', 'email': 'awa@t.local',
                   'telephone': '999', 'service': [self.dep.pk], 'chef': 'on'}
        f = PersonnelForm(donnees, editeur=admin)
        self.assertFalse(f.is_valid())
        self.assertIn('Impossible', f.errors['chef'][0])                      # simple service : pas de chef
        self.assertTrue(PersonnelForm({**donnees, 'service': [self.gc.pk]}, editeur=admin).is_valid())   # programme
        self.assertTrue(PersonnelForm({**donnees, 'service': [self.scol.pk]}, editeur=admin).is_valid())
        f = PersonnelForm({**donnees, 'service': [self.gc.pk]}, editeur=admin)
        self.assertTrue(f.is_valid(), f.errors)
        self.assertEqual(f.save().nom_signature, 'Mme Awa BAH')

    def test_import_des_notes(self):
        from .import_matieres import ImportRefuse
        from .import_notes import importer_notes
        from .models import Deliberation, Note, NoteHistorique
        from .roles import get_profil
        s1, algo, anglais = self._preparer_notes()
        entetes = ['matricule', 'matiere', 'session', 'controle', 'tp', 'examen', 'moyenne', 'rattrapage']
        fichier = lambda *l: fichier_xlsx([entetes, *l])           # noqa: E731
        lignes = [
            ['M0', 'Algorithmique', '2025 - 2026', 6, 7, 8, None, None],          # 2 → 7,15
            ['M1', 'alg1', '2025-2026', None, None, None, '6,75', None],          # 3 code, moyenne déposée
            ['M0', 'ALGORITHMIQUE', '2025-2026', 9, 9, 9, None, None],            # 4 doublon
            ['M9', 'Algorithmique', '2025-2026', None, None, None, 5, None],      # 5 matricule inconnu
            ['M1', 'Inconnue', '2025-2026', None, None, None, 5, None],           # 6 matière introuvable
            ['M0', 'Algorithmique', '2024-2025', None, None, None, 5, None],      # 7 pas inscrit
            ['M0', 'Anglais', '2025-2026', None, None, 'ABS', None, None],        # 8 absent = 0
            ['M1', 'Anglais', '2025-2026', 12, None, None, None, None],           # 9 hors barème
            ['M2', 'Algorithmique', '2025-2026', None, None, None, 5, None],      # 10 autre programme
        ]
        self.connecter('scol')
        scol = get_profil(self.client.get(reverse('notes_importer')).wsgi_request)
        r = importer_notes(fichier(*lignes), scol, simulation=True)
        self.assertEqual((r.lignes, r.creees, r.nb_doublons), (9, 3, 1))
        self.assertEqual([n for n, _ in r.erreurs], [5, 6, 7, 9, 10])
        self.assertFalse(Note.objects.exists())                                   # simulation
        r = importer_notes(fichier(*lignes), scol)
        self.assertEqual(Note.objects.get(etudiant__matricule='M0', matiere=algo).moyenne_finale, 7.15)
        self.assertEqual(Note.objects.get(etudiant__matricule='M1', matiere=algo).moyenne_finale, 6.75)
        self.assertEqual(Note.objects.get(etudiant__matricule='M0', matiere=anglais).moyenne_finale, 0)
        # Réimport : inchangé ; puis mise à jour : refusée sans le PDF de l'enseignant, tracée avec
        import shutil, tempfile
        from django.test import override_settings
        from .models import ImportNotes
        dossier = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, dossier, True)
        pdf = lambda: SimpleUploadedFile('fiche_algo.pdf', b'%PDF-1.4 fiche')   # noqa: E731
        r = importer_notes(fichier(*lignes[:3]), scol)
        self.assertEqual((r.creees, r.mises_a_jour, r.inchangees), (0, 0, 2))
        modif = fichier(['M1', 'Algorithmique', '2025-2026', None, None, None, 5, None])
        r = importer_notes(modif, scol)
        self.assertEqual(r.mises_a_jour, 0)
        self.assertIn("joignez le PDF de l'enseignant", r.erreurs[0][1])
        with override_settings(BASE_DIR=pathlib_path(dossier)):
            r = importer_notes(fichier(['M1', 'Algorithmique', '2025-2026', None, None, None, 5, None]), scol,
                               pdf_enseignant=pdf())
        self.assertEqual((r.mises_a_jour, NoteHistorique.objects.count()), (1, 1))
        self.assertEqual(ImportNotes.objects.count(), 2)                     # chaque import réel est journalisé
        journal = ImportNotes.objects.filter(modifiees=1).get()
        self.assertEqual((journal.modifiees, NoteHistorique.objects.get().justificatif.name), (1, journal.pdf_enseignant.name))
        self.assertTrue((pathlib_path(dossier) / 'prive' / journal.pdf_enseignant.name).exists())
        # Moyenne remplie et contrôle / TP / examen tous à 0 : ces 0 sont vides ; un seul 0 avec d'autres notes reste
        with override_settings(BASE_DIR=pathlib_path(dossier)):
            importer_notes(fichier(['M0', 'Anglais', '2025-2026', 0, 0, 0, '6,5', None],
                                   ['M1', 'Anglais', '2025-2026', 0, 0, 8, None, None]), scol, pdf_enseignant=pdf())
        a0 = Note.objects.get(etudiant__matricule='M0', matiere=anglais)
        self.assertEqual((a0.controle, a0.tp, a0.examen, a0.moyenne_finale), (None, None, None, 6.5))
        a1 = Note.objects.get(etudiant__matricule='M1', matiere=anglais)
        self.assertEqual((a1.controle, a1.tp, a1.examen, a1.moyenne_finale), (0, 0, 8, 3.2))   # sans moyenne : gardés

        # Chef de programme : le PDF de l'enseignant est obligatoire (formulaire et import)
        self.connecter('info')
        info = get_profil(self.client.get(reverse('notes_importer')).wsgi_request)
        with self.assertRaisesMessage(ValueError, "PDF de l'enseignant"):
            importer_notes(fichier(['M0', 'Anglais', '2025-2026', None, None, None, 8, None]), info)
        r = self.client.post(reverse('notes_importer'), {'fichier': SimpleUploadedFile('n.csv', b'matricule;matiere;session;moyenne\n')})
        self.assertContains(r, 'Ce champ est obligatoire')
        # Une ligne d'un autre programme → tout est refusé
        self._matiere('Topographie', self.gc, s1)
        with self.assertRaisesMessage(ImportRefuse, 'ligne 3'):
            importer_notes(fichier(['M0', 'Anglais', '2025-2026', None, None, None, 8, None],
                                   ['M2', 'Topographie', '2025-2026', None, None, None, 6, None]), info,
                           pdf_enseignant=pdf())
        self.assertEqual(Note.objects.get(etudiant__matricule='M0', matiere=anglais).moyenne_finale, 6.5)  # inchangée
        # Semestre délibéré de l'année en cours : la note est mise à jour (avec le PDF)
        Deliberation.objects.create(filiere=self.info, session=self.session, semestre=s1)
        with override_settings(BASE_DIR=pathlib_path(dossier)):
            r = importer_notes(fichier(['M0', 'Anglais', '2025-2026', None, None, None, 8, None]), info,
                               User.objects.get(username='info'), pdf_enseignant=pdf())
        self.assertEqual((r.mises_a_jour, r.erreurs), (1, []))
        # Journal : le chef de programme voit le PDF de son import
        imp = ImportNotes.objects.filter(utilisateur__username='info').get()
        with override_settings(BASE_DIR=pathlib_path(dossier)):
            self.assertEqual(self.client.get(reverse('notes_import_pdf', args=[imp.pk]))['Content-Type'], 'application/pdf')
        self.assertContains(self.client.get(reverse('notes_importer')), 'Derniers imports')

        # Année écoulée : notes manquantes créées ; note déjà enregistrée modifiée seulement par un chef
        from .models import Matiere
        e0 = Etudiant.objects.get(matricule='M0')
        Inscription.objects.create(etudiant=e0, session=self.s2425, licence=self.l2, filiere=self.info)
        reseaux = Matiere.objects.create(nom='Réseaux', filiere=self.info, semestre=s1, licence=self.l2,
                                         session=self.s2425, credit=6)
        ligne = ['M0', 'Réseaux', '2024-2025', None, None, None, 6, None]
        r = importer_notes(fichier(ligne), scol)
        self.assertEqual(r.creees, 1)
        with override_settings(BASE_DIR=pathlib_path(dossier)):
            r = importer_notes(fichier(ligne[:6] + [7, None]), scol, pdf_enseignant=pdf())
        self.assertIn('année écoulée', r.erreurs[0][1])
        self.assertEqual(Note.objects.get(etudiant=e0, matiere=reseaux).moyenne_finale, 6)
        self.assertEqual(self.client.get(reverse('notes_modele')).status_code, 200)

    def test_resultats_et_deliberation(self):
        from .models import Deliberation, Note
        s1, algo, _ = self._preparer_notes()
        Note.objects.create(etudiant=Etudiant.objects.get(matricule='M0'), matiere=algo, moyenne=8, moyenne_deposee=True)
        params = {'session': self.session.pk, 'programme': self.info.pk, 'semestre': s1.pk}
        self.connecter('info')
        r = self.client.get(reverse('notes_resultats'), params)
        self.assertContains(r, 'N0')
        self.assertNotContains(r, 'Délibérer le semestre')                      # réservé à la scolarité
        self.assertEqual(self.client.post(reverse('notes_deliberer'), params).status_code, 403)
        self.connecter('scol')
        self.assertContains(self.client.get(reverse('notes_resultats'), params), 'Délibérer le semestre')
        self.client.post(reverse('notes_deliberer'), params)
        self.assertTrue(Deliberation.objects.filter(filiere=self.info, semestre=s1).exists())
        # Annuler : super utilisateur seulement
        self.assertEqual(self.client.post(reverse('notes_deliberer'), {**params, 'action': 'annuler'}).status_code, 403)
        User.objects.create_superuser('chef', password='x')
        self.client.login(username='chef', password='x')
        self.client.post(reverse('notes_deliberer'), {**params, 'action': 'annuler'})
        self.assertFalse(Deliberation.objects.exists())

    # ---------------------------------------------------------------- personnel et comptes

    def test_gestion_du_personnel_et_invitation(self):
        from django.core import mail
        liste, creer = reverse('personnel_liste'), reverse('personnel_creer')
        fiche = {'prenom': ' mariam ', 'nom': 'Camara', 'fonc': 'Cheffe de programme', 'email': 'Mariam@ist.gn',
                 'telephone': '+224 620  11 22 33', 'service': [self.info.pk], 'envoyer_invitation': 'on'}
        # Réservé aux super utilisateurs : ni les autres services, ni la scolarité (menu masqué)
        self.connecter('dep')
        self.assertEqual(self.client.get(liste).status_code, 403)
        self.connecter('scol')
        r = self.client.get(liste)
        self.assertContains(r, 'réservée aux super utilisateurs', status_code=403)           # page « Accès refusé »
        self.assertContains(r, 'Se connecter avec un autre compte', status_code=403)
        self.assertNotContains(self.client.get(reverse('tableau_de_bord')), 'href="/scolarite/personnel/"')
        User.objects.create_superuser('chef', password='x')
        self.client.login(username='chef', password='x')
        self.assertContains(self.client.get(reverse('tableau_de_bord')), 'href="/scolarite/personnel/"')
        self.assertContains(self.client.get(liste), 'Personnel')
        self.assertContains(self.client.post(creer, {**fiche, 'service': []}), 'Cochez au moins un service')
        self.assertContains(self.client.post(creer, {**fiche, 'email': 'scol@t.local'}), 'déjà utilisée')
        r = self.client.post(creer, fiche)
        self.assertRedirects(r, liste)
        p = Personnel.objects.get(email='mariam@ist.gn')
        u = p.utilisateur
        self.assertEqual((p.prenom, p.nom, p.telephone, u.username, u.is_active, u.has_usable_password()),
                         ('Mariam', 'CAMARA', '+224 620 11 22 33', 'mariam.camara', True, False))
        # E-mail d'invitation : identifiant, rôle, lien d'activation
        self.assertEqual(len(mail.outbox), 1)
        m = mail.outbox[0]
        self.assertEqual((m.to, m.subject), (['mariam@ist.gn'], 'Votre compte gePed a été créé'))
        self.assertIn('mariam.camara', m.body)
        self.assertIn('Programme', m.body)
        lien = re.search(r'https?://\S+/compte/activer/\S+/', m.body).group(0)
        # Activation : choix du mot de passe, connexion directe, droits du service
        self.client.logout()
        r = self.client.get(lien, follow=True)
        self.assertContains(r, 'Activer mon compte')
        r = self.client.post(r.redirect_chain[-1][0], {'new_password1': 'Mamou@2026!', 'new_password2': 'Mamou@2026!'},
                             follow=True)
        self.assertRedirects(r, reverse('tableau_de_bord'))
        self.assertContains(r, 'votre compte est activé')
        u.refresh_from_db()
        self.assertTrue(u.has_usable_password())
        self.assertEqual(self.client.get(reverse('matieres_liste')).status_code, 200)           # chef de programme
        self.assertEqual(self.client.get(liste).status_code, 403)
        self.assertContains(self.client.get(lien, follow=True), 'invalide')                     # lien utilisé
        # Connexion avec l'adresse e-mail
        self.client.logout()
        self.assertTrue(self.client.login(username='MARIAM@ist.gn', password='Mamou@2026!'))

        # Gestion du compte par le super utilisateur : lien de mot de passe, désactivation, réactivation
        self.client.login(username='chef', password='x')
        action = reverse('personnel_action', args=[p.pk])
        self.client.post(action, {'action': 'reinitialiser'})
        self.assertIn('/reinitialiser/', mail.outbox[-1].body)
        self.client.post(action, {'action': 'desactiver'})
        u.refresh_from_db()
        self.assertFalse(u.is_active)
        self.assertFalse(self.client.login(username='mariam.camara', password='Mamou@2026!'))
        self.client.login(username='chef', password='x')
        self.client.post(action, {'action': 'reactiver'})
        u.refresh_from_db()
        self.assertTrue(u.is_active)
        # La scolarité ne peut pas non plus agir sur un compte
        self.connecter('scol')
        self.assertEqual(self.client.post(action, {'action': 'desactiver'}).status_code, 403)
        # Le super utilisateur peut rattacher quelqu'un au Service Scolarité ; sans invitation : aucun e-mail
        self.client.login(username='chef', password='x')
        p_scol = Personnel.objects.get(utilisateur__username='scol')
        self.assertEqual(self.client.get(reverse('personnel_modifier', args=[p_scol.pk])).status_code, 200)
        r = self.client.post(creer, {**fiche, 'prenom': 'Awa', 'email': 'awa@ist.gn', 'telephone': '+224 621',
                                     'service': [self.scol.pk], 'envoyer_invitation': ''})
        awa = Personnel.objects.get(email='awa@ist.gn')
        self.assertEqual((awa.utilisateur.username, len(mail.outbox)), ('awa.camara', 2))   # sans e-mail
        # Personne ne désactive son propre compte (même un super utilisateur)
        chef = User.objects.get(username='chef')
        p_chef = Personnel.objects.create(utilisateur=chef, prenom='Chef', nom='T', telephone='77', email='chef@t.local',
                                          fonc='Admin')
        self.client.post(reverse('personnel_action', args=[p_chef.pk]), {'action': 'desactiver'})
        chef.refresh_from_db()
        self.assertTrue(chef.is_active)

    # ---------------------------------------------------------------- relevés de notes

    def test_releves_de_notes(self):
        from django.utils import timezone
        from .models import Deliberation, Matiere, Note, ReleveNote
        from .notes import lettre
        from .releves import _definitif, contenu
        s1, algo, anglais = self._preparer_notes()
        e0 = Etudiant.objects.get(matricule='M0')
        insc = Inscription.objects.get(etudiant=e0, session=self.session)
        Note.objects.create(etudiant=e0, matiere=algo, moyenne=7.33, moyenne_deposee=True)
        Note.objects.create(etudiant=e0, matiere=anglais, moyenne=5, moyenne_deposee=True)
        self.assertEqual([lettre(v) for v in (8, 7.99, 6, 5, 4.99)], ['A', 'B', 'C', 'D', 'E'])

        self.connecter('info')                                                    # réservé à la scolarité
        self.assertEqual(self.client.get(reverse('releves')).status_code, 403)
        self.connecter('scol')
        r = self.client.get(reverse('releves'), {'q': 'M0'})
        self.assertContains(r, 'Délivrer le relevé de SEMESTRE 1')
        creer = lambda **d: self.client.post(reverse('releve_creer'), d)          # noqa: E731
        r = creer(type='S', inscription=insc.pk, semestre=s1.pk)
        rel = ReleveNote.objects.get()
        self.assertRedirects(r, reverse('releve_pdf', args=[rel.pk]), fetch_redirect_response=False)
        annee = str(timezone.localdate().year)[-2:]
        self.assertEqual(rel.num_ref, f'0001/ISTM/SC/{annee}')
        creer(type='S', inscription=insc.pk, semestre=s1.pk)                      # même relevé : réimprimé
        self.assertEqual(ReleveNote.objects.count(), 1)
        creer(type='A', inscription=insc.pk)
        self.assertEqual(ReleveNote.objects.get(type='A').num_ref, f'0002/ISTM/SC/{annee}')

        # Moyenne pondérée par les crédits et TRONQUÉE : (7,33 × 6 + 5 × 4) / 10 = 6,398 → 6,39
        c = contenu(rel)
        self.assertEqual((c['moyenne'], c['mention'], c['definitif']), (6.39, 'Assez bien', False))
        self.assertEqual([(l['lettre'], l['observation']) for l in c['blocs'][0]['lignes']],
                         [('B', 'Validée'), ('D', 'Validée')])
        r = self.client.get(reverse('releve_pdf', args=[rel.pk]))
        self.assertEqual((r.status_code, r.content[:4]), (200, b'%PDF'))
        # Délibéré : plus « provisoire » ; une session antérieure est considérée comme délibérée
        Deliberation.objects.create(filiere=self.info, session=self.session, semestre=s1)
        self.assertTrue(contenu(rel)['definitif'])
        self.assertTrue(_definitif(self.info.pk, self.s2425, []))

        # Extrait : par programme (le Tronc commun a le sien) ; sans soutenance, générale = moyenne du cycle
        tc = Service.objects.create(nom='Tronc commun', code='TC', programme=True, cycle='Ingénieur')
        Inscription.objects.create(etudiant=e0, session=self.s2425, licence=self.licence, filiere=tc)
        m_tc = Matiere.objects.create(nom='Physique', filiere=tc, semestre=s1, licence=self.licence,
                                      session=self.s2425, credit=6)
        Note.objects.create(etudiant=e0, matiere=m_tc, moyenne=9, moyenne_deposee=True)
        creer(type='E', etudiant=e0.pk, programme=self.info.pk)
        extrait = ReleveNote.objects.get(type='E')
        c = contenu(extrait)
        self.assertEqual([l['matiere'] for b in c['blocs'] for l in b['lignes']], ['Algorithmique', 'Anglais'])
        self.assertEqual((c['moyenne_cycle'], c['moyenne'], c['soutenance']), (6.39, 6.39, None))
        c_tc = contenu(ReleveNote(type='E', etudiant=e0, filiere=tc))
        self.assertEqual([l['matiere'] for b in c_tc['blocs'] for l in b['lignes']], ['Physique'])
        self.assertTrue(c_tc['definitif'])                                         # 2024-2025 : session antérieure
        # Avec une soutenance DANS CE PROGRAMME : 6,39 × 0,87 + 9 × 0,13 = 6,7293 → 6,72 ; pas dans l'extrait du TC
        from .models import Soutenance
        Soutenance.objects.create(etudiant=e0, cycle='Licence', filiere=self.info, session=self.session,
                                  promotion=self.p18, note=9)
        c = contenu(extrait)
        self.assertEqual((c['moyenne_cycle'], c['soutenance'], c['moyenne'], c['mention']), (6.39, 9, 6.72, 'Assez bien'))
        self.assertIsNone(contenu(ReleveNote(type='E', etudiant=e0, filiere=tc))['soutenance'])

        # QR code : page publique de vérification
        r = self.client.get(reverse('verifier_document', args=[extrait.code_verification]))
        self.assertContains(r, 'Extrait du livret universitaire')
        self.assertContains(r, '6,72/10')                                         # moyenne générale, soutenance comprise
        # Cachet et signature : super utilisateur seulement
        normal = self.client.get(reverse('releve_pdf', args=[rel.pk]), {'signature': 1}).content
        User.objects.create_superuser('chef', password='x')
        self.client.login(username='chef', password='x')
        signe = self.client.get(reverse('releve_pdf', args=[rel.pk]), {'signature': 1})
        self.assertIn('_signe', signe['Content-Disposition'])
        self.assertGreater(len(signe.content), len(normal))
        # Suppression : le numéro n'est pas réutilisé
        self.client.post(reverse('releve_supprimer', args=[extrait.pk]))
        self.assertEqual(ReleveNote.prochaine_reference(), f'0004/ISTM/SC/{annee}')

        # Formulaire « Nouveau relevé » : enregistrement, reprise d'un relevé existant, erreurs
        page = reverse('releves')
        self.assertContains(self.client.post(page, {'matricule': 'M0', 'type': 'S', 'licence': self.licence.pk}),
                            'Choisissez le semestre')
        self.assertContains(self.client.post(page, {'matricule': 'M0', 'type': 'A', 'licence': self.l2.pk}),
                            "aucune inscription en Licence 2")
        self.assertContains(self.client.post(page, {'matricule': 'M0', 'type': 'E', 'programme': self.gc.pk}),
                            "aucune inscription en Génie civil")
        r = self.client.post(page, {'matricule': 'M0', 'type': 'E', 'programme': self.info.pk})
        nouveau = ReleveNote.objects.get(type='E', filiere=self.info)
        self.assertRedirects(r, f"{page}?etudiant={e0.pk}&imprimer={nouveau.pk}")
        self.assertEqual(nouveau.num_ref, f'0004/ISTM/SC/{annee}')
        r = self.client.post(page, {'matricule': 'm0', 'type': 'S', 'semestre': s1.pk}, follow=True)   # existant
        self.assertContains(r, 'déjà été délivré')
        self.assertContains(r, f'0001/ISTM/SC/{annee}')
        self.assertEqual(ReleveNote.objects.count(), 3)
        # Aucune note : pas de relevé (formulaire et bouton rapide)
        insc_m1 = Inscription.objects.get(etudiant__matricule='M1')
        self.assertContains(self.client.post(page, {'matricule': 'M1', 'type': 'A', 'licence': self.licence.pk}),
                            "aucune note")
        r = creer(type='A', inscription=insc_m1.pk)
        self.assertRedirects(r, f"{page}?etudiant={insc_m1.etudiant_id}", fetch_redirect_response=False)
        self.assertEqual(ReleveNote.objects.count(), 3)

    # ---------------------------------------------------------------- statistiques

    def test_statistiques_des_inscriptions(self):
        import datetime
        Etudiant.objects.filter(matricule='M1').update(sexe='F', date_de_naissance=datetime.date(2005, 3, 1),
                                                       nationalite='GUINEENNE', option_de_bac='SCIENCES MATHEMATIQUES')
        Etudiant.objects.filter(matricule='M0').update(date_de_naissance=datetime.date(2202, 4, 28),   # incohérente
                                                       nationalite='Guinéenné')
        Inscription.objects.filter(etudiant__matricule='M2').update(redouble=True)
        Etudiant.objects.filter(matricule='M1').update(handicap=True, type_handicap='visuel')
        url = reverse('statistiques')
        self.connecter('sans_service')
        self.assertEqual(self.client.get(url).status_code, 403)
        self.connecter('scol')
        r = self.client.get(url)
        niveaux, total = r.context['niveaux']
        self.assertEqual((total.h, total.f, total.total, total.pct_f), (2, 1, 3, 33.3))
        programmes = {l.libelle: (l.h, l.f) for l in r.context['programmes'][0]}
        self.assertEqual(programmes, {'Informatique': (1, 1), 'Génie civil': (1, 0)})
        self.assertEqual(r.context['feminisation'][0].libelle, 'Informatique')
        ages = r.context['ages']
        self.assertEqual((ages['renseignees'], ages['incoherentes'], ages['moyenne_f']), (1, 1, 20))     # 2025 - 2005
        nat = r.context['origine']['nationalites']
        self.assertEqual(([(l.libelle, l.total) for l in nat['lignes']], nat['taux']), ([('Guinéenne', 2)], 66.7))
        self.assertEqual(r.context['redoublants']['nb'], 1)
        evo = {l.libelle: l for l in r.context['evolution']}
        self.assertEqual((evo['2025-2026'].total, evo['2025-2026'].extra['nouveaux']), (3, 3))      # tous en Licence 1
        self.assertContains(r, 'Statistiques des inscriptions')
        # Handicap : par type, et part des inscrits du programme ; plus d'onglet « Niveau et programme »
        ha = r.context['handicap']
        self.assertEqual((ha['nb'], ha['taux']), (1, 33.3))
        self.assertEqual([(l.libelle, l.h, l.f) for l in ha['types'][0]], [('Visuel', 0, 1)])
        self.assertEqual({l.libelle: l.extra['taux'] for l in ha['programmes'][0]}, {'Informatique': 50.0})
        self.assertEqual(r.context['croise']['total'].total, 3)
        self.assertNotContains(r, 'Niveau et programme')
        # Fiche de l'étudiant : la situation de handicap n'est montrée qu'à la scolarité
        fiche = reverse('etudiant_fiche', args=[Inscription.objects.get(etudiant__matricule='M1').pk])
        self.assertContains(self.client.get(fiche), 'Visuel')
        # Chef de programme : ses programmes seulement
        self.connecter('info')
        self.assertEqual(self.client.get(url).context['niveaux'][1].total, 2)
        self.assertNotContains(self.client.get(fiche), 'Situation de handicap')
        # Export Excel : une feuille par statistique
        r = self.client.get(url, {'export': 'xlsx'})
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertIn('Par niveau', wb.sheetnames)
        self.assertIn('Soutenances', wb.sheetnames)
        self.assertIn('Handicap par type', wb.sheetnames)
        self.assertIn('Programme x niveau', wb.sheetnames)

    # ---------------------------------------------------------------- soutenances

    def _fiche(self, matricule, **extra):
        return {'matricule': matricule, 'filiere': self.info.pk, 'session': self.session.pk,
                'promotion': self.p18.pk, **extra}

    def test_soutenances(self):
        from .models import Note, Resultat, Soutenance, SoutenanceHistorique
        s1, algo, anglais = self._preparer_notes()
        e0 = Etudiant.objects.get(matricule='M0')
        Note.objects.create(etudiant=e0, matiere=algo, moyenne=7, moyenne_deposee=True)
        Note.objects.create(etudiant=e0, matiere=anglais, moyenne=7, moyenne_deposee=True)
        creer = reverse('soutenance_creer')
        self.connecter('dep')                                         # autre service : pas d'accès
        self.assertEqual(self.client.get(reverse('soutenances')).status_code, 403)
        self.connecter('info')
        r = self.client.post(creer, self._fiche('M2', filiere=self.gc.pk, note='8'))
        self.assertContains(r, 'Sélectionnez un choix valide')       # Génie civil : pas son programme
        r = self.client.post(creer, self._fiche('M0', note='9'), follow=True)
        self.assertRedirects(r, reverse('soutenances'))
        self.assertContains(r, 'Ingénieur et Master')                 # cycle Licence : simple avertissement
        s = Soutenance.objects.get()
        self.assertEqual((s.cycle, s.filiere, s.session, s.promotion, s.note_finale, s.verrouillee),
                         ('Licence', self.info, self.session, self.p18, 9, False))
        c = Resultat.objects.get(etudiant=e0, niveau=Resultat.CYCLE)
        self.assertEqual((c.moyenne, c.soutenance, c.moyenne_generale), (7, 9, 7.26))      # 7×0,87 + 9×0,13
        self.assertContains(self.client.get(reverse('soutenances')), '7,26')
        self.assertContains(self.client.post(creer, self._fiche('M0', note='5')), 'déjà une soutenance')
        # Note de session : la meilleure est retenue, avec la mention « Session »
        modifier = reverse('soutenance_modifier', args=[s.pk])
        self.client.post(modifier, self._fiche('M0', note='4', note_session='6,5'))
        s.refresh_from_db()
        self.assertEqual((s.note, s.note_session, s.note_finale, s.par_session), (4, 6.5, 6.5, True))
        self.assertEqual(SoutenanceHistorique.objects.count(), 1)    # modification tracée

        # Verrouillée : le chef de programme ne la change plus ; la scolarité oui, avec un motif
        Soutenance.objects.filter(pk=s.pk).update(verrouillee=True)
        self.client.post(modifier, self._fiche('M0', note='10', note_session='6,5'))
        s.refresh_from_db()
        self.assertEqual(s.note, 4)
        self.assertEqual(self.client.post(reverse('soutenance_supprimer', args=[s.pk])).status_code, 403)
        self.connecter('scol')
        self.assertContains(self.client.post(modifier, self._fiche('M0', note='8', note_session='6,5')),
                            'indiquez le motif')
        self.client.post(modifier, self._fiche('M0', note='8', note_session='6,5', motif='PV du jury'))
        s.refresh_from_db()
        self.assertEqual((s.note_finale, s.par_session), (8, False))
        h = SoutenanceHistorique.objects.first()
        self.assertEqual((h.motif, h.note_verrouillee, h.avant['note'], h.apres['note']), ('PV du jury', True, 4, 8))
        self.client.post(reverse('soutenance_supprimer', args=[s.pk]))
        self.assertFalse(Soutenance.objects.exists())
        # Sans soutenance : la moyenne générale est la moyenne du cycle
        self.assertEqual(Resultat.objects.get(etudiant=e0, niveau=Resultat.CYCLE).moyenne_generale, 7)
        # Aide à la saisie : dernière inscription de l'étudiant
        d = self.client.get(reverse('soutenance_etudiant'), {'matricule': 'M1'}).json()
        self.assertEqual((d['trouve'], d['programme'], d['session']), (True, self.info.pk, self.session.pk))

    def test_departement_du_programme(self):
        """Champ `departement` du Service : département d'un programme, sinon le programme lui-même."""
        from .attestation_soutenance import _paragraphe
        from .attestations import _nom_departement, contenu_inscription
        from .fiche_inscription import pdf_fiche_inscription
        from .models import Attestation
        # Informatique : département ET programme ; Systèmes embarqués : programme du département Informatique
        Service.objects.filter(pk=self.info.pk).update(departement=self.info)
        se = Service.objects.create(nom='Systèmes embarqués', code='SE', programme=True, departement_id=self.info.pk)
        self.info.refresh_from_db()
        self.assertEqual((_nom_departement(self.info), _nom_departement(se), _nom_departement(self.gc)),
                         ('Informatique', 'Informatique', 'Génie Civil'))       # Génie civil : pas de département
        self.assertEqual([p.nom for p in self.info.programmes_rattaches.order_by('nom')],
                         ['Informatique', 'Systèmes embarqués'])
        # Fiche d'inscription : « Département : » = le département du programme
        insc = Inscription.objects.get(etudiant__matricule='M0')
        Inscription.objects.filter(pk=insc.pk).update(filiere=se)
        insc.refresh_from_db()
        from . import fiche_inscription
        lus = []
        original = fiche_inscription._champ
        fiche_inscription._champ = lambda c, x, y, libelle, valeur, fin: (lus.append((libelle, valeur)),
                                                                          original(c, x, y, libelle, valeur, fin))
        try:
            pdf_fiche_inscription(insc)
        finally:
            fiche_inscription._champ = original
        self.assertIn(('Département :', 'Informatique'), lus)
        self.assertIn(('PROGRAMME :', 'Systèmes embarqués'), lus)
        # Attestation de soutenance : département, puis le programme s'il est différent
        a = Attestation(type=Attestation.SOUTENANCE, etudiant=insc.etudiant, filiere=se, session=self.session)
        texte = _paragraphe(insc.etudiant, None, a)
        self.assertIn('du département <b>Informatique</b>, programme <b>Systèmes Embarqués</b>', texte)
        a.filiere = self.info
        self.assertNotIn('programme <b>', _paragraphe(insc.etudiant, None, a))
        # Attestation d'inscription : le programme seulement (pas le département)
        a = Attestation(type=Attestation.INSCRIPTION, etudiant=insc.etudiant, filiere=se, session=self.session,
                        licence=self.licence)
        self.assertIn('Département <b>Systèmes Embarqués</b>', contenu_inscription(a)[1][0])

    def test_attestation_de_soutenance(self):
        import datetime
        from .models import Attestation, ProjetSoutenance, Soutenance
        Service.objects.filter(pk=self.info.pk).update(cycle='Ingénieur')
        p = ProjetSoutenance.objects.create(intitule='Projet 1', theme='Réseau du campus', filiere=self.info,
                                            session=self.session, promotion=self.p18, cycle='Ingénieur',
                                            encadreur='Dr Barry', chef_departement='M. Sow', president_jury='Pr Diallo',
                                            date=datetime.date(2026, 7, 15))
        s = Soutenance.objects.create(etudiant=Etudiant.objects.get(matricule='M0'), cycle='Ingénieur',
                                      filiere=self.info, session=self.session, promotion=self.p18, projet=p, note=4)
        url = reverse('attestations_type', args=['soutenance'])
        self.connecter('info')                                                    # réservé à la scolarité
        self.assertEqual(self.client.get(url).status_code, 403)
        self.connecter('scol')
        self.assertContains(self.client.post(url, {'matricule': 'M1'}), "aucune soutenance")
        self.assertContains(self.client.post(url, {'matricule': 'M0'}), "pas validée")           # 4/10
        Soutenance.objects.filter(pk=s.pk).update(note=8.5, note_finale=8.5)
        r = self.client.post(url, {'matricule': 'M0'})
        a = Attestation.objects.get(type=Attestation.SOUTENANCE)
        self.assertRedirects(r, f"{url}?imprimer={a.pk}")
        self.assertEqual((a.cycle, a.filiere, a.session), ('Ingénieur', self.info, self.session))
        self.assertContains(self.client.post(url, {'matricule': 'M0'}), "déjà une attestation de soutenance")
        # PDF : thème, signataires du projet et du service de la scolarité
        r = self.client.get(reverse('attestation_pdf', args=[a.pk]))
        self.assertEqual((r.status_code, r.content[:4]), (200, b'%PDF'))
        import base64
        import zlib
        from .attestation_soutenance import pdf_soutenance

        def decoder(flux):                      # pages ReportLab : ASCII85 puis Flate
            try:
                return zlib.decompress(base64.a85decode(flux.strip().removesuffix(b'~>')))
            except Exception:
                return b''
        brut = pdf_soutenance(a, 'http://x/v/abc')
        texte = b''.join(decoder(m) for m in re.findall(rb'stream\r?\n(.*?)endstream', brut, re.S))
        for attendu in (b'Diallo', b'Barry', b'Sow', b'campus'):
            self.assertIn(attendu, texte)
        # Page publique de vérification
        r = self.client.get(reverse('verifier_document', args=[a.code_verification]))
        self.assertContains(r, 'Réseau du campus')
        self.assertContains(r, 'Très bien')
        # Bouton « Attestation » dans la liste des soutenances
        self.assertContains(self.client.get(reverse('soutenances')), 'Attestation de soutenance')

    def test_projets_notes_et_verrouillage(self):
        import datetime
        from .models import ProjetSoutenance, Soutenance
        self._semestres()
        self.connecter('info')
        donnees = {'intitule': 'Projet 1', 'theme': 'Réseau du campus', 'filiere': self.info.pk,
                   'session': self.session.pk, 'promotion': self.p18.pk, 'statut': 'programmee',
                   'encadreur': 'Dr Barry', 'chef_departement': 'M. Sow', 'president_jury': 'Pr Diallo',
                   'date': '2026-07-15', 'heure': '09:00', 'salle': 'Amphi A'}
        r = self.client.post(reverse('projet_creer'), {**donnees, 'ajouter': 'M0, M9'})
        self.assertContains(r, 'M9 : matricule inconnu')
        self.assertFalse(ProjetSoutenance.objects.exists())
        self.client.post(reverse('projet_creer'), {**donnees, 'ajouter': 'M0 M1'})
        p = ProjetSoutenance.objects.get()
        self.assertEqual(sorted(p.membres.values_list('etudiant__matricule', flat=True)), ['M0', 'M1'])
        self.assertEqual(p.cycle, 'Licence')                          # repris du programme
        # Même intitulé dans le même programme, session et promotion : refusé
        self.assertContains(self.client.post(reverse('projet_creer'), {**donnees, 'intitule': 'projet 1'}),
                            'existe déjà')

        # Grille des notes : enregistrer, puis verrouiller
        s0, s1 = (Soutenance.objects.get(etudiant__matricule=m) for m in ('M0', 'M1'))
        grille = (f"{reverse('soutenances_notes')}?session={self.session.pk}&programme={self.info.pk}"
                  f"&promotion={self.p18.pk}")
        r = self.client.get(grille)
        self.assertContains(r, 'Projet 1')
        self.assertContains(r, 'Même note pour tout le projet')
        self.assertContains(self.client.post(grille, {f'note_{s0.pk}': '11'}), 'note hors barème')
        self.client.post(grille, {f'note_{s0.pk}': '8', f'note_{s1.pk}': 'ABS', 'action': 'enregistrer'})
        s0.refresh_from_db(), s1.refresh_from_db()
        self.assertEqual((s0.note_finale, s1.note_finale, s0.verrouillee), (8, 0, False))
        self.client.post(grille, {f'note_{s0.pk}': '8', f'note_{s1.pk}': '0', 'action': 'verrouiller'})
        s0.refresh_from_db()
        self.assertTrue(s0.verrouillee)
        self.assertContains(self.client.get(grille), 'disabled')
        self.client.post(grille, {f'note_{s0.pk}': '2', 'action': 'enregistrer'})      # chef : ignoré
        s0.refresh_from_db()
        self.assertEqual(s0.note, 8)
        self.assertEqual(self.client.post(grille, {'action': 'deverrouiller'}).status_code, 403)
        self.connecter('scol')
        self.assertContains(self.client.post(grille, {f'note_{s0.pk}': '9', f'note_{s1.pk}': '0'}), 'indiquez le motif')
        self.client.post(grille, {f'note_{s0.pk}': '9', f'note_{s1.pk}': '0', 'motif': 'Erreur de report'})
        s0.refresh_from_db()
        self.assertEqual(s0.note, 9)

        # PV de soutenance (PDF) et calendrier : conflit de salle et de jury au même créneau
        self.connecter('info')
        r = self.client.get(reverse('projet_pv', args=[p.pk]))
        self.assertEqual((r.status_code, r['Content-Type'], r.content[:4]), (200, 'application/pdf', b'%PDF'))
        ProjetSoutenance.objects.create(intitule='Projet 2', filiere=self.info, session=self.session,
                                        promotion=self.p18, date=datetime.date(2026, 7, 15),
                                        heure=datetime.time(9, 0), salle='amphi a', president_jury='Pr  Diallo')
        r = self.client.get(reverse('soutenances_calendrier'), {'session': self.session.pk})
        self.assertContains(r, 'Salle Amphi A déjà prise')
        self.assertContains(r, 'Pr Diallo est aussi attendu')
        r = self.client.get(reverse('soutenances_calendrier'), {'session': self.session.pk, 'format': 'pdf'})
        self.assertEqual(r.content[:4], b'%PDF')
        # Retirer un membre : il garde sa soutenance notée (sans projet)
        self.client.post(reverse('projet_modifier', args=[p.pk]), {**donnees, 'retirer': [s1.pk]})
        s1.refresh_from_db()
        self.assertIsNone(s1.projet)
        self.client.post(reverse('projet_supprimer', args=[p.pk]))
        self.assertTrue(Soutenance.objects.filter(pk=s0.pk, projet__isnull=True).exists())

    def test_import_des_soutenances(self):
        from .import_matieres import ImportRefuse
        from .import_soutenances import importer_soutenances
        from .models import ProjetSoutenance, Soutenance
        from .roles import get_profil
        self._semestres()
        Service.objects.filter(pk=self.info.pk).update(cycle='Ingénieur')
        entetes = ['matricule', 'programme', 'session', 'promotion', 'note', 'note_session', 'projet', 'theme',
                   'encadreur', 'date', 'heure']
        fichier = lambda *l: fichier_xlsx([entetes, *l])           # noqa: E731
        lignes = [
            ['M0', 'Informatique', '2025 - 2026', 'P18', 8, None, 'Projet 1', 'Réseau du campus', 'Dr Barry',
             '15/07/2026', '9h'],                                                       # 2
            ['M1', 'INFO', '2025-2026', 'p 18', 'ABS', 6, 'projet 1'],                 # 3 même projet, session
            ['M0', 'Informatique', '2025-2026', 'P18', 5],                             # 4 doublon
            ['M9', 'Informatique', '2025-2026', 'P18', 5],                             # 5 matricule inconnu
            ['M1', 'Inconnu', '2025-2026', 'P18', 5],                                  # 6 programme inconnu
            ['M1', 'Informatique', '2025-2026', 'P99', 5],                             # 7 promotion inconnue
            ['M2', 'Génie civil', '2020-2021', 'P18', 7],                              # 8 session inconnue
            ['M2', 'Génie civil', '2025-2026', 'P18', 12],                             # 9 hors barème
            ['M2', 'Génie civil', '2025-2026', 'P18', None, None, None, 'Pont en arc'],  # 10 projet individuel
        ]
        self.connecter('scol')
        scol = get_profil(self.client.get(reverse('soutenances_importer')).wsgi_request)
        r = importer_soutenances(fichier(*lignes), scol, simulation=True)
        self.assertEqual((r.lignes, r.creees, r.nb_doublons, r.projets_crees), (9, 3, 1, 2))
        self.assertEqual([n for n, _ in r.erreurs], [5, 6, 7, 8, 9])
        self.assertTrue(any('Licence' in a for a in r.avertissements))            # Génie civil : cycle Licence
        self.assertFalse(Soutenance.objects.exists())                                # simulation
        importer_soutenances(fichier(*lignes), scol)
        p = ProjetSoutenance.objects.get(intitule='Projet 1')
        self.assertEqual((p.theme, p.encadreur, p.statut, p.heure.hour, p.membres.count()),
                         ('Réseau du campus', 'Dr Barry', 'soutenue', 9, 2))
        s0, s1 = (Soutenance.objects.get(etudiant__matricule=m) for m in ('M0', 'M1'))
        self.assertEqual((s0.cycle, s0.note_finale, s0.verrouillee), ('Ingénieur', 8, True))   # verrouillée
        self.assertEqual((s1.note_finale, s1.par_session), (6, True))
        s2 = Soutenance.objects.get(etudiant__matricule='M2')
        self.assertEqual((s2.projet.theme, s2.note_finale, s2.verrouillee), ('Pont en arc', None, False))

        # Réimport : la note verrouillée ne change pas, les informations du projet oui
        r = importer_soutenances(fichier(['M0', 'Informatique', '2025-2026', 'P18', 3, None, 'Projet 1', None,
                                          'Pr Camara']), scol)
        self.assertIn('verrouillée', r.avertissements[0])
        s0.refresh_from_db(), p.refresh_from_db()
        self.assertEqual((s0.note, p.encadreur, p.theme), (8, 'Pr Camara', 'Réseau du campus'))

        # Chef de programme : une ligne d'un autre programme → tout est refusé
        self.connecter('info')
        info = get_profil(self.client.get(reverse('soutenances_importer')).wsgi_request)
        with self.assertRaisesMessage(ImportRefuse, 'ligne 3'):
            importer_soutenances(fichier(['M1', 'Informatique', '2025-2026', 'P18', 9],
                                         ['M2', 'Génie civil', '2025-2026', 'P18', 9]), info)
        self.assertEqual(self.client.get(reverse('soutenances_modele')).status_code, 200)
        r = self.client.get(reverse('soutenances'), {'export': 'xlsx'})
        self.assertEqual(r.status_code, 200)

    def test_liste_des_notes(self):
        from .models import Note
        s1, algo, anglais = self._preparer_notes()
        topo = self._matiere('Topographie', self.gc, s1)
        e0, e1, e2 = (Etudiant.objects.get(matricule=f'M{k}') for k in range(3))
        Note.objects.create(etudiant=e0, matiere=algo, moyenne=8, moyenne_deposee=True)
        Note.objects.create(etudiant=e1, matiere=algo, controle=2, examen=3)                  # 2,6 : non validée
        Note.objects.create(etudiant=e1, matiere=anglais, examen=4, rattrapage=6)            # 6 en session
        Note.objects.create(etudiant=e2, matiere=topo, moyenne=7, moyenne_deposee=True)
        url = reverse('notes_liste')
        matricules = lambda **p: [n.etudiant.matricule for n in self.client.get(url, p).context['page']]  # noqa: E731

        self.connecter('scol')
        self.assertEqual(sorted(matricules()), ['M0', 'M1', 'M1', 'M2'])
        self.assertEqual(matricules(statut='non_validee'), ['M1'])
        self.assertEqual(matricules(statut='session'), ['M1'])
        self.assertEqual(matricules(tri='finale', ordre='desc'), ['M0', 'M2', 'M1', 'M1'])
        self.assertEqual(matricules(q='topo'), ['M2'])
        r = self.client.get(url)
        self.assertContains(r, 'Session')
        self.assertContains(r, reverse('notes_saisie', args=[algo.pk]))

        # Chef de programme : uniquement ses programmes ; export Excel des notes filtrées
        self.connecter('info')
        self.assertEqual(sorted(matricules()), ['M0', 'M1', 'M1'])
        r = self.client.get(url, {'export': 'xlsx', 'statut': 'validee'})
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        lignes = list(ws.iter_rows(values_only=True))
        self.assertEqual(lignes[0][:4], ('Matricule', 'Nom', 'Prénom', 'Matière'))
        self.assertEqual(sorted((l[0], l[3], l[16], l[18]) for l in lignes[1:]),
                         [('M0', 'Algorithmique', 8, 'non'), ('M1', 'Anglais', 6, 'oui')])
        self.connecter('dep')
        self.assertEqual(self.client.get(url).status_code, 403)

    def test_fichier_xlsx_endommage_repare_automatiquement(self):
        import zipfile
        from .import_matieres import importer_matieres
        from .import_notes import importer_notes
        from .models import Matiere, Note
        from .roles import get_profil
        s1, _ = self._semestres()
        # Fichier « produit par un autre logiciel » : \x02 dans un intitulé, attribut « biltinId » dans les styles
        propre = fichier_xlsx([['nom', 'programme', 'licence', 'semestre', 'session'],
                               ['Probabilité Statistique', 'Informatique', 'Licence 1', 'S1', '2025-2026']]).read()
        abime = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(propre)) as zin, zipfile.ZipFile(abime, 'w') as zout:
            for info in zin.infolist():
                brut = zin.read(info.filename)
                if info.filename == 'xl/sharedStrings.xml':
                    brut = brut.replace('Probabilité Statistique'.encode(), 'Probabilité\x02Statistique'.encode())
                if info.filename == 'xl/styles.xml':
                    brut = brut.replace(b'builtinId=', b'biltinId=')
                zout.writestr(info, brut)
        self.connecter('scol')
        scol = get_profil(self.client.get(reverse('matieres_importer')).wsgi_request)
        r = importer_matieres(SimpleUploadedFile('abime.xlsx', abime.getvalue()), scol)
        self.assertEqual((r.creees, r.nb_erreurs), (1, 0))
        self.assertGreaterEqual(r.reparations, 1)
        self.assertTrue(Matiere.objects.filter(nom='Probabilité Statistique').exists())
        # La page signale la réparation
        r = self.client.post(reverse('matieres_importer'), {'fichier': SimpleUploadedFile('abime.xlsx', abime.getvalue()),
                                                             'simulation': 'on'})
        self.assertContains(r, 'réparé automatiquement')
        # Fichier illisible : message clair, pas d'erreur serveur
        r = self.client.post(reverse('matieres_importer'), {'fichier': SimpleUploadedFile('faux.xlsx', b'pas un zip')})
        self.assertContains(r, 'endommagé et n&#x27;a pas pu être lu')

        # Notes : intitulé retrouvé sans tenir compte de la ponctuation (« _ » , « , »)
        Matiere.objects.filter(nom='Probabilité Statistique').update(nom='Probabilité_Statistique')
        r = importer_notes(fichier_xlsx([['matricule', 'matiere', 'session', 'moyenne'],
                                         ['M0', 'Probabilité Statistique', '2025-2026', 7]]), scol)
        self.assertEqual((r.creees, r.nb_erreurs), (1, 0))
        self.assertEqual(Note.objects.get(etudiant__matricule='M0').matiere.nom, 'Probabilité_Statistique')
