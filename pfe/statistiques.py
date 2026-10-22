"""
Statistiques des inscriptions (menu Étudiants → Statistiques).

Base : les inscriptions visibles par le profil (scolarité et départements : toutes ; chef de
programme : ses programmes), hors préinscrits, filtrées par session, cycle et programme.
Chaque tableau donne Hommes / Femmes / Total / % de femmes. Les statistiques dont les données
sont incomplètes (âge, origine scolaire…) indiquent le taux de fiches renseignées : elles
deviennent plus justes à mesure que les fiches sont complétées.

Rapide même sur toutes les sessions : les inscriptions sont lues UNE fois (charger), puis
toutes les statistiques sont calculées en mémoire. Rien n'est enregistré.
"""
from collections import OrderedDict, defaultdict
from types import SimpleNamespace

from .attestations import est_premiere_annee
from .importer import cle
from .notes import regles

TRANCHES_AGE = [(0, 17, 'Moins de 18 ans'), (18, 20, '18 à 20 ans'), (21, 23, '21 à 23 ans'), (24, 26, '24 à 26 ans'),
                (27, 30, '27 à 30 ans'), (31, 200, '31 ans et plus')]
AGE_MIN, AGE_MAX = 14, 70                 # en dehors : date de naissance incohérente (ex. 2202, 1905)
NON_RENSEIGNE = '— Non renseigné —'

CHAMPS = {                                # attribut de la ligne → champ de l'inscription
    'session_id': 'session_id', 'session': 'session__annee', 'niveau': 'licence__nom', 'programme': 'filiere__nom',
    'promotion': 'promotion__nom', 'redouble': 'redouble', 'sexe': 'etudiant__sexe',
    'naissance': 'etudiant__date_de_naissance', 'nationalite': 'etudiant__nationalite',
    'option': 'etudiant__option_de_bac', 'centre': 'etudiant__centre_de_bac', 'ecole': 'etudiant__ecole_origine',
    'bac': 'etudiant__moyenne_au_bac', 'handicap': 'etudiant__handicap', 'type_handicap': 'etudiant__type_handicap',
}


def charger(qs):
    """Une seule requête : les champs utiles de chaque inscription."""
    noms = list(CHAMPS)
    return [SimpleNamespace(**dict(zip(noms, v))) for v in qs.values_list(*CHAMPS.values()).order_by()]


def pct(n, total):
    return round(100 * n / total, 1) if total else 0


def _f(r):
    return 1 if r.sexe == 'F' else 0


class Ligne:
    """Une ligne H/F : libellé, hommes, femmes (+ colonnes en plus)."""
    def __init__(self, libelle, h=0, f=0, **extra):
        self.libelle, self.h, self.f = libelle, h, f
        self.extra = extra

    @property
    def total(self):
        return self.h + self.f

    @property
    def pct_f(self):
        return pct(self.f, self.total)


def _total(lignes, libelle='Total'):
    return Ligne(libelle, sum(l.h for l in lignes), sum(l.f for l in lignes))


def repartition(lignes_insc, attr):
    """Lignes H/F groupées par l'attribut `attr` (niveau, programme, promotion), plus le total."""
    donnees = defaultdict(lambda: [0, 0])
    for r in lignes_insc:
        v = getattr(r, attr)
        donnees[v if v not in (None, '') else NON_RENSEIGNE][_f(r)] += 1
    lignes = [Ligne(str(k), *donnees[k]) for k in sorted(donnees, key=str)]
    return lignes, _total(lignes)


def croise(lignes_insc):
    """Programme × niveau : programmes en lignes, niveaux en colonnes, H/F dans chaque case."""
    cases = defaultdict(lambda: defaultdict(lambda: [0, 0]))
    for r in lignes_insc:
        cases[r.programme or '—'][r.niveau][_f(r)] += 1
    niveaux = sorted({n for prog in cases.values() for n in prog})
    lignes = []
    for prog in sorted(cases):
        cellules = [Ligne(n, *cases[prog][n]) if n in cases[prog] else None for n in niveaux]
        lignes.append({'programme': prog, 'cellules': cellules, 'total': _total([c for c in cellules if c])})
    totaux = [_total([l['cellules'][i] for l in lignes if l['cellules'][i]]) for i in range(len(niveaux))]
    return {'niveaux': niveaux, 'lignes': lignes, 'totaux': totaux, 'total': _total(totaux)}


def _avec_taux(lignes_insc, selection, attr):
    """Lignes H/F de `selection` groupées par `attr`, avec le taux (%) parmi tous les inscrits du groupe."""
    effectifs = defaultdict(int)
    for r in lignes_insc:
        v = getattr(r, attr)
        effectifs[v if v not in (None, '') else NON_RENSEIGNE] += 1
    lignes, somme = repartition(selection, attr)
    for l in lignes:
        l.extra['taux'] = pct(l.total, effectifs.get(l.libelle, 0))
    somme.extra['taux'] = pct(somme.total, len(lignes_insc))
    return lignes, somme


def evolution(lignes_insc):
    """Par session : H/F, et nouveaux inscrits (première année) / réinscrits."""
    premieres = {}
    par_session = defaultdict(lambda: {'h': 0, 'f': 0, 'nouveaux': 0, 'reinscrits': 0})
    for r in lignes_insc:
        d = par_session[r.session]
        d['f' if r.sexe == 'F' else 'h'] += 1
        if r.niveau not in premieres:
            premieres[r.niveau] = est_premiere_annee(SimpleNamespace(nom=r.niveau))
        d['nouveaux' if premieres[r.niveau] else 'reinscrits'] += 1
    lignes = [Ligne(annee, d['h'], d['f'], nouveaux=d['nouveaux'], reinscrits=d['reinscrits'])
              for annee, d in sorted(par_session.items())]
    maximum = max((l.total for l in lignes), default=0)
    for l in lignes:
        l.extra['hauteur_h'] = pct(l.h, maximum)
        l.extra['hauteur_f'] = pct(l.f, maximum)
    return lignes


def ages(lignes_insc):
    """Tranches d'âge (au 31 décembre de la 1re année de la session), H/F ; âge moyen."""
    tranches = OrderedDict((lib, [0, 0]) for _, _, lib in TRANCHES_AGE)
    somme = {0: [0, 0], 1: [0, 0]}                   # sexe (0 = H, 1 = F) → [somme des âges, effectif]
    incoherentes = 0
    for r in lignes_insc:
        if r.naissance is None or not str(r.session or '')[:4].isdigit():
            continue
        age = int(str(r.session)[:4]) - r.naissance.year
        if not AGE_MIN <= age <= AGE_MAX:
            incoherentes += 1
            continue
        s = _f(r)
        somme[s][0] += age
        somme[s][1] += 1
        for mini, maxi, lib in TRANCHES_AGE:
            if mini <= age <= maxi:
                tranches[lib][s] += 1
                break
    lignes = [Ligne(lib, *hf) for lib, hf in tranches.items()]
    moyenne = lambda s: round(somme[s][0] / somme[s][1], 1) if somme[s][1] else None   # noqa: E731
    renseignees = somme[0][1] + somme[1][1]
    return {'lignes': lignes, 'total': _total(lignes), 'renseignees': renseignees, 'incoherentes': incoherentes,
            'taux': pct(renseignees, len(lignes_insc)), 'moyenne_h': moyenne(0), 'moyenne_f': moyenne(1),
            'moyenne': round((somme[0][0] + somme[1][0]) / renseignees, 1) if renseignees else None}


def _nationalite(v):
    k = cle(v)
    if not k:
        return None
    return 'Guinéenne' if k.startswith('GUIN') else k.title()


def _option(v):
    k = cle(v)
    if not k or k.isdigit():
        return None
    if 'MATH' in k:
        return 'Sciences mathématiques'
    if 'EXPERIM' in k:
        return 'Sciences expérimentales' + (' (franco-arabe)' if 'ARABE' in k else '')
    if 'SOCIAL' in k:
        return 'Sciences sociales'
    return k.capitalize()


def _texte(v):
    s = ' '.join(str(v or '').split()).upper()
    return s if s and s not in ('/', '-', '.') else None


def _par_valeur(lignes_insc, attr, normaliser, limite=None):
    """Lignes H/F d'un champ texte de la fiche, normalisé ; taux de fiches renseignées."""
    compte, cache = defaultdict(lambda: [0, 0]), {}
    for r in lignes_insc:
        brut = getattr(r, attr)
        if brut not in cache:
            cache[brut] = normaliser(brut)
        if cache[brut]:
            compte[cache[brut]][_f(r)] += 1
    lignes = sorted((Ligne(k, *hf) for k, hf in compte.items()), key=lambda l: (-l.total, l.libelle))
    if limite and len(lignes) > limite:
        lignes = lignes[:limite] + [_total(lignes[limite:], f"Autres ({len(lignes) - limite})")]
    total = _total(lignes)
    return {'lignes': lignes, 'taux': pct(total.total, len(lignes_insc)), 'total': total}


def origine(lignes_insc):
    """Nationalité, option du bac, centres et écoles d'origine, moyenne au bac par programme."""
    bac = defaultdict(lambda: [0.0, 0])
    for r in lignes_insc:
        if r.bac:                                 # 0 : non renseignée
            bac[r.programme or '—'][0] += float(r.bac)
            bac[r.programme or '—'][1] += 1
    moyennes = [{'programme': p, 'moyenne': round(s / n, 2), 'n': n} for p, (s, n) in sorted(bac.items())]
    return {
        'nationalites': _par_valeur(lignes_insc, 'nationalite', _nationalite),
        'options': _par_valeur(lignes_insc, 'option', _option),
        'centres': _par_valeur(lignes_insc, 'centre', _texte, limite=15),
        'ecoles': _par_valeur(lignes_insc, 'ecole', _texte, limite=15),
        'bac': moyennes, 'taux_bac': pct(sum(m['n'] for m in moyennes), len(lignes_insc)),
    }


def redoublants(lignes_insc):
    """Redoublants par niveau et par programme (H/F) et taux de redoublement."""
    redoub = [r for r in lignes_insc if r.redouble]
    return {'niveaux': _avec_taux(lignes_insc, redoub, 'niveau'),
            'programmes': _avec_taux(lignes_insc, redoub, 'programme'), 'nb': len(redoub)}


def handicap(lignes_insc):
    """Inscrits en situation de handicap : par type, par niveau et par programme (H/F, part des inscrits)."""
    from .models import Etudiant
    types = dict(Etudiant.TYPE_HANDICAP_CHOICES)
    concernes = [r for r in lignes_insc if r.handicap]
    par_type = defaultdict(lambda: [0, 0])
    for r in concernes:
        par_type[types.get(r.type_handicap, 'Type non précisé')][_f(r)] += 1
    ordre = list(types.values()) + ['Type non précisé']
    lignes_type = [Ligne(k, *par_type[k]) for k in ordre if k in par_type]
    total = _total(lignes_type)
    total.extra['taux'] = pct(total.total, len(lignes_insc))
    return {'nb': len(concernes), 'taux': pct(len(concernes), len(lignes_insc)), 'total': total,
            'types': (lignes_type, total), 'niveaux': _avec_taux(lignes_insc, concernes, 'niveau'),
            'programmes': _avec_taux(lignes_insc, concernes, 'programme')}


def reussite(resultats):
    """Résultats annuels calculés (notes) : admis (moyenne ≥ seuil) par programme et par niveau, H/F."""
    seuil = regles()['seuil_validation']
    lignes_res = list(resultats.values_list('filiere__nom', 'licence__nom', 'etudiant__sexe', 'moyenne').order_by())

    def tableau(index):
        d = defaultdict(lambda: {'h': 0, 'f': 0, 'admis_h': 0, 'admis_f': 0})
        for ligne in lignes_res:
            s = 'f' if ligne[2] == 'F' else 'h'
            x = d[ligne[index] or '—']
            x[s] += 1
            if ligne[3] is not None and ligne[3] >= seuil:
                x[f'admis_{s}'] += 1
        lignes = [Ligne(k, v['h'], v['f'], admis_h=v['admis_h'], admis_f=v['admis_f'],
                        taux=pct(v['admis_h'] + v['admis_f'], v['h'] + v['f']),
                        taux_h=pct(v['admis_h'], v['h']), taux_f=pct(v['admis_f'], v['f'])) for k, v in sorted(d.items())]
        somme = _total(lignes)
        ah, af = sum(l.extra['admis_h'] for l in lignes), sum(l.extra['admis_f'] for l in lignes)
        somme.extra.update(admis_h=ah, admis_f=af, taux=pct(ah + af, somme.total), taux_h=pct(ah, somme.h),
                           taux_f=pct(af, somme.f))
        return lignes, somme
    return {'programmes': tableau(0), 'niveaux': tableau(1), 'seuil': seuil, 'nb': len(lignes_res)}


def soutenances(qs_soutenances):
    """Soutenances par programme : nombre (H/F), notées, validées, moyenne."""
    seuil = regles()['seuil_validation']
    d = defaultdict(lambda: {'h': 0, 'f': 0, 'notees': 0, 'validees': 0, 'somme': 0.0})
    for prog, sexe, note in qs_soutenances.values_list('filiere__nom', 'etudiant__sexe', 'note_finale').order_by():
        x = d[prog or '—']
        x['f' if sexe == 'F' else 'h'] += 1
        if note is not None:
            x['notees'] += 1
            x['somme'] += note
            x['validees'] += note >= seuil
    lignes = [Ligne(k, v['h'], v['f'], notees=v['notees'], validees=v['validees'], taux=pct(v['validees'], v['notees']),
                    moyenne=round(v['somme'] / v['notees'], 2) if v['notees'] else None) for k, v in sorted(d.items())]
    somme = _total(lignes)
    notees, validees = sum(v['notees'] for v in d.values()), sum(v['validees'] for v in d.values())
    somme.extra.update(notees=notees, validees=validees, taux=pct(validees, notees),
                       moyenne=round(sum(v['somme'] for v in d.values()) / notees, 2) if notees else None)
    return {'lignes': lignes, 'total': somme}


def feminisation(lignes_programmes):
    """Programmes classés du plus au moins féminisé."""
    return sorted(lignes_programmes, key=lambda l: (-l.pct_f, -l.total))


def calculer(base, session_id, resultats, qs_soutenances):
    """Toutes les statistiques. `base` : inscriptions de toutes les sessions (pour l'évolution)."""
    toutes = charger(base)
    lignes_insc = [r for r in toutes if r.session_id == session_id] if session_id else toutes
    par_programme = repartition(lignes_insc, 'programme')
    return {
        'niveaux': repartition(lignes_insc, 'niveau'), 'programmes': par_programme,
        'promotions': repartition(lignes_insc, 'promotion'), 'croise': croise(lignes_insc),
        'evolution': evolution(toutes), 'feminisation': feminisation(par_programme[0]),
        'ages': ages(lignes_insc), 'origine': origine(lignes_insc), 'redoublants': redoublants(lignes_insc),
        'handicap': handicap(lignes_insc),
        'reussite': reussite(resultats), 'soutenances': soutenances(qs_soutenances),
    }


# ---------------------------------------------------------------- export Excel

def classeur(s, filtres_texte):
    """Classeur Excel : une feuille par statistique (s = contexte calculé par la vue)."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    gras, entete = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='41627C')

    def feuille(titre, colonnes, lignes):
        ws = wb.create_sheet(titre[:31])
        ws.append([f"{titre} · {filtres_texte}"])
        ws['A1'].font = Font(bold=True, size=12)
        ws.append([])
        ws.append(colonnes)
        for c in ws[3]:
            c.font, c.fill = gras, entete
        for l in lignes:
            ws.append(l)
        for col in ws.iter_cols(min_row=3):
            ws.column_dimensions[col[0].column_letter].width = max(10, min(50, max(len(str(c.value or ''))
                                                                                   for c in col) + 2))
        ws.freeze_panes = 'A4'

    hf = lambda l: [l.libelle, l.h, l.f, l.total, l.pct_f]                       # noqa: E731
    HF = ['Hommes', 'Femmes', 'Total', '% femmes']
    for titre, cle_, nom in (('Par niveau', 'niveaux', 'Niveau'), ('Par programme', 'programmes', 'Programme'),
                             ('Par promotion', 'promotions', 'Promotion')):
        lignes, total = s[cle_]
        feuille(titre, [nom] + HF, [hf(l) for l in lignes] + [hf(total)])
    c = s['croise']
    feuille('Programme x niveau', ['Programme'] + [f"{n} ({x})" for n in c['niveaux'] for x in ('H', 'F')] + ['Total'],
            [[l['programme']] + [v for cel in l['cellules'] for v in ((cel.h, cel.f) if cel else (0, 0))]
             + [l['total'].total] for l in c['lignes']])
    feuille('Évolution', ['Session'] + HF + ['Nouveaux inscrits', 'Réinscrits'],
            [hf(l) + [l.extra['nouveaux'], l.extra['reinscrits']] for l in s['evolution']])
    feuille('Féminisation', ['Programme'] + HF, [hf(l) for l in s['feminisation']])
    a = s['ages']
    feuille('Âge', ["Tranche d'âge"] + HF, [hf(l) for l in a['lignes']] + [hf(a['total'])]
            + [[], ['Âge moyen', a['moyenne_h'], a['moyenne_f'], a['moyenne']],
               ['Fiches renseignées (%)', a['taux']], ['Dates incohérentes', a['incoherentes']]])
    o = s['origine']
    for titre, cle_ in (('Nationalité', 'nationalites'), ('Option du bac', 'options'), ('Centre du bac', 'centres'),
                        ("École d'origine", 'ecoles')):
        feuille(titre, [titre] + HF, [hf(l) for l in o[cle_]['lignes']] + [hf(o[cle_]['total']),
                                                                             [], ['Renseigné (%)', o[cle_]['taux']]])
    feuille('Moyenne au bac', ['Programme', 'Moyenne au bac', 'Étudiants renseignés'],
            [[b['programme'], b['moyenne'], b['n']] for b in o['bac']])
    for titre, cle_ in (('Redoublants par niveau', 'niveaux'), ('Redoublants par programme', 'programmes')):
        lignes, total = s['redoublants'][cle_]
        feuille(titre, ['', 'Hommes', 'Femmes', 'Total', '% femmes', 'Taux de redoublement (%)'],
                [hf(l) + [l.extra['taux']] for l in lignes] + [hf(total) + [total.extra['taux']]])
    ha = s['handicap']
    lignes, total = ha['types']
    feuille('Handicap par type', ['Type'] + HF, [hf(l) for l in lignes] + [hf(total), [],
                                                  ['Part des inscrits (%)', ha['taux']]])
    for titre, cle_ in (('Handicap par niveau', 'niveaux'), ('Handicap par programme', 'programmes')):
        lignes, total = ha[cle_]
        feuille(titre, ['', 'Hommes', 'Femmes', 'Total', '% femmes', 'Part des inscrits (%)'],
                [hf(l) + [l.extra['taux']] for l in lignes] + [hf(total) + [total.extra['taux']]])
    for titre, cle_ in (('Réussite par programme', 'programmes'), ('Réussite par niveau', 'niveaux')):
        lignes, total = s['reussite'][cle_]
        feuille(titre, ['', 'Hommes', 'Femmes', 'Total', 'Admis H', 'Admis F', 'Taux (%)', 'Taux H (%)', 'Taux F (%)'],
                [[l.libelle, l.h, l.f, l.total, l.extra['admis_h'], l.extra['admis_f'], l.extra['taux'],
                  l.extra['taux_h'], l.extra['taux_f']] for l in lignes + [total]])
    so = s['soutenances']
    feuille('Soutenances', ['Programme', 'Hommes', 'Femmes', 'Total', 'Notées', 'Validées', 'Taux (%)', 'Moyenne'],
            [[l.libelle, l.h, l.f, l.total, l.extra['notees'], l.extra['validees'], l.extra['taux'], l.extra['moyenne']]
             for l in so['lignes'] + [so['total']]])
    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()
