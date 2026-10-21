"""
Import des matières (Excel .xlsx ou CSV) : scolarité et chefs de programme.

Une ligne = une matière. Colonnes repérées par leur nom, dans n'importe quel ordre :
  - obligatoires : nom, programme, licence, semestre, session ;
  - facultatives : code, credit (vide → 6), cycle (vide → cycle du programme), promotion, description.
Contrôles : le programme, la licence, le semestre, la session et la promotion existent ;
le semestre appartient à la licence ; les lignes vides sont ignorées.

Une matière est identifiée par son nom (sans tenir compte des majuscules ni des accents),
son programme et sa session :
  - déjà en base → MISE À JOUR (code, crédits, cycle, licence, semestre, promotion, description ;
    un code, un cycle, une promotion ou une description vides dans le fichier ne l'effacent pas) ;
  - répétée dans le fichier → la ligne est rejetée.

Chef de programme : si une seule ligne porte un programme qui n'est pas le sien, TOUT
l'import est refusé et rien n'est enregistré.
"""
import io

from django.db import transaction

from .importer import (_index_programmes, cle, cle_compacte, cle_licence, cle_session, en_cycle, lire_lignes, texte)
from .models import Licence, Matiere, Promotion, Semestre, Session

OBLIGATOIRES = ['nom', 'programme', 'licence', 'semestre', 'session']
COLONNES = OBLIGATOIRES + ['code', 'credit', 'cycle', 'promotion', 'description']
CREDIT_PAR_DEFAUT = 6

# En-têtes alternatifs (après normalisation : minuscules, sans accents, « _ » pour les espaces)
ALIAS = {
    'matiere': 'nom', 'matieres': 'nom', 'nom_matiere': 'nom', 'intitule': 'nom', 'libelle': 'nom',
    'credits': 'credit', 'code_matiere': 'code', 'filiere': 'programme', 'module': 'semestre',
    'niveau': 'licence', 'annee': 'session', 'annee_academique': 'session',
}
LONGUEURS = {'nom': Matiere._meta.get_field('nom').max_length, 'code': Matiere._meta.get_field('code').max_length}


class ImportRefuse(ValueError):
    """Tout l'import est refusé (ex. programme hors du périmètre d'un chef de programme)."""


class RapportMatieres:
    def __init__(self):
        self.lignes = 0
        self.creees = 0
        self.mises_a_jour = 0
        self.inchangees = 0
        self.erreurs = []           # [(numéro de ligne, message)]
        self.simulation = False
        self.colonnes_ignorees = []
        self.reparations = 0

    @property
    def nb_erreurs(self):
        return len(self.erreurs)

    def erreur(self, ligne, message):
        self.erreurs.append((ligne, message))


def en_credit(v):
    s = texte(v).replace(',', '.')
    if not s:
        return CREDIT_PAR_DEFAUT
    try:
        f = float(s)
    except ValueError:
        raise ValueError(f"crédit invalide « {s} » (nombre entier de 1 à 60)")
    if not f.is_integer() or not 1 <= f <= 60:
        raise ValueError(f"crédit invalide « {s} » (nombre entier de 1 à 60)")
    return int(f)


def _index_semestres():
    """Semestres par nom (« SEMESTRE 1 », « Module 1 ») ou par abrégé (« S1 »), sans espaces ni casse."""
    index = {}
    for s in Semestre.objects.select_related('licence'):
        index[cle_compacte(s.nom)] = s
    for s in Semestre.objects.select_related('licence'):
        if s.abrege:
            index.setdefault(cle_compacte(s.abrege), s)
    return index


def importer_matieres(fichier, profil, simulation=False):
    rapport = RapportMatieres()
    rapport.simulation = simulation
    entetes, lignes = lire_lignes(fichier)
    rapport.reparations = getattr(fichier, 'gped_reparations', 0)
    entetes = [ALIAS.get(e, e) for e in entetes]
    rapport.colonnes_ignorees = [e for e in entetes if e and e not in COLONNES]
    manquantes = [c for c in OBLIGATOIRES if c not in entetes]
    if manquantes:
        raise ValueError("Colonnes obligatoires manquantes : " + ", ".join(f"« {c} »" for c in manquantes) + ".")

    autorises = None if profil.est_scolarite else set(profil.programme_ids)
    programmes = _index_programmes()
    licences = {cle_licence(l.nom): l for l in Licence.objects.all()}
    semestres = _index_semestres()
    sessions = {cle_session(s.annee): s for s in Session.objects.all()}
    promotions = {cle_compacte(p.nom): p for p in Promotion.objects.all()}

    # 1) Lecture et contrôle de chaque ligne ---------------------------------
    valides, vues, hors_perimetre = [], {}, []
    for numero, brut in enumerate(lignes, start=2):
        if brut is None or all(texte(c) == '' for c in brut):
            continue                                   # ligne vide : ignorée
        rapport.lignes += 1
        ligne = {entetes[i]: brut[i] for i in range(min(len(entetes), len(brut))) if entetes[i]}

        # Chef de programme : le programme de CHAQUE ligne doit être l'un des siens
        # (un programme vide est une simple ligne incomplète, signalée plus bas)
        if autorises is not None and texte(ligne.get('programme')):
            prog = programmes.get(cle(ligne.get('programme')))
            if prog is None or prog.pk not in autorises:
                hors_perimetre.append((numero, texte(ligne.get('programme'))))
                continue
        try:
            vides = [c for c in OBLIGATOIRES if not texte(ligne.get(c))]
            if vides:
                raise ValueError("valeur manquante : " + ", ".join(vides))
            nom = ' '.join(texte(ligne['nom']).split())
            code = texte(ligne.get('code')) or None
            for champ, valeur in (('nom', nom), ('code', code or '')):
                if len(valeur) > LONGUEURS[champ]:
                    raise ValueError(f"« {champ} » dépasse {LONGUEURS[champ]} caractères")
            programme = programmes.get(cle(ligne['programme']))
            if programme is None:
                raise ValueError(f"programme inconnu « {texte(ligne['programme'])} »")
            licence = licences.get(cle_licence(ligne['licence']))
            if licence is None:
                raise ValueError(f"licence inconnue « {texte(ligne['licence'])} »")
            semestre = semestres.get(cle_compacte(ligne['semestre']))
            if semestre is None:
                raise ValueError(f"semestre inconnu « {texte(ligne['semestre'])} »")
            if semestre.licence_id != licence.pk:
                raise ValueError(f"{semestre.nom} n'appartient pas à {licence.nom} (il est en {semestre.licence.nom})")
            session = sessions.get(cle_session(ligne['session']))
            if session is None:
                raise ValueError(f"session inconnue « {texte(ligne['session'])} »")
            promotion = None
            if texte(ligne.get('promotion')):
                promotion = promotions.get(cle_compacte(ligne['promotion']))
                if promotion is None:
                    raise ValueError(f"promotion inconnue « {texte(ligne['promotion'])} »")
            credit = en_credit(ligne.get('credit'))
            cycle = en_cycle(ligne.get('cycle'))           # vide → None (cycle du programme plus bas)
            if cycle is None and not programme.cycle:
                raise ValueError(f"cycle manquant : le programme {programme} n'en a pas, précisez-le dans la colonne « cycle »")
            description = texte(ligne.get('description')) or None
        except ValueError as e:
            rapport.erreur(numero, str(e))
            continue

        identite = (cle(nom), programme.pk, session.pk)
        if identite in vues:
            rapport.erreur(numero, f"doublon dans le fichier : même matière que la ligne {vues[identite]} "
                                   f"({nom} · {programme} · {session})")
            continue
        vues[identite] = numero
        valides.append((identite, {'nom': nom, 'code': code, 'credit': credit, 'cycle': cycle, 'filiere': programme,
                                   'licence': licence, 'semestre': semestre, 'session': session,
                                   'promotion': promotion, 'description': description}))

    if hors_perimetre:
        mes_programmes = ', '.join(p.nom for p in programmes.values() if p.pk in autorises) or 'aucun'
        details = ', '.join(f"ligne {n} ({p})" for n, p in hors_perimetre[:10])
        if len(hors_perimetre) > 10:
            details += f"… et {len(hors_perimetre) - 10} autre(s)"
        raise ImportRefuse(
            f"Import refusé : le fichier contient des matières d'un programme qui n'est pas le vôtre "
            f"({mes_programmes}) : {details}. Aucune matière n'a été enregistrée.")
    if not valides:
        return rapport

    # 2) Écriture en base : création ou mise à jour ---------------------------
    with transaction.atomic():
        existantes = {}
        filieres = {d['filiere'].pk for _, d in valides}
        sessions_ids = {d['session'].pk for _, d in valides}
        for m in Matiere.objects.filter(filiere_id__in=filieres, session_id__in=sessions_ids):
            existantes[(cle(m.nom), m.filiere_id, m.session_id)] = m

        a_creer, a_maj = [], []
        champs_maj = ['code', 'credit', 'cycle', 'licence', 'semestre', 'promotion', 'description']
        for identite, d in valides:
            m = existantes.get(identite)
            if m is None:
                a_creer.append(Matiere(**{**d, 'cycle': d['cycle'] or d['filiere'].cycle}))
                continue
            avant = [getattr(m, f'{c}_id' if c in ('licence', 'semestre', 'promotion') else c) for c in champs_maj]
            m.credit, m.licence, m.semestre = d['credit'], d['licence'], d['semestre']
            for c in ('code', 'cycle', 'promotion', 'description'):   # vides dans le fichier : on garde l'existant
                if d[c] is not None:
                    setattr(m, c, d[c])
            if not m.cycle:                                  # ni dans le fichier ni en base : celui du programme
                m.cycle = d['filiere'].cycle
            apres = [getattr(m, f'{c}_id' if c in ('licence', 'semestre', 'promotion') else c) for c in champs_maj]
            if apres != avant:
                a_maj.append(m)
            else:
                rapport.inchangees += 1
        Matiere.objects.bulk_create(a_creer, batch_size=500)
        if a_maj:
            Matiere.objects.bulk_update(a_maj, champs_maj, batch_size=500)
        rapport.creees, rapport.mises_a_jour = len(a_creer), len(a_maj)
        rapport.erreurs.sort()
        if simulation:
            transaction.set_rollback(True)
    return rapport


def modele_matieres_xlsx():
    """Fichier Excel avec les en-têtes attendus (obligatoires en jaune) et deux lignes d'exemple."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Matieres'
    ws.append(COLONNES)
    ws.append(['Algorithmique et programmation', 'GENIE INFORMATIQUE', 'LICENCE 1', 'SEMESTRE 1', '2025 - 2026',
               'INF101', 6, 'Licence', 'P20', ''])
    ws.append(['Mathématiques discrètes', 'GENIE INFORMATIQUE', 'LICENCE 1', 'S1', '2025 - 2026',
               'MAT101', '', '', '', 'Crédit vide : 6 ; cycle vide : celui du programme'])
    for cell in ws[1]:
        obligatoire = cell.value in OBLIGATOIRES
        cell.font = Font(bold=True, color='1F2937' if obligatoire else 'FFFFFF')
        cell.fill = PatternFill('solid', fgColor='FBBD08' if obligatoire else '41627C')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, max(len(str(c.value or '')) for c in col) + 2)
    ws.freeze_panes = 'A2'
    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()
