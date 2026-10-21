"""
Menu Soutenances : liste des soutenances, projets de soutenance, saisie des notes (avec
verrouillage), calendrier, import. Règles : voir pfe/soutenances.py.
Droits : scolarité (tout) et chefs de programme (leurs programmes).
"""
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Prefetch
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST

from .forms import ImportSoutenancesForm, ProjetSoutenanceForm, SoutenanceForm
from .import_matieres import ImportRefuse
from .import_soutenances import OBLIGATOIRES, importer_soutenances, modele_soutenances_xlsx
from .models import Etudiant, Inscription, ProjetSoutenance, Promotion, Resultat, Session, Soutenance
from .notes import en_note, recalculer, regles
from .roles import gestion_matieres_requise, get_profil, programmes_du_profil
from .soutenances import (CHAMPS_NOTE, conflits, deverrouiller, enregistrer_notes, filtrer, notes_de,
                          peut_modifier_note, projets_visibles, recherche_projets, recherche_soutenances,
                          soutenances_visibles, verrouiller)

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

STATUTS_SOUTENANCE = [('sans_note', 'Sans note'), ('validee', 'Validées (≥ 5)'), ('non_validee', 'Non validées'),
                      ('session', 'Validées en session'), ('sans_projet', 'Sans projet'),
                      ('verrouillee', 'Notes verrouillées'), ('ouverte', 'Notes non verrouillées')]


def _entier(valeur):
    return int(valeur) if str(valeur or '').isdigit() else None


def _session_courante():
    return Session.objects.filter(en_cours=True).first() or Session.objects.order_by('-annee').first()


def _filtres(request, profil, session_par_defaut=False):
    """Session (« Toutes » par défaut, ou la session en cours), programme, promotion, statut, recherche."""
    g = request.GET
    programmes = list(programmes_du_profil(profil))
    if 'session' in g or not session_par_defaut:
        session_id = _entier(g.get('session'))
    else:
        courante = _session_courante()
        session_id = courante.pk if courante else None
    f = {'session': session_id, 'programme': _entier(g.get('programme')), 'promotion': _entier(g.get('promotion')),
         'projet': _entier(g.get('projet')), 'statut': g.get('statut') or '', 'q': (g.get('q') or '').strip()}
    if len(programmes) == 1:
        f['programme'] = programmes[0].pk
    if f['programme'] and f['programme'] not in {p.pk for p in programmes}:
        f['programme'] = None
    return f, {'filtres': f, 'sessions': Session.objects.order_by('-annee'), 'programmes': programmes,
               'promotions': Promotion.objects.order_by('nom')}


def _params(request, *retirer):
    p = request.GET.copy()
    for cle_ in ('page', 'export', 'format') + retirer:
        p.pop(cle_, None)
    return p.urlencode()


def _retour(request, defaut):
    suivant = request.POST.get('suivant') or request.GET.get('suivant')
    if suivant and url_has_allowed_host_and_scheme(suivant, allowed_hosts={request.get_host()}):
        return suivant
    return defaut


def _avec_resultats(soutenances):
    """Ajoute à chaque soutenance son résultat de cycle (moyenne du cycle et moyenne générale)."""
    cycles = {(r.etudiant_id, r.cycle): r for r in Resultat.objects.filter(
        niveau=Resultat.CYCLE, etudiant_id__in=[s.etudiant_id for s in soutenances])}
    for s in soutenances:
        s.resultat = cycles.get((s.etudiant_id, s.cycle))
    return soutenances


# ------------------------------------------------------------------ liste des soutenances

def _soutenances_filtrees(profil, f):
    qs = recherche_soutenances(filtrer(soutenances_visibles(profil), f), f['q'])
    seuil = regles()['seuil_validation']
    statut = f['statut']
    if statut == 'sans_note':
        qs = qs.filter(note_finale__isnull=True)
    elif statut == 'validee':
        qs = qs.filter(note_finale__gte=seuil)
    elif statut == 'non_validee':
        qs = qs.filter(note_finale__lt=seuil)
    elif statut == 'session':
        qs = qs.filter(par_session=True)
    elif statut == 'sans_projet':
        qs = qs.filter(projet__isnull=True)
    elif statut == 'verrouillee':
        qs = qs.filter(verrouillee=True)
    elif statut == 'ouverte':
        qs = qs.filter(verrouillee=False)
    else:
        f['statut'] = ''
    return qs.order_by('-session__annee', 'filiere__nom', 'promotion__nom', 'etudiant__nom', 'etudiant__prenom', 'pk')


@login_required
@gestion_matieres_requise
def soutenances_liste(request):
    profil = get_profil(request)
    f, contexte = _filtres(request, profil)
    qs = _soutenances_filtrees(profil, f)
    if request.GET.get('export') == 'xlsx':
        return _export_soutenances(qs)
    page = Paginator(qs, 50).get_page(request.GET.get('page'))
    liste = _avec_resultats(list(page))
    return render(request, 'pfe/soutenances_liste.html', {
        **contexte, 'page': page, 'liste': liste, 'params': _params(request), 'statuts': STATUTS_SOUTENANCE,
        'poids': regles(), 'seuil': regles()['seuil_validation'], 'actif': 'sout_liste',
    })


def _export_soutenances(qs):
    import io
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet('Soutenances')
    entetes = ['Matricule', 'Nom', 'Prénom', 'Programme', 'Cycle', 'Session', 'Promotion', 'Projet', 'Thème',
               'Encadreur', 'Chef de département', 'Maître de stage', 'Entreprise', 'Président du jury',
               'Examinateur', 'Rapporteur', 'Date', 'Heure', 'Salle', 'Statut du projet', 'Note', 'Note de session',
               'Note retenue', 'Mention', 'Validée en session', 'Note verrouillée']
    ligne = []
    for titre in entetes:
        cell = WriteOnlyCell(ws, value=titre)
        cell.font, cell.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='41627C')
        ligne.append(cell)
    ws.append(ligne)
    for s in qs.iterator(chunk_size=2000):
        e, p = s.etudiant, s.projet
        projet = [p.intitule, p.theme, p.encadreur, p.chef_departement, p.maitre_stage, p.entreprise,
                  p.president_jury, p.examinateur, p.rapporteur, p.date, p.heure, p.salle,
                  p.get_statut_display()] if p else [''] * 13
        ws.append([e.matricule, e.nom, e.prenom, str(s.filiere), s.cycle, s.session.annee, str(s.promotion),
                   *projet, s.note, s.note_session, s.note_finale, s.mention, 'oui' if s.par_session else 'non',
                   'oui' if s.verrouillee else 'non'])
    flux = io.BytesIO()
    wb.save(flux)
    reponse = HttpResponse(flux.getvalue(), content_type=XLSX)
    reponse['Content-Disposition'] = 'attachment; filename="soutenances_geped.xlsx"'
    return reponse


# ------------------------------------------------------------------ fiche d'une soutenance

def _formulaire_soutenance(request, instance=None):
    profil = get_profil(request)
    if request.method == 'POST':
        form = SoutenanceForm(request.POST, request.FILES, instance=instance, profil=profil)
        if form.is_valid():
            with transaction.atomic():
                s = form.save(utilisateur=request.user)
                recalculer([s.etudiant_id])
            if form.avertissement:
                messages.warning(request, form.avertissement)
            messages.success(request, f"Soutenance de {s.etudiant.prenom} {s.etudiant.nom} enregistrée.")
            if 'continuer' in request.POST:
                return redirect('soutenance_creer')
            return redirect(_retour(request, reverse('soutenances')))
    else:
        initial = {}
        if instance is None and request.GET.get('projet'):
            p = projets_visibles(profil).filter(pk=_entier(request.GET['projet'])).first()
            if p:
                initial = {'projet': p.pk, 'filiere': p.filiere_id, 'session': p.session_id,
                           'promotion': p.promotion_id, 'cycle': p.cycle}
        form = SoutenanceForm(instance=instance, profil=profil, initial=initial)
    historique = list(instance.historique.select_related('utilisateur')[:20]) if instance else []
    retour = _retour(request, reverse('soutenances'))
    return render(request, 'pfe/soutenance_form.html', {
        'form': form, 'soutenance': instance, 'historique': historique, 'actif': 'sout_liste',
        'titre': 'Modifier la soutenance' if instance else 'Nouvelle soutenance',
        'retour': retour, 'suivant': retour if retour != reverse('soutenances') else '',
        'cycles_des_programmes': {str(p.pk): p.cycle or '' for p in programmes_du_profil(profil)},
    })


@login_required
@gestion_matieres_requise
def soutenance_creer(request):
    return _formulaire_soutenance(request)


@login_required
@gestion_matieres_requise
def soutenance_modifier(request, pk):
    s = get_object_or_404(soutenances_visibles(get_profil(request)), pk=pk)
    return _formulaire_soutenance(request, s)


@login_required
@gestion_matieres_requise
@require_POST
def soutenance_supprimer(request, pk):
    profil = get_profil(request)
    s = get_object_or_404(soutenances_visibles(profil), pk=pk)
    if not peut_modifier_note(profil, s):
        raise PermissionDenied("Note verrouillée : seule la scolarité peut supprimer cette soutenance.")
    etudiant_id = s.etudiant_id
    s.delete()
    recalculer([etudiant_id])
    messages.success(request, "Soutenance supprimée.")
    return redirect(_retour(request, reverse('soutenances')))


@login_required
@gestion_matieres_requise
def soutenance_etudiant(request):
    """Aide à la saisie : nom de l'étudiant et sa dernière inscription (programme, session, promotion)."""
    e = Etudiant.objects.filter(matricule=(request.GET.get('matricule') or '').strip()).first()
    if e is None:
        return JsonResponse({'trouve': False})
    insc = (Inscription.objects.filter(etudiant=e, filiere__isnull=False).select_related('session')
            .order_by('-session__annee').first())
    return JsonResponse({'trouve': True, 'nom': f"{e.prenom} {e.nom}", 'preinscrit': e.preinscrit,
                         'programme': insc.filiere_id if insc else None, 'session': insc.session_id if insc else None,
                         'promotion': insc.promotion_id if insc else None, 'cycle': insc.cycle if insc else None})


# ------------------------------------------------------------------ projets de soutenance

STATUTS_PROJET = ProjetSoutenance.STATUTS


@login_required
@gestion_matieres_requise
def projets_liste(request):
    profil = get_profil(request)
    f, contexte = _filtres(request, profil)
    qs = recherche_projets(filtrer(projets_visibles(profil), f), f['q'])
    if f['statut'] in dict(STATUTS_PROJET):
        qs = qs.filter(statut=f['statut'])
    else:
        f['statut'] = ''
    qs = (qs.annotate(nb_membres=Count('membres', distinct=True))
          .prefetch_related(Prefetch('membres', queryset=Soutenance.objects.select_related('etudiant')
                                     .order_by('etudiant__nom')))
          .order_by('-session__annee', 'filiere__nom', 'promotion__nom', 'intitule'))
    page = Paginator(qs, 30).get_page(request.GET.get('page'))
    return render(request, 'pfe/projets_liste.html', {
        **contexte, 'page': page, 'params': _params(request), 'statuts': STATUTS_PROJET, 'actif': 'sout_projets',
    })


def _formulaire_projet(request, projet=None):
    profil = get_profil(request)
    membres = list(projet.membres.select_related('etudiant').order_by('etudiant__nom')) if projet else []
    if request.method == 'POST':
        form = ProjetSoutenanceForm(request.POST, instance=projet, profil=profil)
        if form.is_valid():
            with transaction.atomic():
                p = form.save()
                retires = [s for s in membres if str(s.pk) in request.POST.getlist('retirer')]
                for s in retires:
                    if s.note_finale is None and not s.verrouillee:
                        s.delete()                     # soutenance créée seulement pour l'appartenance au projet
                    else:
                        Soutenance.objects.filter(pk=s.pk).update(projet=None)
                recalculer([s.etudiant_id for s in membres] + [e.pk for e, _ in form.a_ajouter])
            n = len(form.a_ajouter)
            messages.success(request, f"Projet « {p.intitule} » enregistré"
                                      + (f" : {n} membre{'s' if n > 1 else ''} ajouté{'s' if n > 1 else ''}" if n else '')
                                      + '.')
            if 'continuer' in request.POST:
                return redirect('projet_creer')
            return redirect('projet_modifier', pk=p.pk)
    else:
        initial = {}
        if projet is None:
            for cle_, champ in (('programme', 'filiere'), ('session', 'session'), ('promotion', 'promotion')):
                if _entier(request.GET.get(cle_)):
                    initial[champ] = _entier(request.GET[cle_])
        form = ProjetSoutenanceForm(instance=projet, profil=profil, initial=initial)
    return render(request, 'pfe/projet_form.html', {
        'form': form, 'projet': projet, 'membres': _avec_resultats(membres), 'actif': 'sout_projets',
        'titre': f"Projet « {projet.intitule} »" if projet else 'Nouveau projet de soutenance',
        'cycles_des_programmes': {str(p.pk): p.cycle or '' for p in programmes_du_profil(profil)},
    })


@login_required
@gestion_matieres_requise
def projet_creer(request):
    return _formulaire_projet(request)


@login_required
@gestion_matieres_requise
def projet_modifier(request, pk):
    return _formulaire_projet(request, get_object_or_404(projets_visibles(get_profil(request)), pk=pk))


@login_required
@gestion_matieres_requise
@require_POST
def projet_supprimer(request, pk):
    """Supprime le projet ; les soutenances notées restent (sans projet), les autres sont retirées."""
    p = get_object_or_404(projets_visibles(get_profil(request)), pk=pk)
    with transaction.atomic():
        p.membres.filter(note_finale__isnull=True, verrouillee=False).delete()
        p.membres.update(projet=None)
        p.delete()
    messages.success(request, f"Projet « {p.intitule} » supprimé (les notes déjà saisies sont conservées).")
    return redirect(_retour(request, reverse('projets_liste')))


@login_required
@gestion_matieres_requise
def projet_pv(request, pk):
    from .pdf_soutenances import pdf_pv
    p = get_object_or_404(projets_visibles(get_profil(request)), pk=pk)
    membres = list(p.membres.select_related('etudiant').order_by('etudiant__nom', 'etudiant__prenom'))
    reponse = HttpResponse(pdf_pv(p, membres), content_type='application/pdf')
    reponse['Content-Disposition'] = f'inline; filename="PV_soutenance_{p.pk}.pdf"'
    reponse['Cache-Control'] = 'no-store'
    return reponse


# ------------------------------------------------------------------ saisie des notes

@login_required
@gestion_matieres_requise
def soutenances_notes(request):
    """
    Grille des notes d'un programme (session, promotion, projet facultatifs), groupées par projet.
    Boutons : Enregistrer, Verrouiller les notes (après l'enregistrement), Déverrouiller (super utilisateur).
    """
    profil = get_profil(request)
    f, contexte = _filtres(request, profil, session_par_defaut=True)
    contexte.update({'actif': 'sout_notes', 'regles': regles()})
    if not (f['programme'] and f['session']):
        return render(request, 'pfe/soutenances_notes.html', contexte)
    qs = filtrer(soutenances_visibles(profil), f)
    if f['projet']:
        qs = qs.filter(projet_id=f['projet'])
    soutenances = list(qs.order_by('projet__intitule', 'etudiant__nom', 'etudiant__prenom'))
    erreurs, postees, erreur_motif = {}, {}, ''

    if request.method == 'POST':
        action = request.POST.get('action', 'enregistrer')
        if action == 'deverrouiller':
            if not request.user.is_superuser:
                raise PermissionDenied("Seul un super utilisateur peut déverrouiller des notes.")
            n = deverrouiller(soutenances)
            messages.success(request, f"{n} note{'s' if n > 1 else ''} déverrouillée{'s' if n > 1 else ''}.")
            return redirect(request.get_full_path())
        motif = (request.POST.get('motif') or '').strip()
        changements, verrouillees_modifiees = [], False
        for s in soutenances:
            if not peut_modifier_note(profil, s):
                continue
            valeurs = {}
            for c in CHAMPS_NOTE:
                brut = request.POST.get(f'{c}_{s.pk}', '')
                postees[(s.pk, c)] = brut
                try:
                    valeurs[c] = en_note(brut)
                except ValueError as e:
                    erreurs[(s.pk, c)] = str(e)
            if any((s.pk, c) in erreurs for c in CHAMPS_NOTE) or valeurs == notes_de(s):
                continue
            changements.append((s, valeurs))
            verrouillees_modifiees |= s.verrouillee
        if verrouillees_modifiees and not motif:
            erreur_motif = "Notes verrouillées modifiées : indiquez le motif (et joignez le justificatif)."
        if not erreurs and not erreur_motif:
            with transaction.atomic():
                n = enregistrer_notes(changements, request.user, motif=motif,
                                      justificatif=request.FILES.get('justificatif'))
                recalculer({s.etudiant_id for s, _ in changements})
                nv = verrouiller([s for s in soutenances], request.user) if action == 'verrouiller' else 0
            if n:
                messages.success(request, f"{n} note{'s' if n > 1 else ''} de soutenance enregistrée{'s' if n > 1 else ''}.")
            if action == 'verrouiller':
                messages.success(request, f"{nv} note{'s' if nv > 1 else ''} verrouillée{'s' if nv > 1 else ''} : "
                                          f"seule la scolarité peut encore les modifier, avec un motif.")
            elif not n:
                messages.info(request, "Aucune modification.")
            return redirect(request.get_full_path())
        # erreurs : les valeurs postées sont réaffichées, rien n'est enregistré

    # Lignes groupées par projet (les soutenances sans projet à la fin)
    groupes, courant = [], None
    for s in sorted(soutenances, key=lambda x: (x.projet is None, x.projet.intitule if x.projet else '',
                                                 x.etudiant.nom, x.etudiant.prenom)):
        if courant is None or courant['projet'] != s.projet:
            courant = {'projet': s.projet, 'lignes': []}
            groupes.append(courant)
        s.modifiable = peut_modifier_note(profil, s)
        s.champs = [{'nom': c, 'valeur': postees.get((s.pk, c), '' if getattr(s, c) is None else '%g' % getattr(s, c)),
                     'erreur': erreurs.get((s.pk, c))} for c in CHAMPS_NOTE]
        courant['lignes'].append(s)
    contexte.update({
        'groupes': groupes, 'nb': len(soutenances), 'erreurs': bool(erreurs), 'erreur_motif': erreur_motif,
        'motif': request.POST.get('motif', ''), 'nb_notees': sum(1 for s in soutenances if s.note_finale is not None),
        'nb_verrouillees': sum(1 for s in soutenances if s.verrouillee),
        'a_verrouiller': sum(1 for s in soutenances if s.note_finale is not None and not s.verrouillee),
        'projets_du_filtre': projets_visibles(profil).filter(filiere_id=f['programme'], session_id=f['session'],
                                                             **({'promotion_id': f['promotion']} if f['promotion'] else {}))
                                                     .order_by('intitule'),
        'programme_choisi': next((p for p in contexte['programmes'] if p.pk == f['programme']), None),
    })
    return render(request, 'pfe/soutenances_notes.html', contexte)


# ------------------------------------------------------------------ calendrier

@login_required
@gestion_matieres_requise
def soutenances_calendrier(request):
    profil = get_profil(request)
    f, contexte = _filtres(request, profil, session_par_defaut=True)
    qs = (filtrer(projets_visibles(profil), f)
          .prefetch_related(Prefetch('membres', queryset=Soutenance.objects.select_related('etudiant')
                                     .order_by('etudiant__nom'))))
    if f['statut'] in dict(STATUTS_PROJET):
        qs = qs.filter(statut=f['statut'])
    else:
        f['statut'] = ''
    projets = list(qs.filter(date__isnull=False).order_by('date', 'heure', 'salle', 'intitule'))
    sans_date = qs.filter(date__isnull=True).count()
    problemes = conflits(projets)
    for p in projets:
        p.conflits = problemes.get(p.pk, [])
    if request.GET.get('format') == 'pdf':
        from .pdf_soutenances import pdf_calendrier
        titre = ' · '.join(str(x) for x in (
            next((p for p in contexte['programmes'] if p.pk == f['programme']), None),
            Session.objects.filter(pk=f['session']).first() if f['session'] else None,
            Promotion.objects.filter(pk=f['promotion']).first() if f['promotion'] else None) if x)
        reponse = HttpResponse(pdf_calendrier(projets, titre), content_type='application/pdf')
        reponse['Content-Disposition'] = 'inline; filename="calendrier_soutenances.pdf"'
        return reponse
    jours = []
    for p in projets:
        if not jours or jours[-1]['date'] != p.date:
            jours.append({'date': p.date, 'projets': []})
        jours[-1]['projets'].append(p)
    return render(request, 'pfe/soutenances_calendrier.html', {
        **contexte, 'jours': jours, 'nb': len(projets), 'sans_date': sans_date,
        'nb_conflits': len(problemes), 'statuts': STATUTS_PROJET, 'params': _params(request), 'actif': 'sout_calendrier',
    })


# ------------------------------------------------------------------ import

@login_required
@gestion_matieres_requise
def soutenances_importer(request):
    profil = get_profil(request)
    rapport = None
    if request.method == 'POST':
        form = ImportSoutenancesForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                rapport = importer_soutenances(form.cleaned_data['fichier'], profil, request.user,
                                               simulation=form.cleaned_data['simulation'])
            except (ValueError, ImportRefuse) as e:
                form.add_error('fichier', str(e))
            else:
                if not rapport.simulation and (rapport.creees or rapport.mises_a_jour):
                    messages.success(request, "Import des soutenances terminé et enregistré.")
    else:
        form = ImportSoutenancesForm()
    return render(request, 'pfe/soutenances_importer.html', {
        'form': form, 'rapport': rapport, 'obligatoires': OBLIGATOIRES, 'actif': 'sout_importer',
        'mes_programmes': None if profil.est_scolarite else programmes_du_profil(profil),
    })


@login_required
@gestion_matieres_requise
def soutenances_modele(request):
    reponse = HttpResponse(modele_soutenances_xlsx(), content_type=XLSX)
    reponse['Content-Disposition'] = 'attachment; filename="modele_import_soutenances.xlsx"'
    return reponse
