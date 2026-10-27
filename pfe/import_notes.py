"""
Import des notes (Excel .xlsx ou CSV) : scolarité et chefs de programme. Barème sur 10.

Une ligne = une note d'un étudiant dans une matière. Colonnes repérées par leur nom :
  - obligatoires : matricule, matiere (nom ou code), session ;
  - facultatives : programme (si besoin de préciser), controle, tp, examen, moyenne, rattrapage.
Il faut au moins une note sur la ligne. « ABS » = 0 (un absent a 0).
La colonne moyenne est la moyenne déposée par le programme : elle est prise telle quelle ;
sinon la moyenne est calculée (voir pfe/notes.py). Si la moyenne est remplie et que contrôle,
TP et examen valent tous 0, ces 0 sont considérés comme vides (cases remplies par défaut).

La matière est cherchée parmi celles de l'inscription de l'étudiant dans cette session
(son programme et sa licence). Contrôles : étudiant inscrit (pas préinscrit), matière trouvée,
notes dans le barème. Une note déjà enregistrée est MISE À JOUR (et la modification tracée) ;
une ligne répétée dans le fichier est ignorée (doublon). Année en cours (délibérée ou non) : les
notes sont créées ou mises à jour. Année écoulée : les notes manquantes sont créées, mais une note
déjà enregistrée n'est modifiée que par un chef (scolarité ou département du programme).
PDF de l'enseignant (fiche de notes signée) : obligatoire pour un département ou un programme, et
pour la scolarité dès que l'import modifie des notes déjà enregistrées ; gardé dans le journal des
imports (ImportNotes) et comme justificatif de chaque note modifiée.
Chef de programme : si une seule ligne concerne un autre programme, TOUT l'import est refusé.

INSCRIPTION AUTOMATIQUE (scolarité seulement) : un étudiant qui a des notes dans une session où il n'a
AUCUNE inscription y est inscrit. La matière est cherchée parmi toutes celles de la session (nom ou
code ; la colonne « programme » lève le doute) et doit être UNIQUE ; toutes ses notes de la session
doivent mener au même programme et à la même licence. Promotion : celle prévue pour la licence dans la
session (Promotion_Licence_Session), sinon celle de sa dernière inscription ; cycle : celui de sa
dernière inscription, sinon celui du programme ; redoublant s'il avait déjà été inscrit dans cette
licence. Jamais pour un préinscrit, ni pour un étudiant déjà inscrit dans la session (matières en dette).
"""
import io
import re
from collections import defaultdict

from django.db import transaction

from .import_matieres import ImportRefuse
from .importer import _index_programmes, cle, cle_session, lire_lignes, texte
from .models import Etudiant, Inscription, Matiere, Note, Session
from .notes import CHAMPS_SAISIE, en_note, enregistrer_notes, recalculer, saisie_de, session_ecoulee, verifier_pdf

OBLIGATOIRES = ['matricule', 'matiere', 'session']
COLONNES = OBLIGATOIRES + ['programme'] + list(CHAMPS_SAISIE)
ALIAS = {
    'matieres': 'matiere', 'nom_matiere': 'matiere', 'code_matiere': 'matiere', 'intitule': 'matiere',
    'cc': 'controle', 'controle_continu': 'controle', 'devoir': 'controle',
    'travaux_pratiques': 'tp', 'exam': 'examen', 'examen_final': 'examen',
    'moy': 'moyenne', 'moyenne_matiere': 'moyenne', 'rattrapages': 'rattrapage', 'note_rattrapage': 'rattrapage',
}


def cle_intitule(texte_brut):
    """Intitulé sans casse, accents ni ponctuation : « Probabilité_Statistique » = « Probabilité Statistique »."""
    return ' '.join(re.sub(r'[^A-Z0-9]+', ' ', cle(texte_brut)).split())


class RapportNotes:
    def __init__(self):
        self.lignes = 0
        self.creees = 0
        self.mises_a_jour = 0
        self.inchangees = 0
        self.erreurs = []
        self.doublons = []
        self.simulation = False
        self.colonnes_ignorees = []
        self.reparations = 0
        self.journal = None         # ImportNotes enregistré (import réel)
        self.inscriptions_creees = []  # [(1re ligne, matricule, étudiant, session, licence, programme, promotion)]

    @property
    def nb_erreurs(self):
        return len(self.erreurs)

    @property
    def nb_doublons(self):
        return len(self.doublons)

    def erreur(self, ligne, message):
        self.erreurs.append((ligne, message))


def _inscrire(a_inscrire, etudiants):
    """
    Crée les inscriptions automatiques : {(étudiant, session): [(n° de ligne, matière)]}.
    Renvoie [(1re ligne, matricule, étudiant, session, licence, programme, promotion)].
    """
    from .models import Promotion_Licence_Session
    par_pk = {e.pk: e for e in etudiants.values()}
    prevues = defaultdict(set)                                 # (licence, session) -> {promotions prévues}
    for lic, ses, promo in Promotion_Licence_Session.objects.filter(
            session_id__in={s for _, s in a_inscrire}).values_list('licence_id', 'session_id', 'promotion_id'):
        prevues[(lic, ses)].add(promo)
    crees = []
    for (etu_id, session_id), lignes_auto in sorted(a_inscrire.items(), key=lambda x: min(n for n, _ in x[1])):
        numero, matiere = min(lignes_auto, key=lambda x: x[0])
        etu = par_pk[etu_id]
        derniere = (Inscription.objects.filter(etudiant_id=etu_id).select_related('promotion')
                    .order_by('-session__annee', '-licence__nom').first())
        promos = prevues.get((matiere.licence_id, session_id), set())
        promotion_id = (next(iter(promos)) if len(promos) == 1
                        else (derniere.promotion_id if derniere else None))
        insc = Inscription.objects.create(
            etudiant_id=etu_id, session_id=session_id, licence_id=matiere.licence_id, filiere_id=matiere.filiere_id,
            promotion_id=promotion_id, cycle=(derniere.cycle if derniere and derniere.cycle else matiere.filiere.cycle),
            redouble=Inscription.objects.filter(etudiant_id=etu_id, licence_id=matiere.licence_id).exclude(
                session_id=session_id).exists())
        insc = Inscription.objects.select_related('session', 'licence', 'filiere', 'promotion').get(pk=insc.pk)
        crees.append((numero, etu.matricule, f"{etu.prenom} {etu.nom}", insc.session, insc.licence, insc.filiere,
                       insc.promotion))
    return crees


def importer_notes(fichier, profil, utilisateur=None, simulation=False, pdf_enseignant=None):
    rapport = RapportNotes()
    rapport.simulation = simulation
    if pdf_enseignant:
        verifier_pdf(pdf_enseignant, "Le PDF de l'enseignant")
    elif not profil.est_scolarite:
        raise ValueError("Le PDF de l'enseignant (fiche de notes signée) est obligatoire pour importer des notes.")
    entetes, lignes = lire_lignes(fichier)
    rapport.reparations = getattr(fichier, 'gped_reparations', 0)
    entetes = [ALIAS.get(e, e) for e in entetes]
    rapport.colonnes_ignorees = [e for e in entetes if e and e not in COLONNES]
    manquantes = [c for c in OBLIGATOIRES if c not in entetes]
    if manquantes:
        raise ValueError("Colonnes obligatoires manquantes : " + ", ".join(f"« {c} »" for c in manquantes) + ".")
    if not any(c in entetes for c in CHAMPS_SAISIE):
        raise ValueError("Aucune colonne de note : ajoutez au moins « moyenne », « controle », « tp », « examen » "
                         "ou « rattrapage ».")

    autorises = None if profil.est_scolarite else set(profil.programme_ids)
    programmes = _index_programmes()
    sessions = {cle_session(s.annee): s for s in Session.objects.all()}

    # 1) Lecture des lignes ---------------------------------------------------
    brutes = []
    for numero, brut in enumerate(lignes, start=2):
        if brut is None or all(texte(c) == '' for c in brut):
            continue
        rapport.lignes += 1
        ligne = {entetes[i]: brut[i] for i in range(min(len(entetes), len(brut))) if entetes[i]}
        try:
            vides = [c for c in OBLIGATOIRES if not texte(ligne.get(c))]
            if vides:
                raise ValueError("valeur manquante : " + ", ".join(vides))
            session = sessions.get(cle_session(ligne['session']))
            if session is None:
                raise ValueError(f"session inconnue « {texte(ligne['session'])} »")
            programme = None
            if texte(ligne.get('programme')):
                programme = programmes.get(cle(ligne['programme']))
                if programme is None:
                    raise ValueError(f"programme inconnu « {texte(ligne['programme'])} »")
            valeurs = {}
            for c in CHAMPS_SAISIE:
                try:
                    valeurs[c] = en_note(ligne.get(c))
                except ValueError as e:
                    raise ValueError(f"{c} : {e}")
            if all(v is None for v in valeurs.values()):
                raise ValueError("aucune note sur la ligne")
            # Moyenne déposée avec contrôle, TP et examen à 0 : ces 0 ne sont que des cases remplies
            # par défaut dans le fichier → considérés comme vides
            composantes = [valeurs[c] for c in ('controle', 'tp', 'examen') if valeurs[c] is not None]
            if valeurs['moyenne'] is not None and composantes and all(v == 0 for v in composantes):
                valeurs.update(controle=None, tp=None, examen=None)
        except ValueError as e:
            rapport.erreur(numero, str(e))
            continue
        brutes.append((numero, texte(ligne['matricule']), session, programme, texte(ligne['matiere']), valeurs))

    # 2) Étudiants, inscriptions et matières concernés --------------------------
    matricules = list({b[1] for b in brutes})
    etudiants = {}
    for i in range(0, len(matricules), 500):
        etudiants.update({e.matricule: e for e in Etudiant.objects.filter(matricule__in=matricules[i:i + 500])})
    inscriptions = defaultdict(list)                       # (etudiant, session) -> [inscriptions]
    ids = [e.pk for e in etudiants.values()]
    for i in range(0, len(ids), 500):
        for insc in Inscription.objects.filter(etudiant_id__in=ids[i:i + 500], filiere__isnull=False):
            inscriptions[(insc.etudiant_id, insc.session_id)].append(insc)
    triplets = {(x.filiere_id, x.session_id, x.licence_id) for lst in inscriptions.values() for x in lst}
    matieres = defaultdict(dict)                           # triplet -> {clé nom/code: matière}
    if triplets:
        for m in Matiere.objects.filter(filiere_id__in={t[0] for t in triplets},
                                        session_id__in={t[1] for t in triplets},
                                        licence_id__in={t[2] for t in triplets}).select_related('filiere', 'semestre'):
            t = (m.filiere_id, m.session_id, m.licence_id)
            if t in triplets:
                matieres[t][cle_intitule(m.nom)] = m
                if m.code:
                    matieres[t].setdefault(cle_intitule(m.code), m)

    # Inscription automatique (scolarité) : matières de la session pour les étudiants sans inscription
    sans_inscription = {(e.pk, b[2].pk) for b in brutes if (e := etudiants.get(b[1])) is not None
                        and not e.preinscrit and not inscriptions.get((e.pk, b[2].pk))}
    matieres_session = defaultdict(lambda: defaultdict(list))   # session -> intitulé -> [matières]
    if profil.est_scolarite and sans_inscription:
        for m in Matiere.objects.filter(session_id__in={s for _, s in sans_inscription}).select_related(
                'filiere', 'licence', 'semestre'):
            matieres_session[m.session_id][cle_intitule(m.nom)].append(m)
            if m.code and cle_intitule(m.code) != cle_intitule(m.nom):
                matieres_session[m.session_id][cle_intitule(m.code)].append(m)

    def matiere_sans_inscription(numero, matricule, session, programme, nom_matiere):
        """Matière unique de la session (pour inscrire l'étudiant) ; None + erreur sinon."""
        trouvees = [m for m in matieres_session[session.pk].get(cle_intitule(nom_matiere), [])
                    if programme is None or m.filiere_id == programme.pk]
        if not trouvees:
            rapport.erreur(numero, f"{matricule} n'est pas inscrit en {session} et la matière « {nom_matiere} » "
                                   f"n'existe pas dans cette session" + (f" en {programme}" if programme else ''))
            return None
        if len({(m.filiere_id, m.licence_id) for m in trouvees}) > 1:
            ou = ', '.join(sorted({f"{m.filiere.nom} {m.licence.nom}" for m in trouvees})[:3])
            rapport.erreur(numero, f"{matricule} n'est pas inscrit en {session} : « {nom_matiere} » existe dans "
                                   f"plusieurs programmes ou licences ({ou}) ; ajoutez la colonne « programme »")
            return None
        return trouvees[0]

    valides, vues, hors_perimetre = [], {}, []
    a_inscrire = defaultdict(list)          # (étudiant, session) -> [(n° de ligne, matière)] : inscription automatique
    for numero, matricule, session, programme, nom_matiere, valeurs in brutes:
        etu = etudiants.get(matricule)
        if etu is None:
            rapport.erreur(numero, f"matricule inconnu « {matricule} »")
            continue
        if etu.preinscrit:
            rapport.erreur(numero, f"{matricule} est préinscrit : finalisez d'abord son inscription")
            continue
        inscrit = bool(inscriptions.get((etu.pk, session.pk)))
        candidates = [x for x in inscriptions.get((etu.pk, session.pk), [])
                      if programme is None or x.filiere_id == programme.pk]
        if not inscrit:
            if not profil.est_scolarite:
                rapport.erreur(numero, f"{matricule} n'est pas inscrit en {session} : faites inscrire l'étudiant "
                                       f"par la scolarité")
                continue
            matiere = matiere_sans_inscription(numero, matricule, session, programme, nom_matiere)
            if matiere is None:
                continue
            a_inscrire[(etu.pk, session.pk)].append((numero, matiere))
        elif not candidates:
            rapport.erreur(numero, f"{matricule} n'est pas inscrit en {session} en {programme}")
            continue
        else:
            intitule = cle_intitule(nom_matiere)
            matiere = next((matieres[(x.filiere_id, x.session_id, x.licence_id)][intitule]
                            for x in candidates if intitule in matieres[(x.filiere_id, x.session_id, x.licence_id)]),
                           None)
            if matiere is None:
                rapport.erreur(numero, f"matière « {nom_matiere} » introuvable pour {matricule} "
                                       f"(son programme et sa licence en {session})")
                continue
        if autorises is not None and matiere.filiere_id not in autorises:
            hors_perimetre.append((numero, matiere.filiere.nom))
            continue
        identite = (etu.pk, matiere.pk)
        if identite in vues:
            rapport.doublons.append((numero, f"{matricule} · {matiere.nom} : déjà présente à la ligne {vues[identite]} du fichier"))
            continue
        vues[identite] = numero
        valides.append((numero, etu, matiere, valeurs))

    # Inscription automatique : toutes les notes d'un étudiant dans la session doivent mener au même
    # programme et à la même licence ; sinon aucune inscription, et ses lignes de la session sont rejetées
    incoherents = set()
    for (etu_id, session_id), lignes_auto in a_inscrire.items():
        couples = {(m.filiere_id, m.licence_id) for _, m in lignes_auto}
        if len(couples) > 1:
            incoherents.add((etu_id, session_id))
            ou = ', '.join(sorted({f"{m.filiere.nom} {m.licence.nom}" for _, m in lignes_auto}))
            for numero, m in lignes_auto:
                rapport.erreur(numero, f"{m.session} : notes de plusieurs programmes ou licences ({ou}) pour un "
                                       f"étudiant non inscrit : inscrivez-le d'abord")
    if incoherents:
        valides = [v for v in valides if (v[1].pk, v[2].session_id) not in incoherents]
        a_inscrire = {k: v for k, v in a_inscrire.items() if k not in incoherents}

    if hors_perimetre:
        details = ', '.join(f"ligne {n} ({p})" for n, p in hors_perimetre[:10])
        if len(hors_perimetre) > 10:
            details += f"… et {len(hors_perimetre) - 10} autre(s)"
        raise ImportRefuse(f"Import refusé : le fichier contient des notes d'un programme qui n'est pas le vôtre : "
                           f"{details}. Aucune note n'a été enregistrée.")

    retenues = valides
    ecoulees = {}                                          # session -> année écoulée ?

    # 3) Écriture : création ou mise à jour, puis calcul des résultats ------------
    if retenues:
        with transaction.atomic():
            if a_inscrire:
                rapport.inscriptions_creees = _inscrire(a_inscrire, etudiants)
            existantes = {}
            paires_ids = {(e.pk, m.pk) for _, e, m, _ in retenues}
            for n in Note.objects.filter(etudiant_id__in={p[0] for p in paires_ids},
                                         matiere_id__in={p[1] for p in paires_ids}):
                existantes[(n.etudiant_id, n.matiere_id)] = n
            changements = []
            for numero, etu, matiere, valeurs in retenues:
                note = existantes.get((etu.pk, matiere.pk))
                if note is not None:
                    # Colonnes absentes du fichier : on garde ce qui était saisi
                    valeurs = {c: valeurs[c] if c in entetes else saisie_de(note)[c] for c in CHAMPS_SAISIE}
                    if saisie_de(note) == valeurs:
                        rapport.inchangees += 1
                        continue
                    if matiere.session_id not in ecoulees:
                        ecoulees[matiere.session_id] = session_ecoulee(matiere.session)
                    if ecoulees[matiere.session_id] and not profil.est_chef_de(matiere.filiere):
                        rapport.erreur(numero, f"{etu.matricule} · {matiere.nom} ({matiere.session}) : année écoulée, "
                                               f"la note déjà enregistrée ne peut être modifiée que par un chef")
                        continue
                    if not pdf_enseignant and not profil.superuser:       # super utilisateur : dispensé
                        rapport.erreur(numero, f"{etu.matricule} · {matiere.nom} : note déjà enregistrée, joignez le "
                                               f"PDF de l'enseignant pour la modifier")
                        continue
                changements.append((note, etu.pk, matiere, valeurs))
            nom_pdf = ''
            if pdf_enseignant and changements and not simulation:
                from django.utils import timezone
                from .models import stockage_prive
                nom_pdf = stockage_prive().save(
                    f"justificatifs/imports_notes/{timezone.now():%Y}/{pdf_enseignant.name}", pdf_enseignant)
            nom_fichier = getattr(fichier, 'name', '')
            rapport.creees, rapport.mises_a_jour = enregistrer_notes(
                changements, utilisateur, motif=f"Import du fichier {nom_fichier}".strip(), justificatif=nom_pdf)
            recalculer({etu.pk for _, etu, _, _ in retenues})
            if changements and not simulation:
                from .models import ImportNotes
                rapport.journal = ImportNotes.objects.create(
                    utilisateur=utilisateur if getattr(utilisateur, 'pk', None) else None, fichier=nom_fichier[:255],
                    pdf_enseignant=nom_pdf, creees=rapport.creees, modifiees=rapport.mises_a_jour,
                    programmes=', '.join(sorted({m.filiere.nom for _, _, m, _ in changements}))[:255])
            if simulation:
                transaction.set_rollback(True)
    rapport.erreurs.sort()
    rapport.doublons.sort()
    return rapport


def modele_notes_xlsx():
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Notes'
    ws.append(COLONNES)
    ws.append(['2304000123', 'Algorithmique et programmation', '2025 - 2026', '', 6, 7.5, 6.5, '', ''])
    ws.append(['2304000217', 'Algorithmique et programmation', '2025 - 2026', '', '', '', '', 6.75, ''])
    ws.append(['2304000328', 'INF101', '2025 - 2026', '', 3, '', 'ABS', '', 6])
    for cell in ws[1]:
        obligatoire = cell.value in OBLIGATOIRES
        cell.font = Font(bold=True, color='1F2937' if obligatoire else 'FFFFFF')
        cell.fill = PatternFill('solid', fgColor='FBBD08' if obligatoire else '41627C')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, max(len(str(c.value or '')) for c in col) + 2)
    ws.freeze_panes = 'A2'
    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()
